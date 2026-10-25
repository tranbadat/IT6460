"""Xuất master schedule dạng Excel (Gantt theo tuần, sprint, mốc, tải nguồn lực) từ du_lieu/lich.json.

Dùng: python xuat_gantt.py
Quy ước: chữ xanh dương là dữ liệu do lap_lich.py tính (không sửa tay, sửa dữ liệu đầu vào rồi chạy lại),
chữ đen là công thức. Mục cha WBS, tổng sprint, tải nguồn lực, số ngày làm việc đều là công thức.
Sau khi tạo: tính lại công thức (nếu máy có LibreOffice), đối chiếu giá trị Excel với lich.json, đăng ký phiên bản.
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'qlda-loi', 'scripts'))
import loi  # noqa: E402
import xlsx_chung as X  # noqa: E402

lich = loi.doc_json('du_lieu/lich.json')
da = loi.du_an()
wbs = loi.doc_json('du_lieu/wbs.json')
OUT = loi.duong_dan('ket_qua/04_Master_schedule/Master_schedule.xlsx')
os.makedirs(os.path.dirname(OUT), exist_ok=True)
D = lambda s: dt.date.fromisoformat(s)
wb = Workbook()
wb.properties.creator = wb.properties.lastModifiedBy = loi.du_an().get('nguoi_lap', '')

# ---------------- Tong_quan
ws = wb.active; ws.title = 'Tong_quan'
ws['A1'] = 'Master schedule'; ws['A1'].font = X.F_TEN
ws['A2'] = da['ten_du_an']; ws['A2'].font = X.F_THUONG
info = [('Ngày bắt đầu', D(lich['ngay_bat_dau']), X.NGAY, True), ('Ngày cut over', D(lich['cut_over']), X.NGAY, True),
        ('Số ngày làm việc', '=NETWORKDAYS(B4,B5,$F$4:$F$40)', '0', False), ('Tổng khối lượng, MD', "=SUMIFS(Lich_trinh!D:D,Lich_trinh!C:C,1)", X.SO2, False),
        ('Quỹ thời gian, MD', '=SUM(Phan_bo!E:E)', X.SO2, False), ('Tỷ lệ sử dụng quỹ', '=IF(B8=0,0,B7/B8)', X.PT, False),
        ('Hoàn thành giai đoạn Thực hiện (không tính gói dàn trải và gói Kết thúc)', D(lich['ket_thuc_du_kien']), X.NGAY, True),
        ('Dự trữ đến cut over, ngày làm việc (trước các gói gắn ngày cố định)', lich['du_tru_den_cut_over_ngay'], '0', True),
        ('Số sprint', '=COUNTIFS(Sprint!A:A,"S*",Sprint!A:A,"<>Sprint")', '0', False), ('Đường găng theo phụ thuộc, ngày', lich['do_dai_gang_ngay'], '0', True)]
if lich.get('nguong_su_dung_quy_toi_thieu') is not None:
    info.append(('Ngưỡng tỷ lệ sử dụng quỹ tối thiểu', float(lich['nguong_su_dung_quy_toi_thieu']), X.PT, True))
if lich.get('ngay_dong_du_an'):  # thêm cuối bảng để không đổi vị trí các ô công thức B4 đến B14
    info.append(('Ngày đóng dự án (hết giai đoạn Kết thúc)', D(lich['ngay_dong_du_an']), X.NGAY, True))
for i, (k, v, f, nhap) in enumerate(info, 4):
    X.o(ws, i, 1, k)
    X.o(ws, i, 2, v, fmt=f, nhap=nhap)
if lich.get('nguong_su_dung_quy_toi_thieu') is not None:  # tỷ lệ sử dụng quỹ (B9) dưới ngưỡng (B14) tô đỏ nhạt
    ws.conditional_formatting.add('B9', FormulaRule(formula=['B9<B14'], fill=PatternFill('solid', fgColor='F8CBAD')))
X.o(ws, 3, 6, 'Ngày nghỉ', dam=True)
for i, h in enumerate(lich['ngay_nghi'], 4):
    X.o(ws, i, 6, D(h), fmt=X.NGAY, nhap=True)
ws.column_dimensions['A'].width = 36; ws.column_dimensions['B'].width = 16; ws.column_dimensions['F'].width = 14
X.ghi_chu_bang(ws, 16, 'Chữ xanh dương: dữ liệu do bước lập lịch tính từ WBS, giao việc và lịch làm việc. Chữ đen: công thức. '
                       'Không sửa tay dữ liệu xanh dương. Khi cần thay đổi, sửa đầu vào (WBS, giao việc, mốc) và lập lại lịch.')

# ---------------- Lich_trinh
lt = wb.create_sheet('Lich_trinh')
cols = ['Mã WBS', 'Tên công việc', 'Mức', 'MD', 'Người thực hiện', 'Bắt đầu', 'Kết thúc', 'Số ngày làm việc', 'Giai đoạn', 'Sprint',
        'Phụ thuộc', 'Đường găng', 'Mục cha', 'Ràng buộc ngày']
widths = [9, 46, 5, 8, 24, 11, 11, 10, 11, 8, 14, 7, 8, 30]


def cot_giai_doan(g):
    """(Giai đoạn, Sprint) của một dòng: gói Thực hiện ghi sprint, gói quản lý dự án ghi giai đoạn thay cho sprint."""
    if 'giai_doan_hien_thi' not in g:  # lich.json cũ chưa có giai đoạn
        return '', (g['sprint_bat_dau'] if g['sprint_bat_dau'] == g['sprint_ket_thuc'] else f'{g["sprint_bat_dau"]}-{g["sprint_ket_thuc"]}')
    if g.get('giai_doan') == 'thuc_hien':
        return loi.GIAI_DOAN['thuc_hien'], g['giai_doan_hien_thi']
    return g['giai_doan_hien_thi'], ''


def van_ban_rang_buoc(rb):
    """Ràng buộc ngày của gói lá (lich.json rang_buoc) thành chuỗi cho cột Ràng buộc ngày."""
    mau = {'thoi_luong_toi_thieu': lambda x: f'thời lượng tối thiểu {x["so_ngay"]} ngày',
           'phe_duyet': lambda x: f'Chủ đầu tư phê duyệt {x["so_ngay"]} ngày',
           'uat': lambda x: f'kiểm thử chấp nhận {loi.ngay(x["tu"])} - {loi.ngay(x["den"])}',
           'khong_som_hon': lambda x: f'không sớm hơn {loi.ngay(x["ngay"])}',
           'cam_ket': lambda x: (f'trong khoảng cam kết {x["ma"]} {loi.ngay(x["tu"])} - {loi.ngay(x["den"])}' if x.get('trong_khoang')
                                 else f'sau cam kết {x["ma"]}' + (f' hạn {loi.ngay(x["han"])}' if x.get('han') else ''))}
    s = ', '.join(mau[x['loai']](x) for x in rb if x.get('loai') in mau)
    return s[:1].upper() + s[1:]
d0 = D(lich['ngay_bat_dau']); d1 = max([D(lich['cut_over'])] + [D(g['ket_thuc']) for g in lich['goi']]
                                         + ([D(lich['ngay_dong_du_an'])] if lich.get('ngay_dong_du_an') else []))  # đến hết giai đoạn Kết thúc
w0 = d0 - dt.timedelta(days=d0.weekday())
tuan = []
while w0 <= d1:
    tuan.append(w0); w0 += dt.timedelta(days=7)
lt['A1'] = 'Lịch trình theo WBS'; lt['A1'].font = X.F_TEN
X.tieu_de(lt, 4, cols + [f'T{i + 1}' for i in range(len(tuan))], widths + [4.2] * len(tuan))
C0 = len(cols) + 1
for j, w in enumerate(tuan):
    c = lt.cell(row=3, column=C0 + j, value=w); c.number_format = 'dd/mm'; c.font = Font(name=X.FONT, size=8, color='0000FF')
    c.alignment = Alignment(text_rotation=90)
lt.row_dimensions[3].height = 36
tv = {t['id']: t['ho_ten'] for t in loi.doc_json('dau_vao/thanh_vien.json')['thanh_vien']}
r = 5
for g in lich['goi']:
    X.o(lt, r, 1, str(g['ma']), dam=not g['la'])
    X.o(lt, r, 2, g['ten'], dam=not g['la'])
    X.o(lt, r, 3, g['muc'])
    if g['la']:
        X.o(lt, r, 4, float(g['md']), fmt=X.SO2, nhap=True)
        X.o(lt, r, 5, ', '.join(tv.get(p, p) for p in g['nguoi']) + (' (dàn trải)' if g.get('dan_trai') else ''), nhap=True)
        X.o(lt, r, 6, D(g['bat_dau']), fmt=X.NGAY, nhap=True)
        X.o(lt, r, 7, D(g['ket_thuc']), fmt=X.NGAY, nhap=True)
        X.o(lt, r, 11, ', '.join(g.get('phu_thuoc', [])), nhap=True)
        X.o(lt, r, 12, 'Có' if g.get('gang') else '', nhap=True)
    else:
        X.o(lt, r, 4, f'=SUMIFS(D:D,M:M,A{r})', fmt=X.SO2, dam=True)
        X.o(lt, r, 5, '')
        X.o(lt, r, 6, f'=_xlfn.MINIFS(F:F,M:M,A{r})', fmt=X.NGAY, dam=True)
        X.o(lt, r, 7, f'=_xlfn.MAXIFS(G:G,M:M,A{r})', fmt=X.NGAY, dam=True)
        X.o(lt, r, 11, ''); X.o(lt, r, 12, '')
    X.o(lt, r, 8, f'=NETWORKDAYS(F{r},G{r},Tong_quan!$F$4:$F$40)', fmt='0')
    gd_, sp_ = cot_giai_doan(g)
    X.o(lt, r, 9, gd_, nhap=True); X.o(lt, r, 10, sp_, nhap=True)
    X.o(lt, r, 13, str(g.get('cha', '')))
    X.o(lt, r, 14, van_ban_rang_buoc(g.get('rang_buoc', [])), nhap=True)
    r += 1
last = r - 1
rng = f'{L(C0)}5:{L(C0 + len(tuan) - 1)}{last}'
cl = L(C0)
lt.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($C5<3,{cl}$3<=$G5,{cl}$3+4>=$F5)'], fill=PatternFill('solid', fgColor='808080'), stopIfTrue=True))
lt.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($L5="Có",{cl}$3<=$G5,{cl}$3+4>=$F5)'], fill=PatternFill('solid', fgColor='C00000'), stopIfTrue=True))
lt.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({cl}$3<=$G5,{cl}$3+4>=$F5)'], fill=PatternFill('solid', fgColor='4472C4')))
lt.freeze_panes = 'C5'
X.ghi_chu_bang(lt, last + 2, 'Thanh xám: mục cha. Thanh đỏ: gói trên đường găng theo phụ thuộc. Thanh xanh: gói khác. Mỗi cột là một tuần, ngày đầu tuần ở dòng 3. '
                             'Gói quản lý dự án ghi giai đoạn (Khởi động, Suốt dự án, Kết thúc) thay cho sprint, gói Kết thúc nằm sau cut over.')

# ---------------- Sprint
sp = wb.create_sheet('Sprint')
X.tieu_de(sp, 1, ['Sprint', 'Từ ngày', 'Đến ngày', 'Số ngày làm việc', 'MD kế hoạch', 'Quỹ thời gian, MD', 'Tỷ lệ sử dụng', 'Mục tiêu', 'Sản phẩm trình diễn'],
          [8, 12, 12, 10, 12, 14, 11, 44, 44])
for i, s in enumerate(lich['sprint'], 2):
    X.o(sp, i, 1, s['ma'], nhap=True); X.o(sp, i, 2, D(s['tu']), fmt=X.NGAY, nhap=True); X.o(sp, i, 3, D(s['den']), fmt=X.NGAY, nhap=True)
    X.o(sp, i, 4, f'=NETWORKDAYS(B{i},C{i},Tong_quan!$F$4:$F$40)', fmt='0')
    X.o(sp, i, 5, f'=SUMIFS(Phan_bo!D:D,Phan_bo!C:C,A{i})', fmt=X.SO2)
    X.o(sp, i, 6, f'=SUMIFS(Phan_bo!E:E,Phan_bo!C:C,A{i})', fmt=X.SO2)
    X.o(sp, i, 7, f'=IF(F{i}=0,0,E{i}/F{i})', fmt=X.PT)
    X.o(sp, i, 8, s['muc_tieu'], nhap=True, wrap=True); X.o(sp, i, 9, s['san_pham_demo'], nhap=True, wrap=True)
n_sp = len(lich['sprint']) + 1
_q_kt = 0
if any(x['sprint'] == 'KT' for x in lich.get('phan_bo', [])):  # giai đoạn Kết thúc sau cut over: MD có, quỹ huy động bằng 0
    n_sp += 1
    _q_kt = sum(float(x.get('quy') or 0) for x in lich['phan_bo'] if x['sprint'] == 'KT')  # quỹ theo huy_dong_ket_thuc của thanh_vien.json (từ 1.4.6)
    X.o(sp, n_sp, 1, 'KT', nhap=True); X.o(sp, n_sp, 8, 'Giai đoạn Kết thúc sau cut over' + (', quỹ huy động theo tỷ lệ huy động giai đoạn Kết thúc của '
                                                       'thành viên' if _q_kt else ', không có quỹ huy động'), nhap=True, wrap=True)
    X.o(sp, n_sp, 5, f'=SUMIFS(Phan_bo!D:D,Phan_bo!C:C,A{n_sp})', fmt=X.SO2); X.o(sp, n_sp, 6, f'=SUMIFS(Phan_bo!E:E,Phan_bo!C:C,A{n_sp})', fmt=X.SO2)
X.o(sp, n_sp + 1, 1, 'Tổng', dam=True)
for c in (4, 5, 6):
    X.o(sp, n_sp + 1, c, f'=SUM({L(c)}2:{L(c)}{n_sp})', fmt=X.SO2 if c > 4 else '0', dam=True)
X.o(sp, n_sp + 1, 7, f'=IF(F{n_sp + 1}=0,0,E{n_sp + 1}/F{n_sp + 1})', fmt=X.PT, dam=True)

# ---------------- Moc
mc = wb.create_sheet('Moc')
X.tieu_de(mc, 1, ['Mốc', 'Nội dung', 'Ngày dự kiến', 'Hạn', 'Trạng thái', 'Căn cứ (gói hoàn thành)', 'Đợt thanh toán', 'Sprint'], [6, 44, 12, 12, 13, 26, 10, 8])
for i, m in enumerate(lich['moc'], 2):
    X.o(mc, i, 1, m['ma'], nhap=True); X.o(mc, i, 2, m['ten'], nhap=True, wrap=True); X.o(mc, i, 3, D(m['ngay']), fmt=X.NGAY, nhap=True)
    X.o(mc, i, 4, D(m['han']) if m['han'] else None, fmt=X.NGAY, nhap=True)
    X.o(mc, i, 5, f'=IF(D{i}="","Không có hạn",IF(C{i}<=D{i},"Đúng hạn","Trễ hạn"))')
    X.o(mc, i, 6, m['can_cu'], nhap=True, wrap=True); X.o(mc, i, 7, m['dot_thanh_toan'] or None, nhap=True); X.o(mc, i, 8, m['sprint'], nhap=True)
mc.conditional_formatting.add(f'E2:E{len(lich["moc"]) + 1}', FormulaRule(formula=['E2="Trễ hạn"'], fill=PatternFill('solid', fgColor='F8CBAD')))

# ---------------- Giai_doan: Khởi động, Suốt dự án (gói quản lý dàn trải), Thực hiện (các sprint), Kết thúc (sau cut over)
gd_ds = lich.get('giai_doan', [])
if gd_ds:
    gs = wb.create_sheet('Giai_doan')
    gs['A1'] = 'Giai đoạn dự án'; gs['A1'].font = X.F_TEN
    X.tieu_de(gs, 3, ['Giai đoạn', 'Từ ngày', 'Đến ngày', 'Số ngày làm việc', 'Số gói', 'MD', 'Căn cứ', 'Gói lá'], [14, 12, 12, 10, 8, 9, 30, 60])
    for i, x in enumerate(gd_ds, 4):
        X.o(gs, i, 1, x['ten'], nhap=True); X.o(gs, i, 2, D(x['tu']), fmt=X.NGAY, nhap=True); X.o(gs, i, 3, D(x['den']), fmt=X.NGAY, nhap=True)
        X.o(gs, i, 4, f'=NETWORKDAYS(B{i},C{i},Tong_quan!$F$4:$F$40)', fmt='0')
        X.o(gs, i, 5, f'=COUNTIFS(Lich_trinh!$I:$I,A{i},Lich_trinh!$E:$E,"?*")' if x['ma'] != 'thuc_hien' else x['so_goi'], fmt='0', nhap=x['ma'] == 'thuc_hien')
        X.o(gs, i, 6, x['md'], fmt=X.SO2, nhap=True)
        can_cu = {'khoi_dong': 'Quyết định của nhóm về giai đoạn Khởi động', 'thuc_hien': ', '.join(x.get('sprint', [])),
                  'xuyen_suot': 'Gói quản lý dàn trải theo sprint', 'ket_thuc': f'{x.get("so_ngay", "")} ngày làm việc sau cut over'}[x['ma']]
        X.o(gs, i, 7, can_cu + (' (mặc định)' if x.get('mac_dinh') else ''), nhap=True, wrap=True)
        X.o(gs, i, 8, ', '.join(x['goi']) if x['ma'] != 'thuc_hien' else f'{x["so_goi"]} gói theo sprint, xem Lich_trinh', nhap=True, wrap=True)
    n_gd = 3 + len(gd_ds)
    X.o(gs, n_gd + 1, 1, 'Tổng', dam=True); X.o(gs, n_gd + 1, 6, f'=SUM(F4:F{n_gd})', fmt=X.SO2, dam=True)
    X.ghi_chu_bang(gs, n_gd + 3, f'Ngày đóng dự án (hết giai đoạn Kết thúc): {loi.ngay(lich["ngay_dong_du_an"])}. Gói quản lý dự án không gắn sprint '
                                 'phát triển, không nằm trong mục tiêu, sản phẩm trình diễn của sprint. Gói Kết thúc không tính vào dự trữ đến cut over.')

# ---------------- Phan_bo và Tai_nguon
pb = wb.create_sheet('Phan_bo')
X.tieu_de(pb, 1, ['Mã thành viên', 'Họ tên', 'Sprint', 'MD kế hoạch', 'Quỹ thời gian, MD', 'Tỷ lệ'], [12, 26, 8, 12, 14, 9])
for i, x in enumerate(lich['phan_bo'], 2):
    X.o(pb, i, 1, x['thanh_vien'], nhap=True); X.o(pb, i, 2, x['ho_ten'], nhap=True); X.o(pb, i, 3, x['sprint'], nhap=True)
    X.o(pb, i, 4, x['md'], fmt=X.SO2, nhap=True); X.o(pb, i, 5, x['quy'], fmt=X.SO2, nhap=True)
    X.o(pb, i, 6, f'=IF(E{i}=0,0,D{i}/E{i})', fmt=X.PT)
pb.conditional_formatting.add(f'F2:F{len(lich["phan_bo"]) + 1}', FormulaRule(formula=['F2>1'], fill=PatternFill('solid', fgColor='F8CBAD')))
tn = wb.create_sheet('Tai_nguon')
tn['A1'] = 'Tải công việc theo thành viên và sprint, MD'; tn['A1'].font = X.F_TEN
ma_sp = [s['ma'] for s in lich['sprint']] + sorted({x['sprint'] for x in lich['phan_bo']} - {s['ma'] for s in lich['sprint']})  # kỳ KT sau cut over
X.tieu_de(tn, 3, ['Mã', 'Họ tên'] + ma_sp + ['Tổng MD', 'Tổng quỹ', 'Tỷ lệ'], [10, 26] + [8] * len(ma_sp) + [10, 10, 9])
ids = []
for x in lich['phan_bo']:
    if x['thanh_vien'] not in ids:
        ids.append(x['thanh_vien'])
for i, p in enumerate(ids, 4):
    X.o(tn, i, 1, p, nhap=True); X.o(tn, i, 2, tv.get(p, p), nhap=True)
    for j, s in enumerate(ma_sp, 3):
        X.o(tn, i, j, f'=SUMIFS(Phan_bo!$D:$D,Phan_bo!$A:$A,$A{i},Phan_bo!$C:$C,{L(j)}$3)', fmt=X.SO2)
    k = 3 + len(ma_sp)
    X.o(tn, i, k, f'=SUM({L(3)}{i}:{L(k - 1)}{i})', fmt=X.SO2)
    X.o(tn, i, k + 1, f'=SUMIFS(Phan_bo!$E:$E,Phan_bo!$A:$A,$A{i})', fmt=X.SO2)
    X.o(tn, i, k + 2, f'=IF({L(k + 1)}{i}=0,0,{L(k)}{i}/{L(k + 1)}{i})', fmt=X.PT)
# ô đỏ: MD của thành viên trong sprint vượt quỹ thời gian của sprint đó (đối chiếu sheet Phan_bo)
# ô vàng: thành viên được huy động trong sprint nhưng dùng dưới ngưỡng quỹ của sprint (cảnh báo L07 của bước lập lịch)
NGUONG_NGUOI = float(lich.get('nguong_su_dung_nguoi_sprint', 0.3))
_quy_o = f'SUMIFS(Phan_bo!$E:$E,Phan_bo!$A:$A,$A4,Phan_bo!$C:$C,{L(3)}$3)'
_vung = f'{L(3)}4:{L(2 + len(ma_sp))}{3 + len(ids)}'
# cột KT (giai đoạn Kết thúc sau cut over) không có quỹ huy động, MD trả theo công sức (L28 đã được nhóm chấp nhận): không tô đỏ
tn.conditional_formatting.add(_vung, FormulaRule(formula=[f'AND({L(3)}$3<>"KT",{L(3)}4>{_quy_o}+0.01)'], fill=PatternFill('solid', fgColor='F8CBAD'), stopIfTrue=True))
tn.conditional_formatting.add(_vung, FormulaRule(formula=[f'AND({_quy_o}>0,{L(3)}4<{NGUONG_NGUOI}*{_quy_o})'], fill=PatternFill('solid', fgColor='FFE699')))
X.ghi_chu_bang(tn, 5 + len(ids), f'Ô đỏ: vượt quỹ thời gian của sprint. Ô vàng: được huy động trong sprint nhưng dùng dưới {loi.pt(NGUONG_NGUOI)} quỹ của sprint.'
                             + ((' Cột KT: MD giai đoạn Kết thúc sau cut over, quỹ theo tỷ lệ huy động giai đoạn Kết thúc của thành viên.' if _q_kt else
                                 ' Cột KT: MD giai đoạn Kết thúc sau cut over, trả lương theo công sức, không có quỹ huy động (nhóm đã chấp nhận).')
                                if 'KT' in ma_sp else ''))

# ---------------- Nhu_cau: MD theo loại công việc và sprint, dùng để đề xuất huy động theo sprint
if lich.get('nhu_cau_theo_sprint'):
    nc = wb.create_sheet('Nhu_cau')
    nc['A1'] = 'Nhu cầu nhân lực theo loại công việc và sprint, MD'; nc['A1'].font = X.F_TEN
    loai_ds = list(lich['nhu_cau_theo_sprint'][0]['theo_loai'])
    k = 2 + len(loai_ds)
    X.tieu_de(nc, 3, ['Sprint'] + loai_ds + ['Tổng MD', 'Quỹ huy động, MD', 'Tỷ lệ sử dụng'], [8] + [9] * len(loai_ds) + [10, 14, 11])
    for i, s in enumerate(lich['nhu_cau_theo_sprint'], 4):
        X.o(nc, i, 1, s['sprint'], nhap=True)
        for j, lo in enumerate(loai_ds, 2):
            X.o(nc, i, j, s['theo_loai'][lo], fmt=X.SO2, nhap=True)
        X.o(nc, i, k, f'=SUM({L(2)}{i}:{L(k - 1)}{i})', fmt=X.SO2)
        X.o(nc, i, k + 1, f'=SUMIFS(Phan_bo!$E:$E,Phan_bo!$C:$C,$A{i})', fmt=X.SO2)
        X.o(nc, i, k + 2, f'=IF({L(k + 1)}{i}=0,0,{L(k)}{i}/{L(k + 1)}{i})', fmt=X.PT)
    n_nc = 3 + len(lich['nhu_cau_theo_sprint'])
    X.o(nc, n_nc + 1, 1, 'Tổng', dam=True)
    for j in range(2, k + 2):
        X.o(nc, n_nc + 1, j, f'=SUM({L(j)}4:{L(j)}{n_nc})', fmt=X.SO2, dam=True)
    X.o(nc, n_nc + 1, k + 2, f'=IF({L(k + 1)}{n_nc + 1}=0,0,{L(k)}{n_nc + 1}/{L(k + 1)}{n_nc + 1})', fmt=X.PT, dam=True)
    X.ghi_chu_bang(nc, n_nc + 3, 'MD từng loại công việc theo ngày thực hiện trong lịch đã xếp nguồn lực. So sánh với quỹ huy động để đề xuất tu_sprint, den_sprint của từng vị trí.')

# ---------------- Cam_ket: hạn đầu vào của Chủ đầu tư và gói có Chủ đầu tư tham gia
if lich.get('cam_ket_chu_dau_tu') or lich.get('chu_dau_tu_tham_gia'):
    ck = wb.create_sheet('Cam_ket')
    ck['A1'] = 'Cam kết đầu vào của Chủ đầu tư'; ck['A1'].font = X.F_TEN
    X.tieu_de(ck, 3, ['Mã', 'Nội dung cam kết', 'Hạn (cuối khoảng nếu là khoảng tham gia)', 'Gói phụ thuộc', 'Bắt đầu sớm nhất', 'Trạng thái',
                      'Đầu khoảng tham gia', 'Kết thúc muộn nhất'], [8, 50, 14, 22, 14, 20, 14, 14])
    for i, c in enumerate(lich.get('cam_ket_chu_dau_tu', []), 4):
        X.o(ck, i, 1, c['ma'], nhap=True); X.o(ck, i, 2, c['noi_dung'], nhap=True, wrap=True)
        X.o(ck, i, 3, D(c['han']) if c.get('han') else None, fmt=X.NGAY, nhap=True); X.o(ck, i, 4, ', '.join(c['goi']), nhap=True, wrap=True)
        X.o(ck, i, 5, D(c['bat_dau_som_nhat']) if c.get('bat_dau_som_nhat') else None, fmt=X.NGAY, nhap=True)
        if c.get('khoang_tu'):  # cam kết tham gia trong khoảng kiểm thử chấp nhận: so đầu, cuối khoảng
            X.o(ck, i, 7, D(c['khoang_tu']), fmt=X.NGAY, nhap=True); X.o(ck, i, 8, D(c['ket_thuc_muon_nhat']), fmt=X.NGAY, nhap=True)
            X.o(ck, i, 6, f'=IF(AND(E{i}>=G{i},H{i}<=C{i}),"Trong khoảng cam kết","Ngoài khoảng cam kết")')
        else:
            X.o(ck, i, 6, f'=IF(C{i}="","Không có hạn",IF(D{i}="","Chưa gắn gói",IF(E{i}>=C{i},"Đạt","Bắt đầu trước hạn")))')
    r2 = 6 + len(lich.get('cam_ket_chu_dau_tu', []))
    ck.cell(row=r2 - 1, column=1, value='Gói có Chủ đầu tư tham gia').font = X.F_TEN
    X.tieu_de(ck, r2, ['Gói', 'Loại tham gia', 'Thời lượng tối thiểu, ngày', 'Bắt đầu', 'Kết thúc', 'Số ngày làm việc', 'Trạng thái'], [8, 50, 12, 22, 14, 18, 16])
    for i, x in enumerate(lich.get('chu_dau_tu_tham_gia', []), r2 + 1):
        X.o(ck, i, 1, x['goi'], nhap=True)
        X.o(ck, i, 2, 'Kiểm thử chấp nhận' if x['loai'] == 'uat' else 'Phê duyệt tài liệu', nhap=True)
        X.o(ck, i, 3, x['so_ngay'], fmt='0', nhap=True)
        X.o(ck, i, 4, D(x['bat_dau']), fmt=X.NGAY, nhap=True); X.o(ck, i, 5, D(x['ket_thuc']), fmt=X.NGAY, nhap=True)
        X.o(ck, i, 6, f'=NETWORKDAYS(D{i},E{i},Tong_quan!$F$4:$F$40)', fmt='0')
        X.o(ck, i, 7, f'=IF(F{i}>=C{i},"Đủ thời lượng","Thiếu thời lượng")')

# ---------------- Huong_dan
hd = wb.create_sheet('Huong_dan')
for i, t in enumerate(['Hướng dẫn sử dụng master schedule',
                       'Sheet Lich_trinh: lịch từng gói WBS, thanh Gantt theo tuần. Mục cha tính bằng công thức từ mục con.',
                       'Sheet Sprint: mục tiêu, sản phẩm trình diễn, MD kế hoạch và quỹ thời gian từng sprint.',
                       'Sheet Moc: các mốc, ngày dự kiến, hạn, trạng thái, liên kết đợt thanh toán.',
                       'Sheet Giai_doan: giai đoạn Khởi động, Thực hiện (các sprint), Kết thúc (sau cut over) và gói quản lý dự án dàn trải suốt dự án.',
                       'Sheet Phan_bo và Tai_nguon: MD kế hoạch theo thành viên và sprint, ô đỏ là quá tải, ô vàng là dùng dưới ngưỡng quỹ của sprint huy động.',
                       'Sheet Nhu_cau: MD theo loại công việc và sprint so với quỹ huy động, dùng để đề xuất thời gian huy động từng vị trí.',
                       'Sheet Cam_ket: hạn đầu vào của Chủ đầu tư và gói phụ thuộc, gói có Chủ đầu tư tham gia (phê duyệt tài liệu, kiểm thử chấp nhận) và thời lượng tối thiểu.',
                       'Cách cập nhật: sửa WBS, giao việc hoặc mốc và lập lại lịch theo skill master schedule. Không sửa tay các ô chữ xanh dương.',
                       f'Nguồn dữ liệu: lịch dự án do công cụ của nhóm lập, tạo lúc {loi.ngay(loi.bay_gio()[:10])} {loi.bay_gio()[11:16]}.'], 1):
    c = hd.cell(row=i, column=1, value=t); c.font = X.F_TEN if i == 1 else X.F_THUONG
hd.column_dimensions['A'].width = 120
wb.save(OUT)

# ---------------- tính lại và đối chiếu
kq = X.tinh_lai(OUT)
ghi_chu = 'Chưa tính lại công thức (máy không có LibreOffice). Excel tự tính khi mở.'
if kq and kq.get('status'):
    loi_ct = X.dem_loi_cong_thuc(OUT)
    if loi_ct:
        loi.thoat_loi('Công thức lỗi: ' + ', '.join(loi_ct[:10]))
    v = X.doc_gia_tri(OUT)
    tong = v['Tong_quan']['B7'].value
    if tong is None or abs(float(tong) - float(wbs['tong_md'])) > 0.01:
        loi.thoat_loi(f'Tổng MD trong Excel ({tong}) khác WBS ({wbs["tong_md"]})')
    so_ngay = v['Tong_quan']['B6'].value
    if int(so_ngay) != lich['so_ngay_lam_viec']:
        loi.thoat_loi(f'Số ngày làm việc trong Excel ({so_ngay}) khác lịch ({lich["so_ngay_lam_viec"]})')
    for i, s in enumerate(lich['sprint'], 2):
        if abs(float(v['Sprint'][f'E{i}'].value) - s['md_ke_hoach']) > 0.05:
            loi.thoat_loi(f'MD sprint {s["ma"]} trong Excel khác lich.json')
    for i, x in enumerate(lich.get('giai_doan', []), 4):
        if x['ma'] != 'xuyen_suot' and int(v['Giai_doan'][f'D{i}'].value) != int(x['so_ngay']):
            loi.thoat_loi(f'Số ngày làm việc giai đoạn {x["ten"]} trong Excel ({v["Giai_doan"][f"D{i}"].value}) khác lịch ({x["so_ngay"]})')
        if x['ma'] != 'thuc_hien' and int(v['Giai_doan'][f'E{i}'].value) != x['so_goi']:
            loi.thoat_loi(f'Số gói giai đoạn {x["ten"]} trong Excel khác lịch ({x["so_goi"]})')
    if lich.get('nhu_cau_theo_sprint'):
        o_tong = v['Nhu_cau'].cell(row=4 + len(lich['nhu_cau_theo_sprint']), column=2 + len(lich['nhu_cau_theo_sprint'][0]['theo_loai'])).value
        if o_tong is None or abs(float(o_tong) - float(wbs['tong_md'])) > 0.05:
            loi.thoat_loi(f'Tổng MD sheet Nhu_cau ({o_tong}) khác WBS ({wbs["tong_md"]})')
    ghi_chu = f'Đã tính lại công thức, 0 lỗi, đối chiếu tổng MD {wbs["tong_md"]:g}, {so_ngay} ngày làm việc, MD từng sprint khớp lich.json.'
print('Đã tạo ' + loi.tuong_doi(OUT) + '. ' + ghi_chu)
pbs = loi.hoan_tat_buoc('lich', 'Lập lịch có ràng buộc nguồn lực và xuất master schedule Excel. ' + ghi_chu,
                        hash_tep={'ket_qua/04_Master_schedule/Master_schedule.xlsx': loi.bam_obj(lich)})
for k, v in pbs.items():
    print(f'  {k}: phiên bản {v}')
