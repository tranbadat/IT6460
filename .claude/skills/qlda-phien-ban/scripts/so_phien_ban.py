"""Sổ phiên bản và đường cơ sở của bộ tài liệu dự án.

Dùng:
  python so_phien_ban.py tong_hop
      Tổng hợp sổ phiên bản (du_lieu/phien_ban.json), nhật ký chạy, sổ thay đổi, lịch sử git (nếu thư mục là kho git)
      thành du_lieu/phien_ban_tong_hop.json và ket_qua/15_Quan_ly_phien_ban/So_phien_ban.xlsx.
  python so_phien_ban.py trang_thai <tệp> --trang-thai "Đã kiểm tra|Đã duyệt" --nguoi "<họ tên>" [--ghi-chu ...]
      Cập nhật trạng thái phiên bản hiện hành của một tài liệu (sau kiểm tra lớp B, sau phê duyệt).
  python so_phien_ban.py duong_co_so <mã, ví dụ BL1> --ten "<tên đường cơ sở>" [--tep <tệp> ...] [--ngay yyyy-mm-dd] [--ghi-chu ...] [--hop-dong] [--cr CR03] [--bo-qua-kiem-tra] [--lap-lai "<lý do>"]
      --lap-lai: lập lại một đường cơ sở sau thay đổi đã có (cùng mã, cùng ngày) sau khi sửa tài liệu theo kiểm tra lớp B. Bản cũ chuyển vào
      duong_co_so_thay_the, phiên bản cũ giữ trong sổ; phiên bản mới giữ số chính của bản cũ, số phụ kế tiếp (2.2 thành 2.3), lý do ghi vào lam_lai (từ 1.2.5).
      --hop-dong: đường cơ sở khi ký hợp đồng (BL1), HSMT và HSDT chuyển trạng thái Khóa. Chỉ một lần, không ghi đè.
      --cr: đường cơ sở sau yêu cầu thay đổi đã phê duyệt (BL2), mã CR phải có trong du_lieu/thay_doi.json và đã phê duyệt.
      Lập đường cơ sở: mỗi tài liệu được chọn (mặc định mọi tài liệu trong ket_qua) lên số chính kế tiếp (1.0, 2.0),
      trạng thái Đường cơ sở, lưu băm đầu vào tại thời điểm lập (vao_duong_co_so) và bản phát hành hiện hành tại ngày lập.
      Không lập được nếu tài liệu còn dùng đầu vào cũ.
  Phiên bản đường cơ sở, khóa không đổi trạng thái bằng lệnh trang_thai và không bị ghi đè khi đăng ký lại cùng nội dung.
  Sau ký hợp đồng, mỗi phiên bản mới ghi khóa phat_hanh (bản phát hành hiện hành theo ngày tài liệu, du_lieu/phat_hanh.json).
Quy ước: references/quy_uoc.json.
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import argparse
import datetime as dt
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'qlda-loi', 'scripts'))
import loi  # noqa: E402

ap = argparse.ArgumentParser()
ap.add_argument('lenh', nargs='?', default='tong_hop'); ap.add_argument('doi_tuong', nargs='?')
ap.add_argument('--trang-thai'); ap.add_argument('--nguoi', default=''); ap.add_argument('--ghi-chu', default='')
ap.add_argument('--ten', default=''); ap.add_argument('--tep', nargs='*'); ap.add_argument('--ngay'); ap.add_argument('--bo-qua-kiem-tra', action='store_true')
ap.add_argument('--hop-dong', action='store_true', help='đường cơ sở khi ký hợp đồng: khóa HSMT, HSDT')
ap.add_argument('--cr', default='', help='mã yêu cầu thay đổi đã phê duyệt làm căn cứ đường cơ sở (ví dụ BL2)')
ap.add_argument('--lap-lai', default='', help='lý do lập lại đường cơ sở sau thay đổi đã có cùng mã')
a = ap.parse_args()
sd = loi.so_do(); BD = loi.nut_bien_dong(sd)
reg = loi.doc_json('du_lieu/phien_ban.json', {'tai_lieu': {}})
reg.setdefault('duong_co_so', [])
QU = loi.doc_json(os.path.join(HERE, '..', 'references', 'quy_uoc.json'))
TRANG_THAI_CO_DINH = ('Đường cơ sở', 'Khóa')  # phiên bản không đổi trạng thái bằng lệnh trang_thai
buoc = {b['id']: b for b in sd['buoc']}


def cu(rec):
    vao_hien = buoc.get(rec.get('buoc'), {}).get('vao', list(rec.get('vao', {})))
    return sorted(n for n, h in rec.get('vao', {}).items()
                  if n in vao_hien and n not in BD and loi.duong_dan_nut(sd, n) and loi.bam_nut(sd, n) != h)


def gan_ma_cr(reg_, so_td_):
    """Gắn mã CR cho phiên bản tạo do yêu cầu thay đổi (quy ước 2.1): phiên bản đăng ký sau quyết định phê duyệt (thu_tu lớn hơn
    dem_phien_ban_khi_quyet_dinh) và trước khi CR đóng, thuộc bước trong tai_lieu_cap_nhat, gồm cả bản sửa sau đường cơ sở sau thay đổi
    (vẫn thuộc đường cơ sở đó). Trả số phiên bản được gắn."""
    n = 0
    for c in so_td_.get('yeu_cau', []):
        dem = c.get('dem_phien_ban_khi_quyet_dinh')
        if dem is None or not (c.get('ap_dung') or {}).get('da_ghi'):
            continue
        dong = next((x.get('ghi_luc', x['ngay']) for x in reversed(c.get('lich_su', [])) if x.get('trang_thai') == 'Đã đóng'), '')
        buoc_cr = {x.split(' ')[0] for x in c.get('tai_lieu_cap_nhat', [])}
        for lst in reg_['tai_lieu'].values():
            for r in lst:
                if r.get('ma_cr') or r.get('duong_co_so') or r.get('buoc') not in buoc_cr or int(r.get('thu_tu') or 0) <= dem:
                    continue
                if dong and r.get('ngay', '') > dong:
                    continue
                r['ma_cr'] = c['ma']
                r['ghi_chu'] = r.get('ghi_chu') if c['ma'] in str(r.get('ghi_chu', '')) else (f'{c["ma"]}. {r["ghi_chu"]}' if r.get('ghi_chu') else c['ma'])
                n += 1
    return n


def phien_ban_sau_cr(reg_, ma_cr):
    """Phiên bản sau thay đổi của CR: tài liệu của đường cơ sở lập theo CR (không gồm tài liệu khóa), thiếu thì phiên bản mới nhất có mã CR."""
    bl = next((x for x in reg_.get('duong_co_so', []) if x.get('ma_cr') == ma_cr), None)
    if bl:
        return {p: v for p, v in bl['tai_lieu'].items() if any(r.get('duong_co_so') == bl['ma'] for r in reg_['tai_lieu'].get(p, []))}
    return {p: next(r['phien_ban'] for r in reversed(lst) if r.get('ma_cr') == ma_cr) for p, lst in reg_['tai_lieu'].items()
            if any(r.get('ma_cr') == ma_cr for r in lst)}


def tai_lieu_nop():
    return [p for p in reg['tai_lieu'] if p.startswith('ket_qua/') and os.path.exists(loi.duong_dan(p))]


if a.lenh == 'trang_thai':
    p = loi.tuong_doi(loi.duong_dan(a.doi_tuong or ''))
    if p not in reg['tai_lieu']:
        loi.thoat_loi(f'{p} chưa có trong sổ phiên bản')
    hop_le = [x['ma'] for x in QU['trang_thai'] if x['ma'] not in ('Đường cơ sở', 'Khóa')]
    if a.trang_thai not in hop_le:
        loi.thoat_loi('Trạng thái phải là một trong: ' + ', '.join(hop_le) + '. Đường cơ sở dùng lệnh duong_co_so.')
    rec = reg['tai_lieu'][p][-1]
    if rec.get('trang_thai') in TRANG_THAI_CO_DINH:
        loi.thoat_loi(f'{p} phiên bản {rec["phien_ban"]} đang ở trạng thái {rec["trang_thai"]}, không đổi trạng thái. '
                      'Thay đổi sau đường cơ sở đi qua yêu cầu thay đổi, tạo phiên bản mới rồi mới cập nhật trạng thái.')
    if cu(rec):
        loi.thoat_loi(f'{p} đang dùng đầu vào cũ ({", ".join(cu(rec))}), tạo lại trước khi cập nhật trạng thái.')
    rec.setdefault('lich_su_trang_thai', []).append({'tu': rec['trang_thai'], 'den': a.trang_thai, 'nguoi': a.nguoi, 'luc': loi.bay_gio(), 'ghi_chu': a.ghi_chu})
    rec['trang_thai'] = a.trang_thai
    loi.ghi_json('du_lieu/phien_ban.json', reg)
    loi.ghi_nhat_ky('qlda-phien-ban', 'phien_ban.trang_thai', 'script', f'{p} {rec["phien_ban"]}: {a.trang_thai} ({a.nguoi})', ra=['du_lieu/phien_ban.json'])
    print(f'{p} phiên bản {rec["phien_ban"]}: {a.trang_thai}')
    sys.exit(0)

if a.lenh == 'duong_co_so':
    ma = a.doi_tuong
    CU_BL = next((x for x in reg['duong_co_so'] if x['ma'] == ma), None)
    if a.lap_lai and (not CU_BL or CU_BL.get('hop_dong') or a.hop_dong):
        loi.thoat_loi('--lap-lai chỉ dùng cho đường cơ sở sau thay đổi đã có (không dùng cho đường cơ sở hợp đồng)')
    if not ma or (CU_BL and not a.lap_lai):
        loi.thoat_loi('Cần mã đường cơ sở mới, chưa dùng (ví dụ BL1, BL2), hoặc --lap-lai "<lý do>" để lập lại đường cơ sở sau thay đổi')
    if a.hop_dong and reg.get('hop_dong'):
        loi.thoat_loi(f'Đã có đường cơ sở hợp đồng {reg["hop_dong"]["duong_co_so"]} ngày {loi.ngay(reg["hop_dong"]["ngay"])}, không ghi đè. '
                      'Đường cơ sở sau thay đổi dùng --cr <mã CR> thay cho --hop-dong.')
    if a.cr:
        cr_ = next((c for c in loi.doc_json('du_lieu/thay_doi.json', {'yeu_cau': []}).get('yeu_cau', []) if c.get('ma') == a.cr), None)
        if not cr_:
            loi.thoat_loi(f'Không có yêu cầu thay đổi {a.cr} trong du_lieu/thay_doi.json')
        if not (cr_.get('trang_thai') in ('Đã phê duyệt', 'Đã thực hiện') or (cr_.get('trang_thai') == 'Đã đóng' and cr_.get('ket_qua_dong') not in ('Từ chối', 'Hoãn'))):
            loi.thoat_loi(f'{a.cr} đang ở trạng thái {cr_.get("trang_thai")}, chỉ lập đường cơ sở theo yêu cầu thay đổi đã phê duyệt')
    ngay_bl = a.ngay or loi.ngay_hien_hanh()
    ph_bl = loi.ban_phat_hanh_hien_hanh(ngay_bl) if (reg.get('hop_dong') or a.hop_dong) else None
    ds = [loi.tuong_doi(loi.duong_dan(t)) for t in a.tep] if a.tep else tai_lieu_nop()
    ds = [p for p in ds if not p.startswith(('ket_qua/15_', 'ket_qua/16_'))]
    # báo cáo theo kỳ (bước anh_chup, ví dụ báo cáo tiến độ) giữ số liệu tại thời điểm lập, không thuộc đường cơ sở kế hoạch
    ds = [p for p in ds if not (p in reg['tai_lieu'] and buoc.get(reg['tai_lieu'][p][-1]['buoc'], {}).get('anh_chup'))]
    loi_ = [f'{p}: chưa đăng ký' for p in ds if p not in reg['tai_lieu']]
    khoa = lambda p: reg.get('hop_dong') and buoc.get(reg['tai_lieu'][p][-1]['buoc'], {}).get('khoa')
    loi_ += [f'{p}: đầu vào đã đổi ({", ".join(cu(reg["tai_lieu"][p][-1]))})' for p in ds
             if p in reg['tai_lieu'] and cu(reg['tai_lieu'][p][-1]) and not khoa(p)]
    if loi_ and not a.bo_qua_kiem_tra:
        loi.thoat_loi('Chưa lập được đường cơ sở:\n  ' + '\n  '.join(loi_) + '\nChạy lại các bước liên quan (chay.py trang_thai) rồi lập lại.')
    chi_tiet, dung_lai = {}, []
    if CU_BL:  # phiên bản của bản cũ giữ trong sổ, đánh dấu đã thay để tra cứu đường cơ sở chỉ ra bản lập lại
        for lst_ in reg['tai_lieu'].values():
            for r_ in lst_:
                if r_.get('duong_co_so') == ma:
                    r_['duong_co_so_da_thay'] = r_.pop('duong_co_so')
    for p in ds:
        last = reg['tai_lieu'][p][-1]
        if khoa(p) and not a.hop_dong:  # hồ sơ đấu thầu đã khóa từ BL1: giữ phiên bản Khóa, thay đổi ghi bằng phụ lục
            chi_tiet[p] = last['phien_ban']; continue
        major = int(last['phien_ban'].split('.')[0]) + 1
        so_moi = f'{major}.0'
        if CU_BL and p in CU_BL.get('tai_lieu', {}):  # lập lại: giữ số chính của bản cũ, số phụ kế tiếp
            mj = int(CU_BL['tai_lieu'][p].split('.')[0])
            mn = max(int(r['phien_ban'].split('.')[1]) for r in reg['tai_lieu'][p] if int(r['phien_ban'].split('.')[0]) == mj)
            so_moi = f'{mj}.{mn + 1}'
        rec = dict(last, phien_ban=so_moi, ngay=loi.bay_gio(), trang_thai='Đường cơ sở', ghi_chu=f'Đường cơ sở {ma}. {a.ghi_chu}'.strip(),
                   nguoi=a.nguoi or last.get('nguoi', ''), duong_co_so=ma, vao_duong_co_so=dict(last.get('vao', {})))
        if ph_bl is not None:
            rec['phat_hanh'] = ph_bl
        if a.cr:
            rec['ma_cr'] = a.cr
        rec.pop('lich_su_trang_thai', None); rec.pop('ngay_kiem_tra', None)
        reg['tai_lieu'][p].append(rec); chi_tiet[p] = rec['phien_ban']
        if p.endswith('.docx') and rec.get('noi_dung') and os.path.exists(loi.duong_dan(rec['noi_dung'])):
            nd = loi.doc_json(rec['noi_dung']); nd['meta']['phien_ban'] = rec['phien_ban']
            nd_bl = loi.duong_dan(rec['noi_dung'].replace('.json', f'_{ma}.json')); loi.ghi_json(nd_bl, nd)
            r = subprocess.run([sys.executable, os.path.join(loi.THU_MUC_LOI, 'scripts', 'build_docx.py'), nd_bl, loi.duong_dan(p)]
                               + (['--co-dinh', rec['co_dinh']] if rec.get('co_dinh') else []), capture_output=True, text=True, encoding='utf-8', errors='replace')
            if r.returncode:
                print(f'Không dựng lại được {p}: {r.stderr.strip()[-300:]}')
            else:
                dung_lai.append(p)
    if CU_BL:
        reg['duong_co_so'] = [x for x in reg['duong_co_so'] if x is not CU_BL]
        reg.setdefault('duong_co_so_thay_the', []).append(dict(CU_BL, thay_luc=loi.bay_gio(), ly_do=a.lap_lai))
        reg.setdefault('lam_lai', []).append({'ngay': loi.bay_gio()[:10], 'ngay_du_an': ngay_bl,
                                              'noi_dung': f'Lập lại đường cơ sở {ma} cùng ngày {loi.ngay(ngay_bl)}: {a.lap_lai} Bản lập lúc {CU_BL.get("luc", "")} '
                                                          'chuyển vào danh sách đường cơ sở đã thay, các phiên bản cũ giữ trong sổ.'})
    reg['duong_co_so'].append({'ma': ma, 'ten': a.ten, 'ngay': ngay_bl, 'luc': loi.bay_gio(), 'nguoi': a.nguoi, 'ghi_chu': a.ghi_chu,
                               'tai_lieu': chi_tiet, 'bo_qua_kiem_tra': bool(loi_), 'hop_dong': a.hop_dong, 'ma_cr': a.cr,
                               'phat_hanh': ph_bl or ''})
    if a.hop_dong:
        reg['hop_dong'] = {'duong_co_so': ma, 'ngay': ngay_bl}
        for p in ds:
            if buoc.get(reg['tai_lieu'][p][-1]['buoc'], {}).get('khoa'):
                reg['tai_lieu'][p][-1]['trang_thai'] = 'Khóa'
    if a.cr:  # phiên bản tạo do CR trước đường cơ sở được gắn mã CR, CR ghi phiên bản sau thay đổi
        so_td_ = loi.doc_json('du_lieu/thay_doi.json', {'yeu_cau': []})
        gan_ma_cr(reg, so_td_)
        cr_ = next(c for c in so_td_['yeu_cau'] if c.get('ma') == a.cr)
        cr_['phien_ban_sau'] = phien_ban_sau_cr(reg, a.cr)
        cr_.setdefault('lich_su', []).append({'ngay': ngay_bl, 'ghi_luc': loi.bay_gio(), 'trang_thai': cr_.get('trang_thai', ''), 'nguoi': a.nguoi,
                                              'ghi_chu': f'Đường cơ sở {ma} sau thay đổi gồm {len(ds)} tài liệu, trong đó {len(cr_["phien_ban_sau"])} '
                                                         'tài liệu có phiên bản sau thay đổi, còn lại là tài liệu đấu thầu đã khóa'})  # từ 1.2.4
        loi.ghi_json('du_lieu/thay_doi.json', so_td_)
    loi.ghi_json('du_lieu/phien_ban.json', reg)
    loi.ghi_nhat_ky('qlda-phien-ban', f'phien_ban.duong_co_so.{ma}', 'script', f'Lập đường cơ sở {ma} ({a.ten}), {len(chi_tiet)} tài liệu',
                    ra=['du_lieu/phien_ban.json'], chi_tiet=chi_tiet)
    print(f'Đường cơ sở {ma}: {len(chi_tiet)} tài liệu, dựng lại {len(dung_lai)} tệp Word với số phiên bản mới.')
    if os.path.isdir(loi.duong_dan('.git')):
        print(f'Gợi ý: git add -A && git commit -m "Đường cơ sở {ma}: {a.ten}" && git tag {ma}')
    sys.exit(0)

if a.lenh != 'tong_hop':
    loi.thoat_loi('Lệnh không hợp lệ')

# Tổng hợp
ten_buoc = lambda bid: buoc.get(bid, {}).get('ten', bid)
so_td = loi.doc_json('du_lieu/thay_doi.json', {'yeu_cau': []})
if gan_ma_cr(reg, so_td):  # sổ lập trước khi có quy ước tự gắn mã CR: bổ sung một lần, giữ nguyên số phiên bản
    loi.ghi_json('du_lieu/phien_ban.json', reg)


def ten_nut(n):
    if n in sd['dau_vao']:
        return sd['dau_vao'][n]['mo_ta']
    return next((b['ten'] for b in sd['buoc'] if n in b['ra']), n)


WBS_NHOM = {str(g['ma']): g['ten'] for g in loi.doc_json('du_lieu/wbs.json', {'goi': []}).get('goi', []) if '.' not in str(g['ma'])}


def ten_tep_tl(k):
    """Tên tài liệu có dấu theo bước tạo ra tệp. Bước có nhiều tệp: WBS theo nhóm (tên nhóm trong wbs.json), báo cáo tiến độ theo kỳ,
    bước có cả Excel và Word ghi thêm loại tệp."""
    lst = reg['tai_lieu'].get(k, [])
    ten = ten_buoc(lst[-1]['buoc']) if lst else os.path.basename(k)
    goc = os.path.splitext(os.path.basename(k))[0]
    if k.startswith('ket_qua/03_'):
        m = re.match(r'WBS_(\d+)_', goc)
        return f'{ten}, nhóm {m.group(1)} {WBS_NHOM.get(m.group(1), "")}'.strip() if m else (f'{ten}, tổng hợp' if 'tong_hop' in goc else ten)
    if k.startswith('ket_qua/09_'):
        m = re.search(r'(S\d+|KT)$', goc)
        return f'{ten} {m.group(1)}' if m else ten
    cung_buoc = [q for q in reg['tai_lieu'] if q.startswith('ket_qua/') and reg['tai_lieu'][q][-1]['buoc'] == (lst[-1]['buoc'] if lst else '')]
    if len(cung_buoc) > 1:
        return ten + (' (bảng tính Excel)' if k.endswith('.xlsx') else ' (thuyết minh)' if k.endswith('.docx') else '')
    return ten


HSMT_NGAY = (loi.doc_json('du_lieu/hsmt.json', {}).get('lich') or {}).get('phat_hanh', '')


def ngay_tl(p, r):
    """Ngày tài liệu của một phiên bản: ngay_tai_lieu đã ghi, thiếu thì HSMT lấy ngày phát hành HSMT, tệp cùng bước lấy ngày của tệp
    cùng phiên bản (Cost_estimate.xlsx theo thuyết minh), phiên bản đường cơ sở lấy ngày đường cơ sở."""
    if r.get('ngay_tai_lieu'):
        return r['ngay_tai_lieu']
    if r.get('buoc') == 'hsmt' and HSMT_NGAY:
        return HSMT_NGAY
    for q, lst in reg['tai_lieu'].items():
        if q != p and q.startswith('ket_qua/'):
            x = next((y for y in lst if y.get('buoc') == r.get('buoc') and y['phien_ban'] == r['phien_ban'] and y.get('ngay_tai_lieu')), None)
            if x:
                return x['ngay_tai_lieu']
    bl = next((x for x in reg.get('duong_co_so', []) if x['ma'] == r.get('duong_co_so')), None)
    if bl:
        return loi.ngay(bl['ngay'])
    # tệp chưa ghi ngày (Excel, phụ lục): tài liệu cùng bước gần nhất có ngày; master schedule theo ngày bắt đầu dự án như tài liệu kế hoạch;
    # phụ lục bằng chứng ghi thời điểm chạy thực tế
    cung = [y for q, lst in reg['tai_lieu'].items() if q != p and q.startswith('ket_qua/') for y in lst
            if y.get('buoc') == r.get('buoc') and y.get('ngay_tai_lieu') and y.get('ngay', '') <= r.get('ngay', '')]
    if cung:
        return max(cung, key=lambda y: y.get('ngay', ''))['ngay_tai_lieu']
    if r.get('buoc') == 'lich':
        return loi.ngay(loi.du_an().get('ngay_bat_dau'))
    return loi.ngay(str(r.get('ngay', ''))[:10])


def ten_nguon_cap_nhat(x):
    """Tên tài liệu hoặc dữ liệu nguồn trong tai_lieu_cap_nhat của CR (không ghi đường dẫn tệp)."""
    k = x.split(' ')[0]
    if '/' not in k:
        return ten_buoc(k)
    dv = next((v['mo_ta'] for v in sd['dau_vao'].values() if v.get('tep') == k), '')
    if dv:
        return 'dữ liệu đầu vào ' + re.split(r'[:,(]', dv)[0].strip().lower()
    return {'wbs_nhap': 'dữ liệu nhập WBS', 'phan_cong_nhap': 'dữ liệu nhập giao việc', 'moc_nhap': 'dữ liệu nhập mốc của master schedule'}.get(
        os.path.splitext(os.path.basename(k))[0], 'dữ liệu nhập ' + os.path.splitext(os.path.basename(k))[0].replace('_', ' '))


def ngay_gio(v):
    v = str(v or '')
    return (loi.ngay(v[:10]) + (' ' + v[11:16] if len(v) > 10 else '')) if v else ''
danh_muc, du_lieu_ds, lich_su, luu_tru = [], [], [], []
thu_tu = {p: i for i, p in enumerate(p for b in sd['buoc'] for p in b['ra'].values())}


def thu_tu_tep(p):
    for q, i in thu_tu.items():
        if p == q or p.startswith(q.rstrip('/') + '/'):
            return i
    return 999


for p in sorted(reg['tai_lieu'], key=lambda p: (thu_tu_tep(p), p)):
    lst = reg['tai_lieu'][p]; last = lst[-1]; b = buoc.get(last['buoc'], {})
    if last['buoc'] == 'phien_ban':
        continue  # đầu ra của chính bước này không tự liệt kê
    doi = [] if b.get('anh_chup') else cu(last)  # báo cáo theo kỳ giữ số liệu tại thời điểm lập
    tinh_trang = ('Khóa' if b.get('khoa') and reg.get('hop_dong') else ('Cần cập nhật' if doi else 'Cập nhật'))
    con_tep = os.path.exists(loi.duong_dan(p))
    if not con_tep:
        tinh_trang = 'Không còn tệp'
    row = {'tep': p, 'ten_tep': os.path.basename(p), 'buoc': last['buoc'], 'ten': ten_buoc(last['buoc']), 'skill': b.get('skill', ''), 'nhom': b.get('nhom', ''),
           'phien_ban': last['phien_ban'], 'ngay': last['ngay'][:10], 'ngay_tai_lieu': ngay_tl(p, last), 'nguoi': last.get('nguoi', ''), 'trang_thai': last.get('trang_thai', ''),
           'tinh_trang': tinh_trang, 'dau_vao_da_doi': ', '.join(ten_nut(n) for n in doi), 'so_phien_ban': len(lst), 'duong_co_so': last.get('duong_co_so', ''),
           'phat_hanh': last.get('phat_hanh', '')}
    # tài liệu đã chuyển khỏi ket_qua (bản thay thế) không tính vào danh mục nộp, lịch sử vẫn giữ
    (danh_muc if p.startswith('ket_qua/') and con_tep else luu_tru if p.startswith('ket_qua/') else du_lieu_ds).append(row)
    for r in lst:
        lich_su.append({'tep': p, 'ten_tep': os.path.basename(p), 'phien_ban': r['phien_ban'], 'ngay': r['ngay'], 'thoi_diem': ngay_gio(r['ngay']), 'ngay_tai_lieu': ngay_tl(p, r),
                        'nguoi': r.get('nguoi', ''), 'buoc': r['buoc'], 'ten_buoc': ten_buoc(r['buoc']), 'trang_thai': r.get('trang_thai', ''),
                        'ghi_chu': str(r.get('ghi_chu', '')).replace('; ', ', '), 'hash': r['hash'],
                        'la_tai_lieu': p.startswith('ket_qua/'), 'phat_hanh': r.get('phat_hanh', ''), 'duong_co_so': r.get('duong_co_so', '')})
for i, r in enumerate(danh_muc, 1):
    r['ma'] = f'TL{i:02d}'
    r['ten'] = ten_tep_tl(r['tep'])
MA_TL = {r['tep']: r['ma'] for r in danh_muc}

cr = [{'ma': c['ma'], 'tieu_de': c['tieu_de'], 'trang_thai': c['trang_thai'], 'ngay_quyet_dinh': c.get('ngay_quyet_dinh', ''),
       'tai_lieu_cap_nhat': ', '.join(dict.fromkeys(ten_nguon_cap_nhat(x) for x in c.get('tai_lieu_cap_nhat', []))),
       'tai_lieu_khoa': ', '.join(ten_buoc(x.split(' ')[0]) for x in c.get('tai_lieu_khoa', [])),
       'phien_ban_sau': '. '.join(f'{ten_tep_tl(k)} {v}' for k, v in sorted((c.get('phien_ban_sau') or phien_ban_sau_cr(reg, c['ma'])).items(), key=lambda kv: MA_TL.get(kv[0], 'TL99'))
                                  if k.startswith('ket_qua/')),
       'phu_luc_hop_dong': 'Có' if c.get('phu_luc_hop_dong') else 'Không'} for c in so_td['yeu_cau']]

# Bản phát hành theo sprint (du_lieu/phat_hanh.json, qlda-master-schedule/scripts/phat_hanh.py) và tài liệu gắn từng bản
ph_ds = loi.doc_json('du_lieu/phat_hanh.json', {}).get('ban_phat_hanh', [])
NGUON_PH = {'thuc_te': 'Thực tế', 'ke_hoach': 'Kế hoạch'}
phat_hanh = [{'ma': x['ma'], 'sprint': x['sprint'], 'ngay': x['ngay'], 'loai': x['loai'], 'chinh_thuc': 'Có' if x.get('chinh_thuc') else 'Không',
              'nguon': NGUON_PH.get(x.get('nguon'), x.get('nguon', '')), 'so_goi': x.get('so_goi', 0), 'so_hang_muc_01a': len(x.get('hang_muc_01a', [])),
              'so_yeu_cau_rfp': len(x.get('yeu_cau_rfp', [])), 'so_loi_da_sua': x.get('so_loi_da_sua', 0), 'so_loi_con_mo': x.get('so_loi_con_mo', 0),
              'cr': ', '.join(c['ma'] for c in x.get('cr_dua_vao', [])) or 'Không', 'ghi_chu': x.get('ghi_chu', ''),
              'so_phien_ban_tai_lieu': sum(1 for r in lich_su if r['la_tai_lieu'] and r['phat_hanh'] == x['ma'])} for x in ph_ds]
# tài liệu thuộc từng đường cơ sở (BL1 ký hợp đồng, BL2 sau yêu cầu thay đổi được phê duyệt)
bl_tai_lieu = [{'ma_bl': x['ma'], 'ma_tl': MA_TL.get(p, ''), 'ten': ten_tep_tl(p), 'tep': p, 'ten_tep': os.path.basename(p), 'phien_ban': v,
                'phat_hanh': next((r.get('phat_hanh', '') for r in reg['tai_lieu'].get(p, []) if r.get('duong_co_so') == x['ma']), '')}
               for x in reg['duong_co_so'] for p, v in x['tai_lieu'].items()]

# Nhật ký chạy
import json  # noqa: E402
nk = []
pnk = loi.duong_dan('bang_chung/nhat_ky_chay.jsonl')
if os.path.exists(pnk):
    for line in open(pnk, encoding='utf-8'):
        try:
            nk.append(json.loads(line))
        except ValueError:
            pass
theo_skill = {}
for x in nk:
    sk = str(x.get('skill', '')).split()[0] if str(x.get('skill', '')).split() else ''  # bản ghi ghép nhầm mã bước vào tên skill: lấy tên skill (từ 1.2.6)
    t = theo_skill.setdefault(sk, {'skill': sk, 'phien_ban_skill': '', 'cac_phien_ban': set(), 'so_lan_script': 0, 'so_lan_agent': 0,
                                   'so_lan_nguoi_dung': 0, 'prompt': set()})
    t['so_lan_script' if x.get('loai') == 'script' else 'so_lan_agent' if x.get('loai') == 'agent' else 'so_lan_nguoi_dung'] += 1
    if x.get('phien_ban_skill'):
        t['cac_phien_ban'].add(x['phien_ban_skill'])
    if x.get('prompt'):
        t['prompt'].add(x['prompt'])


def khoa_phien_ban(v):
    return [int(p) if p.isdigit() else 0 for p in str(v).split('.')]


def khoang_phien_ban(ds):
    """Phiên bản skill đã dùng: một giá trị, hoặc khoảng từ bản đầu đến bản mới nhất (bản cuối là bản tạo tài liệu hiện hành)."""
    ds = sorted(ds, key=khoa_phien_ban)
    return '' if not ds else ds[0] if len(ds) == 1 else f'{ds[0]} - {ds[-1]}'


nhat_ky = [dict({k: v for k, v in t.items() if k != 'cac_phien_ban'}, phien_ban_skill=khoang_phien_ban(t['cac_phien_ban']),
                phien_ban_moi_nhat=max(t['cac_phien_ban'], key=khoa_phien_ban, default=''),
                prompt=', '.join(sorted(os.path.basename(p) for p in t['prompt']))) for t in theo_skill.values()]

# Git
git = []
try:
    r = subprocess.run(['git', '-C', loi.GOC, 'log', '--date=short', '--pretty=format:%h|%ad|%an|%s', '-n', '300'], capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=20)
    if r.returncode == 0:
        git = [dict(zip(('ma', 'ngay', 'tac_gia', 'noi_dung'), line.split('|', 3))) for line in r.stdout.splitlines() if line.count('|') >= 3]
except Exception:
    pass

_ngs = []
for r in danh_muc:
    try:
        _ngs.append(dt.datetime.strptime(r['ngay_tai_lieu'], '%d/%m/%Y').date())
    except (ValueError, TypeError):
        pass
out = {'ngay_tong_hop': loi.bay_gio()[:10], 'ngay_tai_lieu': loi.ngay(max(_ngs).isoformat()) if _ngs else loi.ngay(loi.bay_gio()[:10]), 'danh_muc': danh_muc, 'du_lieu': du_lieu_ds, 'lich_su': lich_su,
       'lich_su_tai_lieu': [x for x in lich_su if x['la_tai_lieu']], 'duong_co_so': reg['duong_co_so'],
       'duong_co_so_bang': [{'ma': x['ma'], 'ten': x['ten'], 'ngay': x['ngay'], 'so_tai_lieu': len(x['tai_lieu']), 'nguoi': x.get('nguoi', ''),
                             'ghi_chu': x.get('ghi_chu', ''), 'phat_hanh': x.get('phat_hanh', '') or 'Chưa có',
                             'can_cu': 'Ký hợp đồng' if x.get('hop_dong') or reg.get('hop_dong', {}).get('duong_co_so') == x['ma']
                             else (f'Yêu cầu thay đổi {x["ma_cr"]}' if x.get('ma_cr') else 'Quyết định của nhóm')} for x in reg['duong_co_so']],
       'duong_co_so_tai_lieu': bl_tai_lieu, 'phat_hanh': phat_hanh, 'so_ban_phat_hanh': len(phat_hanh),
       'ban_chinh_thuc': next((x['ma'] for x in ph_ds if x.get('chinh_thuc')), ''),
       'thay_doi': cr, 'nhat_ky': nhat_ky, 'so_lan_chay': len(nk), 'so_lan_agent': sum(1 for x in nk if x.get('loai') == 'agent'),
       'so_lan_nguoi_dung': sum(1 for x in nk if x.get('loai') not in ('script', 'agent')),  # quyết định, phụ lục do người dùng ghi (từ 1.2.6)
       'lam_lai': [dict(x, ngay_du_an_hien_thi=loi.ngay(x.get('ngay_du_an') or '')) for x in reg.get('lam_lai', [])],
       'duong_co_so_thay_the': [{'ma': x['ma'], 'ngay': x['ngay'], 'luc': x.get('luc', ''), 'thay_luc': x.get('thay_luc', ''), 'ly_do': x.get('ly_do', ''),
                                 'so_tai_lieu': len(x.get('tai_lieu', {})),
                                 'phien_ban': ', '.join(f'{os.path.basename(p)} {v}' for p, v in x.get('tai_lieu', {}).items())}
                                for x in reg.get('duong_co_so_thay_the', [])],
       'git': git, 'co_git': bool(git), 'hop_dong': reg.get('hop_dong', {}), 'kho_luu_tru': loi.du_an().get('kho_luu_tru', ''),
       'luu_tru': luu_tru, 'so_luu_tru': len(luu_tru),
       'so_tai_lieu': len(danh_muc), 'so_can_cap_nhat': sum(1 for r in danh_muc if r['tinh_trang'] == 'Cần cập nhật'),
       'so_phien_ban_tai_lieu': sum(1 for x in lich_su if x['la_tai_lieu']), 'quy_uoc': QU}
loi.ghi_json('du_lieu/phien_ban_tong_hop.json', out)

# Excel
from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
import xlsx_chung as X  # noqa: E402
OUT_DIR = loi.duong_dan('ket_qua/15_Quan_ly_phien_ban'); os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, 'So_phien_ban.xlsx')
wb = Workbook()
wb.properties.creator = wb.properties.lastModifiedBy = loi.du_an().get('nguoi_lap', '')
ws = wb.active; ws.title = 'Danh_muc'
X.tieu_de(ws, 1, ['Mã', 'Tài liệu', 'Tệp', 'Nhóm', 'Phiên bản', 'Ngày đăng ký', 'Trạng thái', 'Tình trạng', 'Số phiên bản', 'Đầu vào đã đổi', 'Bản phát hành'],
          [7, 36, 52, 18, 10, 12, 14, 14, 12, 30, 12])
for i, r in enumerate(danh_muc, 2):
    for j, k in enumerate(['ma', 'ten', 'tep', 'nhom', 'phien_ban', 'ngay', 'trang_thai', 'tinh_trang'], 1):
        X.o(ws, i, j, loi.ngay(r[k]) if k == 'ngay' else r[k], nhap=True)
    X.o(ws, i, 9, f'=COUNTIF(Lich_su!$A:$A,C{i})'); X.o(ws, i, 10, r['dau_vao_da_doi'], nhap=True); X.o(ws, i, 11, r['phat_hanh'], nhap=True)
if danh_muc:
    n_dm = len(danh_muc) + 1
    X.o(ws, n_dm + 1, 2, 'Tổng số phiên bản của tài liệu nộp hiện hành', nhap=True); X.o(ws, n_dm + 1, 9, f'=SUM(I2:I{n_dm})')
    X.o(ws, n_dm + 2, 2, 'Số phiên bản của tài liệu lưu trữ', nhap=True)
    _ltr = sorted({r['tep'] for r in luu_tru})  # tài liệu lưu trữ: đếm phiên bản trên sheet Lich_su bằng công thức
    X.o(ws, n_dm + 2, 9, '=' + '+'.join(f'COUNTIF(Lich_su!$A:$A,"{p}")' for p in _ltr) if _ltr else 0)
    X.o(ws, n_dm + 3, 2, 'Tổng số phiên bản tài liệu đã đăng ký', nhap=True); X.o(ws, n_dm + 3, 9, f'=I{n_dm + 1}+I{n_dm + 2}')
if danh_muc:
    ws.conditional_formatting.add(f'H2:H{len(danh_muc) + 1}', FormulaRule(formula=['H2="Cần cập nhật"'], fill=PatternFill('solid', fgColor='FCE4D6')))
ls = wb.create_sheet('Lich_su')
X.tieu_de(ls, 1, ['Tệp', 'Phiên bản', 'Ngày giờ', 'Người', 'Bước', 'Trạng thái', 'Ghi chú', 'Băm nội dung', 'Bản phát hành', 'Đường cơ sở'],
          [52, 10, 18, 18, 18, 14, 50, 14, 12, 10])
for i, r in enumerate(lich_su, 2):
    for j, k in enumerate(['tep', 'phien_ban', 'ngay', 'nguoi', 'buoc', 'trang_thai', 'ghi_chu', 'hash', 'phat_hanh', 'duong_co_so'], 1):
        X.o(ls, i, j, r['thoi_diem'] if k == 'ngay' else r[k], nhap=True)
dl = wb.create_sheet('Du_lieu')
X.tieu_de(dl, 1, ['Tệp dữ liệu', 'Bước', 'Phiên bản', 'Ngày', 'Tình trạng', 'Số phiên bản'], [40, 30, 10, 12, 14, 12])
for i, r in enumerate(du_lieu_ds, 2):
    for j, k in enumerate(['tep', 'ten', 'phien_ban', 'ngay', 'tinh_trang'], 1):
        X.o(dl, i, j, loi.ngay(r[k]) if k == 'ngay' else r[k], nhap=True)
    X.o(dl, i, 6, f'=COUNTIF(Lich_su!$A:$A,A{i})')
bl = wb.create_sheet('Duong_co_so')
X.tieu_de(bl, 1, ['Mã', 'Tên', 'Ngày', 'Tệp', 'Phiên bản', 'Bản phát hành', 'Căn cứ'], [8, 36, 12, 52, 10, 12, 26])
can_cu_bl = {x['ma']: x['can_cu'] for x in out['duong_co_so_bang']}
i = 2
for x in reg['duong_co_so']:
    for p, v in x['tai_lieu'].items():
        ph_ = next((y['phat_hanh'] for y in bl_tai_lieu if y['ma_bl'] == x['ma'] and y['tep'] == p), '')
        for j, val in enumerate([x['ma'], x['ten'], loi.ngay(x['ngay']), p, v, ph_, can_cu_bl[x['ma']]], 1):
            X.o(bl, i, j, val, nhap=True)
        i += 1
if reg.get('lam_lai'):  # ghi chú làm lại đường cơ sở, bản ghi bỏ khỏi sổ và nơi lưu sổ cũ (từ 1.2.4)
    i += 1
    X.o(bl, i, 1, 'Ghi chú làm lại', nhap=True)
    for x in reg['lam_lai']:
        i += 1
        X.o(bl, i, 1, loi.ngay(x.get('ngay_du_an') or ''), nhap=True); X.o(bl, i, 2, x.get('noi_dung', ''), nhap=True, wrap=True)
if out['duong_co_so_thay_the']:  # đường cơ sở đã thay khi lập lại: mã, lúc lập, lúc thay, lý do, phiên bản (từ 1.2.6)
    i += 1
    X.o(bl, i, 1, 'Đường cơ sở đã thay', nhap=True)
    for x in out['duong_co_so_thay_the']:
        i += 1
        for j, val in enumerate([x['ma'], f'Lập {ngay_gio(x["luc"])}, thay {ngay_gio(x["thay_luc"])}', loi.ngay(x['ngay']), x['phien_ban'], '', '', x['ly_do']], 1):
            X.o(bl, i, j, val, nhap=True, wrap=j in (4, 7))
phs = wb.create_sheet('Phat_hanh')
X.tieu_de(phs, 1, ['Mã', 'Sprint', 'Ngày', 'Loại', 'Nguồn số liệu', 'Số gói', 'Hạng mục 01A', 'Yêu cầu RFP', 'Lỗi đã sửa', 'Lỗi còn mở', 'CR đưa vào',
                   'Phiên bản tài liệu gắn bản', 'Ghi chú phát hành'], [6, 8, 12, 26, 12, 8, 10, 10, 10, 10, 12, 12, 70])
for i, r in enumerate(phat_hanh, 2):
    for j, k in enumerate(['ma', 'sprint', 'ngay', 'loai', 'nguon', 'so_goi', 'so_hang_muc_01a', 'so_yeu_cau_rfp', 'so_loi_da_sua', 'so_loi_con_mo', 'cr'], 1):
        X.o(phs, i, j, loi.ngay(r[k]) if k == 'ngay' else r[k], nhap=True, wrap=k == 'loai')  # ngày dd/mm/yyyy (từ 1.2.6)
    X.o(phs, i, 12, f'=COUNTIFS(Lich_su!$I:$I,A{i},Lich_su!$A:$A,"ket_qua/*")'); X.o(phs, i, 13, r['ghi_chu'], nhap=True, wrap=True)
td_ = wb.create_sheet('Thay_doi')
X.tieu_de(td_, 1, ['Mã CR', 'Nội dung', 'Trạng thái', 'Ngày quyết định', 'Tài liệu cập nhật', 'Phiên bản sau thay đổi', 'Tài liệu khóa', 'Phụ lục hợp đồng'],
          [8, 40, 14, 14, 50, 50, 22, 12])
for i, r in enumerate(cr, 2):
    for j, k in enumerate(['ma', 'tieu_de', 'trang_thai', 'ngay_quyet_dinh', 'tai_lieu_cap_nhat', 'phien_ban_sau', 'tai_lieu_khoa', 'phu_luc_hop_dong'], 1):
        X.o(td_, i, j, loi.ngay(r[k]) if k == 'ngay_quyet_dinh' and r[k] else r[k], nhap=True, wrap=j in (2, 5, 6))
nks = wb.create_sheet('Nhat_ky')
X.tieu_de(nks, 1, ['Skill', 'Phiên bản skill', 'Số lần chạy script', 'Số lần agent', 'Số lần người dùng ghi', 'Prompt đã dùng'], [26, 14, 16, 14, 14, 60])
for i, r in enumerate(nhat_ky, 2):
    for j, k in enumerate(['skill', 'phien_ban_skill', 'so_lan_script', 'so_lan_agent', 'so_lan_nguoi_dung', 'prompt'], 1):
        X.o(nks, i, j, r[k], nhap=True)
if git:
    gs = wb.create_sheet('Git')
    X.tieu_de(gs, 1, ['Commit', 'Ngày', 'Tác giả', 'Nội dung'], [10, 12, 20, 80])
    for i, r in enumerate(git, 2):
        for j, k in enumerate(['ma', 'ngay', 'tac_gia', 'noi_dung'], 1):
            X.o(gs, i, j, loi.ngay(r[k]) if k == 'ngay' else r[k], nhap=True)
wb.save(OUT)
kq = X.tinh_lai(OUT)
if kq and kq.get('status'):
    lc = X.dem_loi_cong_thuc(OUT)
    if lc:
        loi.thoat_loi('Công thức lỗi: ' + ', '.join(lc[:10]))
    v = X.doc_gia_tri(OUT)
    for i, r in enumerate(danh_muc, 2):
        if v['Danh_muc'].cell(row=i, column=9).value != r['so_phien_ban']:
            loi.thoat_loi(f'Đếm phiên bản trong Excel lệch ở {r["tep"]}')
    for i, r in enumerate(phat_hanh, 2):
        if v['Phat_hanh'].cell(row=i, column=12).value != r['so_phien_ban_tai_lieu']:
            loi.thoat_loi(f'Đếm phiên bản tài liệu gắn bản phát hành {r["ma"]} trong Excel lệch')
print(f'{len(danh_muc)} tài liệu, {out["so_phien_ban_tai_lieu"]} phiên bản tài liệu, {len(reg["duong_co_so"])} đường cơ sở, {len(phat_hanh)} bản phát hành, {len(cr)} yêu cầu thay đổi, '
      f'{len(nk)} lần chạy được ghi nhật ký, {"có" if git else "chưa có"} lịch sử git. Cần cập nhật: {out["so_can_cap_nhat"]}.')
h = loi.bam_obj({k: v for k, v in out.items() if k not in ('ngay_tong_hop', 'nhat_ky', 'so_lan_chay', 'so_lan_agent', 'so_lan_nguoi_dung')})
loi.hoan_tat_buoc('phien_ban', 'Tổng hợp sổ phiên bản', bo_qua=('doc_phien_ban',),
                  hash_tep={'ket_qua/15_Quan_ly_phien_ban/So_phien_ban.xlsx': h, 'du_lieu/phien_ban_tong_hop.json': h})
print('Tiếp theo: python .claude/skills/qlda-loi/scripts/chay.py chuan_bi phien_ban')
