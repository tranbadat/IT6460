"""Lập bảng phân chia lương theo cách trả lương do nhóm khai báo (don_gia.json cach_tra_luong).

Dùng: python chia_luong.py [--den yyyy-mm-dd]
  --den  tính đến ngày, mặc định toàn bộ nhật ký.
Đầu vào: dau_vao/thuc_te/nhat_ky_cong_viec.csv, dau_vao/thanh_vien.json, dau_vao/don_gia.json (luong_tham_chieu_thang,
         ngay_lam_viec_moi_thang, he_so_bao_hiem, cach_tra_luong, quy_chia_theo_dong_gop), du_lieu/lich.json (sprint, MD kế hoạch và
         MD huy động theo người, sprint), du_lieu/chi_phi.json (A, tong_chi_phi_nguoi), dau_vao/du_an.json.
Cách tính:
- Đơn giá ngày của vai trò = lương tham chiếu tháng / số ngày làm việc mỗi tháng.
- Ngày huy động theo lich.json phan_bo[].quy (cùng nguồn với cost estimate), tính đến ngày --den (qlda-bao-cao-tien-do/scripts/thuc_te.py).
- cach_tra_luong (nhóm quyết định, thiếu thì dừng):
    huy_dong: lương = ngày huy động x đơn giá ngày;
    cong_suc: lương = MD thực tế trong nhật ký x đơn giá ngày;
    ket_hop:  lương = ngày huy động x đơn giá ngày, cộng phần chia = quy_chia_theo_dong_gop.so_tien x tỷ trọng đóng góp
              (phần chia không nhân hệ số bảo hiểm).
- Chi phí doanh nghiệp = lương x hệ số bảo hiểm. Tỷ trọng đóng góp = MD thực tế của thành viên / tổng MD thực tế.
- Ngày sẵn sàng không ghi công = ngày huy động - MD thực tế, trình bày riêng. Đối chiếu MD kế hoạch (lich.json phan_bo, tính đến cùng ngày),
  khoản A và chi phí lương kế hoạch của cost estimate theo cùng cách trả lương. Huy động thực tế cộng dau_vao/thuc_te/huy_dong_bo_sung.json (từ 1.4.1).
Kiểm tra dữ liệu (lỗi chặn): L01 thành viên không có trong danh sách, L02 vai trò không có đơn giá, L04 vượt tỷ lệ tham gia trong ngày,
  L05 MD không hợp lệ, L09 chưa khai báo cach_tra_luong hoặc giá trị không hợp lệ, L10 ket_hop thiếu quy_chia_theo_dong_gop.so_tien.
  Cảnh báo: L03 ngày ngoài lịch dự án, L06 ngày nghỉ, L07 ghi công ngoài thời gian huy động, L08 chi phí lương kế hoạch lệch cost estimate.
Kỳ KT: lịch có giai đoạn Kết thúc thì ngày làm việc tính đến ngay_dong_du_an (thuc_te.ngay_lam_viec_du_an), kỳ KT là một cột như sprint
  (sheet Sprint, Theo_sprint), nhật ký sau cut over không bị L03. Có du_lieu/mo_phong.json (qlda-mo-phong) thì sheet Thuyet_minh ghi
  căn cứ số liệu mô phỏng.
Ra: du_lieu/luong.json, ket_qua/14_Bang_phan_chia_luong/Bang_phan_chia_luong.xlsx (công thức SUMIFS, INDEX/MATCH, đã tính lại và đối chiếu).
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import argparse
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as CL

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'qlda-loi', 'scripts'))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'qlda-bao-cao-tien-do', 'scripts'))
import loi  # noqa: E402
import thuc_te  # noqa: E402
import xlsx_chung as X  # noqa: E402

CACH = thuc_te.CACH_TRA_LUONG  # cùng danh mục với cost estimate, báo cáo tiến độ, tổng kết
ap = argparse.ArgumentParser(); ap.add_argument('--den'); ap.add_argument('--ghi-chu', default='', help='ghi chú phiên bản (lý do lập lại)')
a = ap.parse_args()
D = dt.date.fromisoformat
da = loi.du_an(); dg = loi.doc_json('dau_vao/don_gia.json'); lich = loi.doc_json('du_lieu/lich.json')
cp = loi.doc_json('du_lieu/chi_phi.json', {}); tvs = loi.doc_json('dau_vao/thanh_vien.json')['thanh_vien']; tv = {t['id']: t for t in tvs}
cach = thuc_te.cach_tra_luong(dg)
if not cach:
    loi.thoat_loi(thuc_te.thong_bao_thieu_cach('L09'))
quy_dg = dg.get('quy_chia_theo_dong_gop') or {}
if cach == 'ket_hop' and (not isinstance(quy_dg, dict) or quy_dg.get('so_tien') in (None, '')):
    loi.thoat_loi('L10: cach_tra_luong ket_hop cần dau_vao/don_gia.json quy_chia_theo_dong_gop {so_tien, can_cu}: tổng số tiền chia theo tỷ trọng '
                  'đóng góp (MD thực tế), không nhân hệ số bảo hiểm. Số tiền do nhóm quyết định.')
QUY = float(quy_dg.get('so_tien') or 0) if cach == 'ket_hop' else 0.0
nk = loi.doc_csv('dau_vao/thuc_te/nhat_ky_cong_viec.csv')
if not nk:
    loi.thoat_loi('Chưa có nhật ký công việc thực tế (dau_vao/thuc_te/nhat_ky_cong_viec.csv), xem qlda-bao-cao-tien-do/references/nhap_lieu.md.')
sai_ngay = []
for i, r in enumerate(nk, 2):
    try:
        D(r['ngay'])
    except (ValueError, KeyError):
        sai_ngay.append(f'dòng {i}: "{r.get("ngay")}"')
if sai_ngay:
    loi.thoat_loi('nhat_ky_cong_viec.csv có ngày không đúng dạng yyyy-mm-dd: ' + ', '.join(sai_ngay[:10]))
den = D(a.den) if a.den else max(D(r['ngay']) for r in nk)
nk = [r for r in nk if D(r['ngay']) <= den]
F = []
add = lambda ma, mo, md='loi': F.append({'ma': ma, 'muc_do': md, 'mo_ta': mo})  # noqa: E731
ds_ngay = thuc_te.ngay_lam_viec_du_an(da, lich); days = set(ds_ngay); nghi = set(da.get('ngay_nghi', []))
sp = thuc_te.cac_ky(lich, da)  # sprint, cộng kỳ KT khi lịch có giai đoạn Kết thúc
sp_idx = {s['ma']: i for i, s in enumerate(sp, 1)}


def sprint_cua(d):
    for s in sp:
        if D(s['tu']) <= d <= D(s['den']):
            return s['ma']
    return ''


ngay_tv = {}
for i, r in enumerate(nk, 2):
    p = r['thanh_vien']
    try:
        v = float(r['md'])
        if v <= 0 or v > 1:
            raise ValueError
    except ValueError:
        add('L05', f'Dòng {i}: MD "{r["md"]}" không hợp lệ, ghi số trong khoảng (0, 1]'); continue
    if p not in tv:
        add('L01', f'Dòng {i}: thành viên "{p}" không có trong thanh_vien.json'); continue
    d = D(r['ngay'])
    if d.isoformat() in nghi or d.weekday() >= 5:
        add('L06', f'Dòng {i}: {tv[p]["ho_ten"]} ghi công ngày nghỉ {loi.ngay(d)}', 'canh_bao')
    elif d not in days:
        add('L03', f'Dòng {i}: ngày {loi.ngay(d)} ngoài lịch dự án', 'canh_bao')
    s = sprint_cua(d); t = tv[p]
    if s and ((t.get('tu_sprint') and sp_idx[s] < int(t['tu_sprint'])) or (t.get('den_sprint') and sp_idx[s] > int(t['den_sprint']))):
        add('L07', f'Dòng {i}: {t["ho_ten"]} ghi công sprint {s}, ngoài thời gian huy động', 'canh_bao')
    ngay_tv[(p, d)] = ngay_tv.get((p, d), 0) + v
for (p, d), v in ngay_tv.items():
    if v > float(tv[p].get('ty_le', 1)) + 1e-6:
        add('L04', f'{tv[p]["ho_ten"]} ghi {loi.so(v, 2)} MD ngày {loi.ngay(d)}, vượt tỷ lệ tham gia {loi.pt(float(tv[p].get("ty_le", 1)), 0)}')
hd, hd_sp = thuc_te.md_huy_dong(lich, tvs, ds_ngay, den)
HD_BS = thuc_te.huy_dong_bo_sung_tom_tat()
# L07 theo quỹ sprint: MD ghi công của người trong sprint vượt quỹ huy động (kế hoạch cộng huy động bổ sung), từ 1.4.1
md_sp_tv = {}
for r in nk:
    if r['thanh_vien'] in tv:
        try:
            k_ = (r['thanh_vien'], sprint_cua(D(r['ngay'])))
            md_sp_tv[k_] = md_sp_tv.get(k_, 0) + float(r['md'])
        except (ValueError, TypeError):
            pass
for (p_, s_), v_ in sorted(md_sp_tv.items(), key=lambda kv: (kv[0][1] or '', kv[0][0])):
    if s_ and v_ > hd_sp.get((p_, s_), 0) + 1e-6:
        add('L07', f'{tv[p_]["ho_ten"]} ghi {loi.so(v_, 2)} MD trong {s_}, vượt huy động {loi.so(hd_sp.get((p_, s_), 0), 2)} MD (kế hoạch và huy động bổ sung)', 'canh_bao')
vai = sorted({tv[r['thanh_vien']]['vai_tro'] for r in nk if r['thanh_vien'] in tv} | {tv[p]['vai_tro'] for p, q in hd.items() if q and p in tv})
for v in vai:
    if loi.don_gia_ngay_vai_tro(dg, v) is None:
        add('L02', f'Vai trò {v} chưa có lương tham chiếu trong don_gia.json')
loi_ = [f for f in F if f['muc_do'] == 'loi']
if not loi_:
    lech_ce = thuc_te.doi_chieu_cost_estimate(cp, thuc_te.chi_phi_luong_ke_hoach(cach, lich, tvs, dg, ds_ngay, lich.get('tong_md')), cach)
    if lech_ce:
        add('L08', lech_ce, 'canh_bao')
for f in F:
    print(f"[{f['muc_do']}] {f['ma']}: {f['mo_ta']}")
if loi_:
    loi.ghi_json('lam_viec/bang_luong/bao_cao_kiem_tra.json', {'ket_luan': 'chua_dat', 'phat_hien': F})
    loi.thoat_loi(f'{len(loi_)} lỗi dữ liệu nhật ký, sửa rồi chạy lại.')

# Tính bằng Python (đối chiếu với Excel)
hs = float(dg.get('he_so_bao_hiem', 1.0)); nlv = float(dg.get('ngay_lam_viec_moi_thang', 22))
dgn = {v: dg['luong_tham_chieu_thang'][v]['luong'] / nlv for v in dg['luong_tham_chieu_thang']}
kh = {}
tl_sp = thuc_te.ty_le_sprint_da_qua(lich, ds_ngay, den)  # MD kế hoạch tính đến cùng ngày với nhật ký
for x in lich['phan_bo']:
    kh[x['thanh_vien']] = kh.get(x['thanh_vien'], 0) + float(x['md']) * tl_sp.get(x['sprint'], 0)
md_tt = {}; md_sp = {}
for r in nk:
    if r['thanh_vien'] in tv:
        md_tt[r['thanh_vien']] = md_tt.get(r['thanh_vien'], 0) + float(r['md'])
        k = (r['thanh_vien'], sprint_cua(D(r['ngay'])))
        md_sp[k] = md_sp.get(k, 0) + float(r['md'])
tong_md_tt = sum(md_tt.values())
THEO_HD = cach in ('huy_dong', 'ket_hop')


def ngay_luong_sprint(p, s):
    return hd_sp.get((p, s), 0) if THEO_HD else md_sp.get((p, s), 0)


nguoi = []
for t in tvs:
    if not kh.get(t['id']) and not md_tt.get(t['id']) and not hd.get(t['id']):
        continue
    m = md_tt.get(t['id'], 0); q = hd.get(t['id'], 0); g = dgn[t['vai_tro']]; nl = q if THEO_HD else m
    tt_dg = m / tong_md_tt if tong_md_tt else 0
    nguoi.append({'id': t['id'], 'ho_ten': t['ho_ten'], 'loai': 'Thành viên nhóm' if t['loai'] == 'nhom' else 'Vị trí bổ sung', 'vai_tro': t['vai_tro'],
                  'ten_vai_tro': dg['luong_tham_chieu_thang'][t['vai_tro']]['ten'], 'md_ke_hoach': round(kh.get(t['id'], 0), 2), 'md_thuc_te': round(m, 2),
                  'md_huy_dong': round(q, 2), 'md_san_sang': round(q - m, 2), 'chenh_lech_md': round(m - kh.get(t['id'], 0), 2),
                  'ty_trong_dong_gop': round(tt_dg, 4), 'ngay_tinh_luong': round(nl, 2), 'don_gia_ngay': round(g),
                  'luong': nl * g, 'bao_hiem': nl * g * (hs - 1), 'chi_phi': nl * g * hs, 'phan_chia_dong_gop': QUY * tt_dg,
                  'thu_nhap': nl * g + QUY * tt_dg, 'luong_ke_hoach': (q if THEO_HD else kh.get(t['id'], 0)) * g,
                  'theo_sprint': [{'sprint': s['ma'], 'md': round(md_sp.get((t['id'], s['ma']), 0), 2),
                                   'md_huy_dong': round(hd_sp.get((t['id'], s['ma']), 0), 2)} for s in sp]})
tong_luong = sum(x['luong'] for x in nguoi)
for x in nguoi:
    x['ty_trong'] = x['ty_trong_dong_gop']
    for k in ('luong', 'bao_hiem', 'chi_phi', 'phan_chia_dong_gop', 'thu_nhap', 'luong_ke_hoach'):
        x[k] = round(x[k])
tong = {'md_ke_hoach': round(sum(x['md_ke_hoach'] for x in nguoi), 2), 'md_thuc_te': round(tong_md_tt, 2),
        'md_huy_dong': round(sum(hd.get(x['id'], 0) for x in nguoi), 2), 'md_san_sang': round(sum(hd.get(x['id'], 0) for x in nguoi) - tong_md_tt, 2),
        'luong': round(tong_luong), 'bao_hiem': round(tong_luong * (hs - 1)), 'chi_phi': round(tong_luong * hs),
        'phan_chia_dong_gop': round(QUY if tong_md_tt else 0), 'thu_nhap': round(tong_luong + (QUY if tong_md_tt else 0)),
        'luong_ke_hoach': round(sum((hd.get(x['id'], 0) if THEO_HD else kh.get(x['id'], 0)) * dgn[x['vai_tro']] for x in nguoi)),
        'luong_theo_md_ke_hoach': round(sum(kh.get(x['id'], 0) * dgn[x['vai_tro']] for x in nguoi))}  # theo MD công việc kế hoạch, để tham khảo
A = float(cp.get('A', 0)); ce_hd = float(cp.get('tong_chi_phi_nguoi') or 0)
doi_chieu = {'A': A, 'ty_le_chi_phi_tren_A': round((tong['chi_phi'] + tong['phan_chia_dong_gop']) / A, 4) if A else 0,
             'chi_phi_ke_hoach': round(tong['luong_ke_hoach'] * hs), 'chenh_chi_phi': round(tong['chi_phi'] - tong['luong_ke_hoach'] * hs),
             'chi_phi_theo_md_ke_hoach': round(tong['luong_theo_md_ke_hoach'] * hs),
             'chi_phi_luong_cost_estimate': ce_hd, 'ty_le_su_dung_huy_dong': round(tong_md_tt / tong['md_huy_dong'], 4) if tong['md_huy_dong'] else 0,
             'he_so_bao_hiem': hs, 'ngay_lam_viec_moi_thang': nlv}
theo_sprint = [{'sprint': s['ma'], 'tu': s['tu'], 'den': s['den'], 'md': round(sum(md_sp.get((x['id'], s['ma']), 0) for x in nguoi), 2),
                'md_huy_dong': round(sum(hd_sp.get((x['id'], s['ma']), 0) for x in nguoi), 2),
                'luong': round(sum(ngay_luong_sprint(x['id'], s['ma']) * dgn[x['vai_tro']] for x in nguoi))} for s in sp]

# Excel
OUT_DIR = loi.duong_dan('ket_qua/14_Bang_phan_chia_luong'); os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, 'Bang_phan_chia_luong.xlsx')
wb = Workbook()
wb.properties.creator = wb.properties.lastModifiedBy = loi.du_an().get('nguoi_lap', '')
tm = wb.active; tm.title = 'Thuyet_minh'
MP = thuc_te.mo_phong()
CAN_CU_MP = ([('Số liệu thực tế trong bảng là số liệu mô phỏng. ' + MP.get('_ghi_chu', '')).strip()] if MP else [])
nl_txt = 'ngày huy động' if THEO_HD else 'MD thực tế'
TEN_NHOM = ', '.join(dict.fromkeys(x for x in (da.get('ma_nhom', ''), da.get('don_vi_thuc_hien', '')) if x))  # tên nhóm trùng tên nhà thầu tạm thì ghi một lần (từ 1.4.1)
dong = [('Bảng phân chia lương', X.F_TEN), (da['ten_du_an'], None), (f'Nhóm {TEN_NHOM}', None),
        (f'Số liệu tính đến ngày {loi.ngay(den)}. Cách trả lương do nhóm khai báo: {CACH[cach]}.', None), ('', None),
        ('Căn cứ', X.F_TIEU_DE),
        ('1. Công sức thực tế theo nhật ký công việc (sheet Nhat_ky), mỗi dòng là MD của một thành viên cho một gói WBS trong một ngày.', None),
        ('2. Ngày huy động theo kế hoạch huy động của master schedule (sheet Huy_dong), cùng nguồn với cost estimate, tính đến ngày trên'
         + (f', cộng huy động bổ sung thực tế theo nhu cầu {loi.md(HD_BS["tong"])} MD theo {HD_BS["nguon"]}.' if HD_BS['co'] else '.'), None),
        ('3. Đơn giá ngày của vai trò bằng lương tham chiếu tháng chia số ngày làm việc mỗi tháng (sheet Tham_so).', None),
        (f'4. Lương bằng ngày tính lương nhân đơn giá ngày, ngày tính lương là {nl_txt}. Chi phí doanh nghiệp gồm lương và phần bảo hiểm theo hệ số.', None),
        ('5. Tỷ trọng đóng góp là tỷ lệ MD thực tế của thành viên trên tổng MD thực tế.'
         + (' Phần chia theo đóng góp bằng quỹ chia (sheet Tham_so) nhân tỷ trọng đóng góp, không nhân hệ số bảo hiểm.' if cach == 'ket_hop' else ''), None),
        ('6. Ngày sẵn sàng không ghi công bằng ngày huy động trừ MD thực tế, trình bày riêng để đối chiếu, không dùng để trừ lương.', None),
        ('7. Bảng đối chiếu MD kế hoạch theo master schedule, khoản A và chi phí lương kế hoạch của cost estimate theo cùng cách trả lương ('
         + ('MD công việc' if cach == 'cong_suc' else 'ngày huy động') + ', sheet Doi_chieu).', None),
        ] + [('8. ' + t, None) for t in CAN_CU_MP] + [
        ('', None), ('Ký hiệu', X.F_TIEU_DE), ('Chữ xanh dương: số liệu nhập. Chữ xanh lá: tham chiếu sheet khác. Chữ đen: công thức.', None),
        ('', None), ('Xác nhận', X.F_TIEU_DE), ('Thành viên ký xác nhận tại cột cuối của sheet Bang_luong.', None),
        (f'Người lập: {da.get("nguoi_lap", "")}', None)]
for i, (t, f) in enumerate(dong, 1):
    c = tm.cell(row=i, column=1, value=t); c.font = f or X.F_THUONG
tm.column_dimensions['A'].width = 120

ts = wb.create_sheet('Tham_so')
X.tieu_de(ts, 1, ['Vai trò', 'Tên vai trò', 'Lương tham chiếu tháng, đồng', 'Đơn giá ngày, đồng'], [16, 34, 22, 20])
vt = list(dg['luong_tham_chieu_thang'].items())
nv = len(vt) + 1
for i, (k, v) in enumerate(vt, 2):
    X.o(ts, i, 1, k, nhap=True); X.o(ts, i, 2, v['ten'], nhap=True); X.o(ts, i, 3, v['luong'], fmt=X.SO, nhap=True)
    X.o(ts, i, 4, f'=C{i}/$B${nv + 3}', fmt=X.SO)
X.o(ts, nv + 2, 1, 'Hệ số bảo hiểm', dam=True); X.o(ts, nv + 2, 2, hs, fmt='0.000', nhap=True)
X.o(ts, nv + 3, 1, 'Ngày làm việc mỗi tháng', dam=True); X.o(ts, nv + 3, 2, nlv, fmt='0', nhap=True)
X.o(ts, nv + 4, 1, 'Cách trả lương', dam=True); X.o(ts, nv + 4, 2, cach, nhap=True); X.o(ts, nv + 4, 3, CACH[cach])
X.o(ts, nv + 5, 1, 'Quỹ chia theo đóng góp, đồng', dam=True); X.o(ts, nv + 5, 2, QUY, fmt=X.SO, nhap=True)
X.ghi_chu_bang(ts, nv + 7, 'Nguồn lương tham chiếu: ' + str(dg.get('nguon_luong', '')))
X.ghi_chu_bang(ts, nv + 8, 'Nguồn hệ số bảo hiểm: ' + str(dg.get('nguon_he_so_bao_hiem', '')))
if cach == 'ket_hop':
    X.ghi_chu_bang(ts, nv + 9, 'Căn cứ quỹ chia theo đóng góp: ' + str(quy_dg.get('can_cu', '')))
HS, CA, QC = f'Tham_so!$B${nv + 2}', f'Tham_so!$B${nv + 4}', f'Tham_so!$B${nv + 5}'

ss = wb.create_sheet('Sprint')
X.tieu_de(ss, 1, ['Sprint', 'Từ ngày', 'Đến ngày'], [10, 14, 14])
for i, s in enumerate(sp, 2):
    X.o(ss, i, 1, s['ma'], nhap=True); X.o(ss, i, 2, D(s['tu']), fmt=X.NGAY, nhap=True); X.o(ss, i, 3, D(s['den']), fmt=X.NGAY, nhap=True)
ns = len(sp) + 1

nks = wb.create_sheet('Nhat_ky')
X.tieu_de(nks, 1, ['Ngày', 'Sprint', 'Mã thành viên', 'Mã gói WBS', 'MD', 'Ghi chú'], [12, 9, 14, 12, 8, 40])
for i, r in enumerate(sorted(nk, key=lambda r: (r['ngay'], r['thanh_vien'], r['ma_goi'])), 2):
    X.o(nks, i, 1, D(r['ngay']), fmt=X.NGAY, nhap=True)
    X.o(nks, i, 2, f'=IFERROR(INDEX(Sprint!$A$2:$A${ns},MATCH(A{i},Sprint!$B$2:$B${ns},1)),"")', lien_ket=True)
    X.o(nks, i, 3, r['thanh_vien'], nhap=True); X.o(nks, i, 4, r['ma_goi'], nhap=True)
    X.o(nks, i, 5, float(r['md']), fmt=X.SO2, nhap=True); X.o(nks, i, 6, r.get('ghi_chu', ''), nhap=True)
nn = len(nk) + 1

hds = wb.create_sheet('Huy_dong')
X.tieu_de(hds, 1, ['Mã thành viên', 'Sprint', 'MD huy động đến ngày tính', 'Theo kế hoạch', 'Huy động bổ sung'], [14, 9, 16, 14, 14])
HD_BS_SP = thuc_te.huy_dong_bo_sung()  # phần huy động bổ sung thực tế theo người và sprint, tách cột cho người ký xác nhận (từ 1.4.2)
hd_rows = [(p, s['ma'], round(hd_sp.get((p, s['ma']), 0), 4)) for p in (x['id'] for x in nguoi) for s in sp if hd_sp.get((p, s['ma']))]
for i, (p, s, v) in enumerate(hd_rows, 2):
    bs = min(v, HD_BS_SP.get((p, s), 0))
    X.o(hds, i, 1, p, nhap=True); X.o(hds, i, 2, s, nhap=True); X.o(hds, i, 3, f'=D{i}+E{i}', fmt=X.SO2)
    X.o(hds, i, 4, round(v - bs, 4), fmt=X.SO2, nhap=True); X.o(hds, i, 5, bs, fmt=X.SO2, nhap=True)
nh = max(len(hd_rows) + 1, 2)
X.ghi_chu_bang(hds, nh + 2, 'Nguồn: kế hoạch huy động của master schedule (phan_bo quỹ theo sprint), sprint đang chạy tính theo tỷ lệ ngày làm việc đã qua; '
               'huy động bổ sung theo sổ huy động bổ sung thực tế (dau_vao/thuc_te/huy_dong_bo_sung.json).')

bl = wb.create_sheet('Bang_luong', 1)
HDR = ['STT', 'Mã thành viên', 'Họ tên', 'Loại', 'Vai trò', 'MD kế hoạch', 'MD huy động', 'MD thực tế', 'Ngày sẵn sàng không ghi công',
       'Chênh lệch MD so với kế hoạch', 'Tỷ trọng đóng góp', 'Đơn giá ngày, đồng', 'Ngày tính lương', 'Lương, đồng', 'Bảo hiểm phần doanh nghiệp, đồng',
       'Chi phí doanh nghiệp, đồng', 'Phần chia theo đóng góp, đồng', 'Thu nhập, đồng', 'Ký xác nhận']
X.tieu_de(bl, 1, HDR, [5, 13, 24, 15, 11, 10, 10, 10, 12, 12, 10, 14, 10, 16, 18, 18, 16, 16, 16])
r0 = 2; rl = r0 + len(nguoi) - 1; rt = rl + 1
for i, x in enumerate(nguoi):
    r = r0 + i
    X.o(bl, r, 1, i + 1); X.o(bl, r, 2, x['id'], nhap=True); X.o(bl, r, 3, x['ho_ten'], nhap=True); X.o(bl, r, 4, x['loai'], nhap=True)
    X.o(bl, r, 5, x['vai_tro'], nhap=True); X.o(bl, r, 6, x['md_ke_hoach'], fmt=X.SO2, nhap=True)
    X.o(bl, r, 7, f'=SUMIFS(Huy_dong!$C$2:$C${nh},Huy_dong!$A$2:$A${nh},B{r})', fmt=X.SO2, lien_ket=True)
    X.o(bl, r, 8, f'=SUMIFS(Nhat_ky!$E$2:$E${nn},Nhat_ky!$C$2:$C${nn},B{r})', fmt=X.SO2, lien_ket=True)
    X.o(bl, r, 9, f'=G{r}-H{r}', fmt=X.SO2); X.o(bl, r, 10, f'=H{r}-F{r}', fmt=X.SO2)
    X.o(bl, r, 11, f'=IF($H${rt}=0,0,H{r}/$H${rt})', fmt=X.PT)
    X.o(bl, r, 12, f'=INDEX(Tham_so!$D$2:$D${nv},MATCH(E{r},Tham_so!$A$2:$A${nv},0))', fmt=X.SO, lien_ket=True)
    X.o(bl, r, 13, f'=IF({CA}="cong_suc",H{r},G{r})', fmt=X.SO2)
    X.o(bl, r, 14, f'=M{r}*L{r}', fmt=X.SO); X.o(bl, r, 15, f'=N{r}*({HS}-1)', fmt=X.SO); X.o(bl, r, 16, f'=N{r}+O{r}', fmt=X.SO)
    X.o(bl, r, 17, f'=IF({CA}="ket_hop",{QC}*K{r},0)', fmt=X.SO); X.o(bl, r, 18, f'=N{r}+Q{r}', fmt=X.SO); X.o(bl, r, 19, '')
X.o(bl, rt, 3, 'Tổng cộng', dam=True)
for c in (6, 7, 8, 9, 10, 11, 13, 14, 15, 16, 17, 18):
    L = CL(c)
    X.o(bl, rt, c, f'=SUM({L}{r0}:{L}{rl})', fmt=X.PT if c == 11 else (X.SO2 if c in (6, 7, 8, 9, 10, 13) else X.SO), dam=True)
bl.conditional_formatting.add(f'J{r0}:J{rl}', FormulaRule(formula=[f'ABS(J{r0})>0.1*F{r0}'], fill=PatternFill('solid', fgColor='FCE4D6')))
X.ghi_chu_bang(bl, rt + 2, 'Ô chênh lệch tô màu khi MD thực tế lệch quá 10% so với kế hoạch. Ngày tính lương theo cách trả lương ở sheet Tham_so.')

tsp = wb.create_sheet('Theo_sprint', 2)
hdr = ['Mã thành viên', 'Họ tên'] + [f'MD {s["ma"]}' for s in sp] + ['Tổng MD'] + [f'Lương {s["ma"]}' for s in sp] + ['Tổng lương']
X.tieu_de(tsp, 1, hdr, [13, 24] + [9] * len(sp) + [10] + [14] * len(sp) + [16])
for i, x in enumerate(nguoi):
    r = 2 + i
    X.o(tsp, r, 1, x['id'], nhap=True); X.o(tsp, r, 2, x['ho_ten'], nhap=True)
    for j, s in enumerate(sp):
        X.o(tsp, r, 3 + j, f'=SUMIFS(Nhat_ky!$E$2:$E${nn},Nhat_ky!$C$2:$C${nn},$A{r},Nhat_ky!$B$2:$B${nn},"{s["ma"]}")', fmt=X.SO2, lien_ket=True)
    ct = 3 + len(sp)
    X.o(tsp, r, ct, f'=SUM({CL(3)}{r}:{CL(ct - 1)}{r})', fmt=X.SO2)
    dg_o = f'INDEX(Bang_luong!$L${r0}:$L${rl},MATCH($A{r},Bang_luong!$B${r0}:$B${rl},0))'
    for j, s in enumerate(sp):
        hd_o = f'SUMIFS(Huy_dong!$C$2:$C${nh},Huy_dong!$A$2:$A${nh},$A{r},Huy_dong!$B$2:$B${nh},"{s["ma"]}")'
        X.o(tsp, r, ct + 1 + j, f'=IF({CA}="cong_suc",{CL(3 + j)}{r},{hd_o})*{dg_o}', fmt=X.SO)
    X.o(tsp, r, ct + 1 + len(sp), f'=SUM({CL(ct + 1)}{r}:{CL(ct + len(sp))}{r})', fmt=X.SO)
rtt = 2 + len(nguoi)
X.o(tsp, rtt, 2, 'Tổng cộng', dam=True)
for c in range(3, len(hdr) + 1):
    X.o(tsp, rtt, c, f'=SUM({CL(c)}2:{CL(c)}{rtt - 1})', fmt=X.SO2 if c <= 3 + len(sp) else X.SO, dam=True)

dc = wb.create_sheet('Doi_chieu', 3)
X.tieu_de(dc, 1, ['Chỉ tiêu', 'Giá trị', 'Ghi chú'], [56, 20, 60])
rows = [('Khoản A của cost estimate (chi phí nhân công theo đơn giá trọn gói), đồng', A, 'du_lieu/chi_phi.json', True, X.SO),
        ('Chi phí lương kế hoạch cả dự án trong cost estimate (cùng cách trả lương), đồng', ce_hd, 'du_lieu/chi_phi.json tong_chi_phi_nguoi', True, X.SO),
        ('Tổng MD kế hoạch (master schedule)', f'=Bang_luong!F{rt}', '', False, X.SO2),
        ('Tổng MD huy động đến ngày tính', f'=Bang_luong!G{rt}', 'Sheet Huy_dong', False, X.SO2),
        ('Tổng MD thực tế (nhật ký)', f'=Bang_luong!H{rt}', '', False, X.SO2),
        ('Ngày sẵn sàng không ghi công', f'=Bang_luong!I{rt}', 'MD huy động trừ MD thực tế', False, X.SO2),
        ('Tỷ lệ sử dụng nhân lực huy động', '=IF(B5=0,0,B6/B5)', 'MD thực tế trên MD huy động', False, X.PT),
        ('Tổng lương theo cách trả lương, đồng', f'=Bang_luong!N{rt}', CACH[cach], False, X.SO),
        ('Chi phí doanh nghiệp, đồng', f'=Bang_luong!P{rt}', 'Lương nhân hệ số bảo hiểm', False, X.SO),
        ('Phần chia theo đóng góp, đồng', f'=Bang_luong!Q{rt}', 'Chỉ khi cách trả lương là ket_hop', False, X.SO),
        ('Chi phí doanh nghiệp theo MD công việc kế hoạch (tham khảo), đồng', f'=SUMPRODUCT(Bang_luong!F{r0}:F{rl},Bang_luong!L{r0}:L{rl})*{HS}', '', False, X.SO),
        ('Chênh lệch chi phí doanh nghiệp so với chi phí lương kế hoạch của cost estimate, đồng', '=B10-B3',
         'Tính đến hết dự án: bằng 0 khi trả theo huy động, là chênh lệch MD thực tế so với MD giao khi trả theo công sức', False, X.SO),
        ('Tỷ lệ chi phí doanh nghiệp và phần chia trên khoản A', '=IF(B2=0,0,(B10+B11)/B2)', 'Phần còn lại của A bù chi phí quản lý và lợi nhuận', False, X.PT)]
for i, (t, v, g, nh_, f) in enumerate(rows, 2):
    X.o(dc, i, 1, t); X.o(dc, i, 2, v, fmt=f, nhap=nh_, lien_ket=not nh_); X.o(dc, i, 3, g)
wb.save(OUT)

kq = X.tinh_lai(OUT)
ghi = 'Máy không có LibreOffice, Excel tự tính khi mở.'
if kq and kq.get('status'):
    lc = X.dem_loi_cong_thuc(OUT)
    if lc:
        loi.thoat_loi('Công thức lỗi: ' + ', '.join(lc[:10]))
    v = X.doc_gia_tri(OUT)
    so_ = [('tổng MD thực tế', v['Bang_luong'][f'H{rt}'].value, tong['md_thuc_te']), ('tổng MD huy động', v['Bang_luong'][f'G{rt}'].value, tong['md_huy_dong']),
           ('tổng lương', v['Bang_luong'][f'N{rt}'].value, tong['luong']), ('chi phí doanh nghiệp', v['Bang_luong'][f'P{rt}'].value, tong['chi_phi']),
           ('thu nhập', v['Bang_luong'][f'R{rt}'].value, tong['thu_nhap']),
           ('tổng lương theo sprint', v['Theo_sprint'].cell(row=rtt, column=len(hdr)).value, tong['luong']),
           ('tổng MD theo sprint', v['Theo_sprint'].cell(row=rtt, column=3 + len(sp)).value, tong['md_thuc_te'])]
    for ten, a_, b_ in so_:
        if a_ is None or abs(float(a_) - float(b_)) > 2:
            loi.thoat_loi(f'Đối chiếu Excel và Python lệch ở {ten}: {a_} khác {b_}. Kiểm tra ngày nhật ký nằm trong sprint.')
    for x in nguoi:
        r = r0 + nguoi.index(x)
        if abs(float(v['Bang_luong'][f'R{r}'].value) - x['thu_nhap']) > 2:
            loi.thoat_loi(f'Đối chiếu thu nhập của {x["ho_ten"]} lệch giữa Excel và Python')
    ghi = 'Đã tính lại công thức, 0 lỗi, giá trị Excel khớp tính toán Python (MD, ngày huy động, lương, thu nhập từng người, tổng theo sprint).'

out = {'den': den.isoformat(), 'so_lieu_mo_phong': bool(MP), 'can_cu_mo_phong': CAN_CU_MP, 'cach_tra_luong': cach, 'cach_tra_luong_ten': CACH[cach], 'quy_chia_theo_dong_gop': QUY,
       'nguoi': nguoi, 'tong': tong, 'doi_chieu': doi_chieu, 'theo_sprint': theo_sprint,
       'huy_dong_bo_sung': [{'thanh_vien': p, 'sprint': s, 'md': m} for (p, s), m in sorted(HD_BS_SP.items())],
       'don_gia': [{'vai_tro': k, 'ten': v['ten'], 'luong_thang': v['luong'], 'don_gia_ngay': round(dgn[k])} for k, v in vt],
       'kiem_tra': {'ket_luan': 'dat', 'phat_hien': F, 'doi_chieu_excel': ghi}}
loi.ghi_json('du_lieu/luong.json', out)
print(f'Đến {loi.ngay(den)}, {CACH[cach]}: {len(nguoi)} người, {loi.md(tong["md_thuc_te"])} MD thực tế trên {loi.md(tong["md_huy_dong"])} MD huy động, '
      f'tổng lương {loi.trieu(tong["luong"])} triệu, chi phí doanh nghiệp {loi.trieu(tong["chi_phi"])} triệu ({loi.pt(doi_chieu["ty_le_chi_phi_tren_A"])} khoản A). {ghi}')
pbs = loi.hoan_tat_buoc('bang_luong', f'Lập bảng phân chia lương, {CACH[cach]}. ' + ghi, ghi_chu=a.ghi_chu,
                        hash_tep={'ket_qua/14_Bang_phan_chia_luong/Bang_phan_chia_luong.xlsx': loi.bam_obj(out)})
for k, v in pbs.items():
    print(f'  {k}: phiên bản {v}')
