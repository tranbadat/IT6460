"""Lập cost estimate từ WBS, đơn giá, giao việc, lịch và sổ rủi ro.

Dùng: python tinh_chi_phi.py
Đầu vào: du_lieu/wbs.json, dau_vao/don_gia.json, dau_vao/thanh_vien.json, du_lieu/phan_cong.json, du_lieu/lich.json,
         du_lieu/rui_ro.json, du_lieu/hsmt.json
Phương pháp:
  A. Nhân công theo loại công việc = MD / md_moi_mm x đơn giá trọn gói theo người-tháng.
  B. Chi phí khác theo danh mục trong don_gia.json.
  D. Dự phòng rủi ro theo du_phong_rui_ro.cach_tinh (nằm trong B, tỷ lệ, hoặc bằng tổng EMV), so với tổng EMV của sổ rủi ro.
  Giá chào: mục tiêu (C là phần còn lại) hoặc cộng dồn (C = tỷ lệ x chi phí). Thuế GTGT theo thuế suất và căn cứ khai báo.
  Đối chiếu: chi phí lương kế hoạch theo từng người = lương tham chiếu x hệ số bảo hiểm / ngày làm việc mỗi tháng x MD tính lương.
  MD tính lương theo don_gia.json cach_tra_luong (nhóm quyết định, thiếu thì dừng C31): cong_suc là MD công việc giao theo lich.json
  phan_bo (tổng bằng tổng WBS), huy_dong và ket_hop là MD huy động (phan_cong.json tai[].quy). Lương của toàn bộ ngày huy động và tỷ lệ
  sử dụng quỹ vẫn tính để theo dõi năng lực. Cùng cơ sở với báo cáo tiến độ, tổng kết, bảng lương (qlda-bao-cao-tien-do/scripts/thuc_te.py).
  Khoản còn lại sau chi phí lương kế hoạch = giá trước thuế - chi phí lương theo người - B - D, kèm MD tính lương hòa vốn.
  Dự phòng hiện có khi dự phòng nằm trong B: du_phong_rui_ro.so_tien trống hoặc null là cả khoản B (du_phong_ca_khoan_b true),
  0 là không có dự phòng, số từ 0 đến B là phần dự phòng, ngoài khoảng đó là lỗi C33. Cách tính khác thì là D.
  Quy tắc dùng dự phòng lấy từ .claude/skills/qlda-change/references/phan_loai.json khóa du_phong (nguồn chung, tệp tĩnh).
  Phân bổ: giá theo nhóm WBS và theo hạng mục Mẫu 01A (Mẫu 11A), dòng tiền theo sprint với các đợt thanh toán
  (chi lương theo cách trả lương, thu theo ngày sự kiện cộng thời hạn thanh toán của HSMT). MD của kỳ sau cut over trong lich.json phan_bo
  (kỳ KT của giai đoạn Kết thúc) vào cột Sau cut over, nên tổng chi lương theo MD công việc bằng chi phí lương cả dự án; giá chào không đổi.
Phụ lục hợp đồng (CR đã phê duyệt, yeu_cau_thay_doi.py ap_dung cộng giá điều chỉnh vào gia_chao_muc_tieu và ghi dieu_chinh_gia_theo_phu_luc):
  giá gói thầu là trần của giá dự thầu ban đầu (giá chào trừ giá điều chỉnh, C20), không áp cho giá sau phụ lục. Giá chào so với giá hợp đồng
  sau phụ lục cuối (du_lieu/thay_doi.json, khóa phu_luc của CR): vượt là lỗi C20, thấp hơn là cảnh báo C42. Vết điều chỉnh không khớp
  phụ lục của CR là lỗi C42. chi_phi.json ghi gia_chao_ban_dau, gia_dieu_chinh_phu_luc, gia_hop_dong_sau_phu_luc, phu_luc_hop_dong.
Ra: du_lieu/chi_phi.json, ket_qua/06_Cost_estimate/Cost_estimate.xlsx (công thức). Mã thoát 0 đạt, 1 lỗi.
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'qlda-loi', 'scripts'))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'qlda-bao-cao-tien-do', 'scripts'))
import loi  # noqa: E402
import thuc_te  # noqa: E402
import xlsx_chung as X  # noqa: E402

wbs = loi.doc_json('du_lieu/wbs.json'); dg = loi.doc_json('dau_vao/don_gia.json')
tvs = loi.doc_json('dau_vao/thanh_vien.json')['thanh_vien']; pc = loi.doc_json('du_lieu/phan_cong.json')
lich = loi.doc_json('du_lieu/lich.json'); rr = loi.doc_json('du_lieu/rui_ro.json'); hsmt = loi.doc_json('du_lieu/hsmt.json')
by, la = loi.nut_wbs(wbs)
F = []
add = lambda ma, mo, md='loi': F.append({'ma': ma, 'muc_do': md, 'nhom': 'chi_phi', 'vi_tri': '', 'mo_ta': mo, 'cach_sua': 'can_nguoi_dung', 'de_xuat': ''})
LCV = dg['loai_cong_viec']; MDMM = float(dg['md_moi_mm'])
for k, v in LCV.items():
    if not v.get('don_gia_tron_goi_thang'):
        add('C01', f'Thiếu đơn giá trọn gói cho loại {k} trong don_gia.json')
for g in la:
    if g.get('loai') not in LCV:
        add('C02', f'Gói {g["ma"]} có loại {g.get("loai")} chưa có đơn giá')
for t in tvs:
    if t.get('vai_tro') not in dg['luong_tham_chieu_thang']:
        add('C03', f'{t["ho_ten"]}: vai trò {t.get("vai_tro")} chưa có lương tham chiếu')
CACH = thuc_te.cach_tra_luong(dg)
if not CACH:
    add('C31', thuc_te.thong_bao_thieu_cach())
THEO_HD = thuc_te.theo_huy_dong(CACH)
if [f for f in F if f['muc_do'] == 'loi']:
    for f in F:
        print(f'[{f["muc_do"]}] {f["ma"]}: {f["mo_ta"]}')
    sys.exit(1)

# A. nhân công
nc = []
for k, v in LCV.items():
    m = sum(float(g['md']) for g in la if g['loai'] == k)
    if m:
        nc.append({'loai': k, 'ten': v['ten'], 'md': round(m, 2), 'mm': round(m / MDMM, 4), 'don_gia_thang': v['don_gia_tron_goi_thang'],
                   'thanh_tien': round(m / MDMM * v['don_gia_tron_goi_thang'])})
A = sum(x['thanh_tien'] for x in nc)
cost_goi = {g['ma']: float(g['md']) / MDMM * LCV[g['loai']]['don_gia_tron_goi_thang'] for g in la}
# B, D
B = sum(float(x['so_tien']) for x in dg.get('chi_phi_khac', []))
# EMV còn lại: rủi ro đã xảy ra hoặc đã đóng (theo dau_vao/thuc_te/theo_doi_rui_ro.csv) không cần dự phòng nữa, tác động thực tế ghi ở sổ
emv = float(rr['tong_hop'].get('tong_emv_tien_con_lai', rr['tong_hop']['tong_emv_tien']))
emv_ban_dau = float(rr['tong_hop']['tong_emv_tien'])
cdp = dg.get('du_phong_rui_ro', {}).get('cach_tinh', 'nam_trong_chi_phi_khac')
D = round(emv) if cdp == 'emv' else (round((A + B) * float(dg['du_phong_rui_ro'].get('ty_le', 0))) if cdp == 'ty_le' else 0)
so_tien_dp = dg.get('du_phong_rui_ro', {}).get('so_tien')  # trống hoặc null: cả khoản B, 0: không có dự phòng, từ 0 đến B: phần dự phòng
dp_b, dp_tach, dp_loi = thuc_te.du_phong_khoan_b(dg.get('du_phong_rui_ro', {}), B)
if cdp == 'nam_trong_chi_phi_khac' and dp_loi:
    add('C33', dp_loi)
du_phong_co = D if cdp != 'nam_trong_chi_phi_khac' else dp_b
du_phong_ca_khoan_b = cdp == 'nam_trong_chi_phi_khac' and not dp_tach  # dự phòng bằng toàn bộ khoản B
if du_phong_co < emv:
    add('C10', f'Dự phòng ({loi.tien(du_phong_co)} đồng) thấp hơn EMV còn lại của sổ rủi ro ({loi.tien(emv)} đồng)', 'canh_bao')
# giá


def phu_luc_da_ap_dung():
    """Phụ lục hợp đồng đã cộng giá điều chỉnh vào gia_chao_muc_tieu (don_gia.json dieu_chinh_gia_theo_phu_luc, do
    qlda-change/scripts/yeu_cau_thay_doi.py ap_dung ghi), lấy giá theo khóa phu_luc của CR trong du_lieu/thay_doi.json, theo thứ tự ngày.
    Vết điều chỉnh không có CR hoặc phụ lục tương ứng thì báo lỗi C42."""
    vet = dg.get('dieu_chinh_gia_theo_phu_luc') or []
    if not vet:
        return []
    cr = {c['ma']: c for c in loi.doc_json('du_lieu/thay_doi.json', {'yeu_cau': []}).get('yeu_cau', [])}
    out = []
    for v in vet:
        pl = (cr.get(v.get('ma_cr')) or {}).get('phu_luc') or {}
        if not pl or 'gia_hop_dong_sau' not in pl:
            add('C42', f'don_gia.json ghi điều chỉnh giá theo {v.get("ma_cr")} nhưng du_lieu/thay_doi.json chưa có phụ lục của CR này')
            continue
        if abs(float(pl.get('gia_dieu_chinh', 0)) - float(v.get('gia_dieu_chinh', 0))) > 1:
            add('C42', f'Giá điều chỉnh của {v["ma_cr"]} trong don_gia.json ({loi.tien(v.get("gia_dieu_chinh", 0))} đồng) khác phụ lục '
                       f'{pl.get("so", "")} ({loi.tien(pl.get("gia_dieu_chinh", 0))} đồng)')
        out.append(dict(pl, ma_cr=v['ma_cr']))
    return sorted(out, key=lambda x: (x.get('ngay', ''), x['ma_cr']))


vat = float(dg['thue_gtgt']['thue_suat'])
chi_phi_biet = A + B + D
if dg.get('cach_tinh_gia') == 'muc_tieu':
    gia = float(dg['gia_chao_muc_tieu'])
    truoc_thue = gia / (1 + vat)
    C = truoc_thue - chi_phi_biet
else:
    C = chi_phi_biet * float(dg['ty_le_chi_phi_chung_va_loi_nhuan'])
    truoc_thue = chi_phi_biet + C
    gia = truoc_thue * (1 + vat)
thue = gia - truoc_thue
ggt = float(hsmt['tien']['gia_goi_thau'])
PHU_LUC = phu_luc_da_ap_dung()
gia_dc = sum(float(x['gia_dieu_chinh']) for x in PHU_LUC)  # giá điều chỉnh theo phụ lục đã cộng vào gia_chao_muc_tieu
gia_ban_dau = gia - gia_dc  # giá dự thầu ban đầu (giá hợp đồng khi ký), giá gói thầu là trần của giá này
if gia_ban_dau > ggt:
    add('C20', f'Giá chào {loi.tien(gia_ban_dau)} đồng vượt giá gói thầu {loi.tien(ggt)} đồng'
        + (f' (giá dự thầu ban đầu, trước {len(PHU_LUC)} phụ lục hợp đồng)' if PHU_LUC else ''))
gia_hd_sau = float(PHU_LUC[-1]['gia_hop_dong_sau']) if PHU_LUC else None
if PHU_LUC and gia > gia_hd_sau + 1:
    add('C20', f'Giá chào {loi.tien(gia)} đồng vượt giá hợp đồng sau phụ lục {PHU_LUC[-1]["so"]} ({loi.tien(gia_hd_sau)} đồng)')
elif PHU_LUC and abs(gia - gia_hd_sau) > 1:
    add('C42', f'Giá chào {loi.tien(gia)} đồng khác giá hợp đồng sau phụ lục {PHU_LUC[-1]["so"]} ({loi.tien(gia_hd_sau)} đồng), '
        'kiểm tra dieu_chinh_gia_theo_phu_luc của don_gia.json và phu_luc của du_lieu/thay_doi.json', 'canh_bao')
if C / truoc_thue < float(dg.get('ty_le_chi_phi_chung_va_loi_nhuan_toi_thieu', 0)):
    add('C21', f'Khoản chi phí chung và lợi nhuận chỉ {C / truoc_thue:.1%} giá trước thuế, dưới mức tối thiểu', 'loi' if C < 0 else 'canh_bao')
vat2 = float(dg['thue_gtgt'].get('thue_suat_do_nhay', vat))
tt2 = gia / (1 + vat2); C2 = tt2 - chi_phi_biet

# đối chiếu theo người
ngay_thang = float(dg['ngay_lam_viec_moi_thang']); hs = float(dg['he_so_bao_hiem'])
quy = {x['id']: x['quy'] for x in pc['tai']}
giao = thuc_te.md_cong_viec(lich, tong_wbs=wbs['tong_md'])  # MD công việc giao theo lịch, tổng bằng tổng WBS
dcn = []
for t in tvs:
    luong = float(dg['luong_tham_chieu_thang'][t['vai_tro']]['luong'])
    dgn = luong * hs / ngay_thang
    q, m = float(quy.get(t['id'], 0)), giao.get(t['id'], 0)
    md_tl = q if THEO_HD else m
    dcn.append({'id': t['id'], 'ho_ten': t['ho_ten'], 'loai': 'Thành viên nhóm' if t['loai'] == 'nhom' else 'Vị trí bổ sung', 'vai_tro': t['vai_tro'],
                'luong': luong, 'don_gia_ngay': round(dgn), 'md_huy_dong': q, 'md_giao': m, 'md_tinh_luong': md_tl, 'md_san_sang': round(q - m, 2),
                'chi_phi': round(dgn * md_tl), 'chi_phi_huy_dong': round(dgn * q), 'chi_phi_cong_viec': round(dgn * m), '_dgn': dgn})
# chi phí từng người làm tròn sao cho tổng bằng tổng chưa làm tròn đã làm tròn (cùng số với lũy kế chi của dòng tiền và khoản còn lại sau lương)
for x, v in zip(dcn, loi.lam_tron_tong([x['_dgn'] * x['md_tinh_luong'] for x in dcn], 0)):
    x['chi_phi'] = int(v)
tong_nguoi = sum(x['chi_phi'] for x in dcn)
CO_SO = thuc_te.CO_SO_LUONG.get(CACH, '')
if tong_nguoi > A:
    lo = tong_nguoi + B + D > truoc_thue
    add('C30', f'Chi phí lương kế hoạch theo người, {thuc_te.CACH_TRA_LUONG.get(CACH, "")}, ({loi.tien(tong_nguoi)} đồng) vượt dòng A ({loi.tien(A)} đồng)'
               + ('. Giá trước thuế không đủ bù chi phí lương, cần giảm MD tính lương hoặc xem lại giá' if lo else '. Giá chào vẫn bù được nhưng khoản C giảm'),
        'loi' if lo else 'canh_bao')

tong_md_hd = sum(x['md_huy_dong'] for x in dcn); tong_md_tl = sum(x['md_tinh_luong'] for x in dcn)
luong_tho = sum(x['_dgn'] * x['md_tinh_luong'] for x in dcn)  # chưa làm tròn từng người, cùng cách tính với Excel
luong_hd_tho = sum(x['_dgn'] * x['md_huy_dong'] for x in dcn)
luong_md = luong_tho / tong_md_tl if tong_md_tl else 0  # chi phí lương bình quân mỗi MD tính lương
luong_md_hd = luong_hd_tho / tong_md_hd if tong_md_hd else 0  # bình quân mỗi MD huy động, theo dõi năng lực
C_luong = truoc_thue - luong_tho - B - D
C2_luong = tt2 - luong_tho - B - D
md_hoa_von = A / luong_md if luong_md else 0
md_hd_hoa_von = A / luong_md_hd if luong_md_hd else 0
if 0 <= C_luong < truoc_thue * float(dg.get('ty_le_chi_phi_chung_va_loi_nhuan_toi_thieu', 0)):
    add('C22', f'Khoản còn lại sau chi phí lương kế hoạch chỉ {C_luong / truoc_thue:.1%} giá trước thuế, dưới mức tối thiểu', 'canh_bao')
ghi_chu_hd = (f'Theo cách trả lương đã khai báo, ngày sẵn sàng không ghi công ({loi.md(tong_md_hd - tong_md_tl)} MD) không tính lương. '
              f'Lương của toàn bộ ngày huy động ({loi.tien(sum(x["chi_phi_huy_dong"] for x in dcn))} đồng) chỉ dùng để theo dõi năng lực.' if not THEO_HD else
              f'Theo cách trả lương đã khai báo, lương tính cho toàn bộ ngày huy động. Dòng A đủ chi trả khi tổng MD huy động không vượt '
              f'{loi.md(md_hd_hoa_von)} MD, tức tỷ lệ sử dụng quỹ từ {loi.pt(wbs["tong_md"] / md_hd_hoa_von if md_hd_hoa_von else 0)} trở lên.')
PL_TD = loi.doc_json(os.path.join(loi.THU_MUC_SKILLS, 'qlda-change', 'references', 'phan_loai.json'), {})
quy_tac_dp = {'cot': ['Nội dung', 'Quy định'], 'widths': [2300, 7000],
              'dong': [[x['noi_dung'], x['quy_dinh'].format_map(PL_TD.get('vai_ban', {}))] for x in PL_TD.get('du_phong', {}).get('quy_tac', [])]}
if not quy_tac_dp['dong']:
    add('C32', 'qlda-change/references/phan_loai.json chưa có du_phong.quy_tac, thuyết minh dự toán không có bảng quy tắc dùng dự phòng')

# phân bổ theo WBS và Mẫu 01A
he_so_gia = truoc_thue / A if A else 0
a_nhom = [sum(cost_goi[g['ma']] for g in la if g['ma'].split('.')[0] == n['ma']) for n in wbs['nhom']]
a_tron = loi.lam_tron_tong(a_nhom, 0); g_tron = loi.lam_tron_tong([v / A * gia if A else 0 for v in a_nhom], 0)  # tổng các dòng bằng đúng A và giá chào
theo_wbs = [{'ma': n['ma'], 'ten': n['ten'], 'md': n['md'], 'chi_phi_a': int(a_tron[i]), 'gia_phan_bo': int(g_tron[i])} for i, n in enumerate(wbs['nhom'])]
p01 = {}
for g in la:
    hm = g.get('hang_muc_01a') or []
    for h in hm:
        p01[h] = p01.get(h, 0) + cost_goi[g['ma']] / len(hm)
m11 = []
a_hm = [p01.get(h['stt'], 0) for h in hsmt['mau_01a']]
a_hm_tron = loi.lam_tron_tong(a_hm, 0)
for i, h in enumerate(hsmt['mau_01a']):
    m11.append({'stt': h['stt'], 'danh_muc': h['danh_muc'], 'khoi_luong': h['khoi_luong'], 'don_vi': h['don_vi'], 'chi_phi_a': int(a_hm_tron[i]),
                'thanh_tien': round(a_hm[i] / A * gia) if A else 0})
m11[-1]['thanh_tien'] = round(gia) - sum(x['thanh_tien'] for x in m11[:-1])  # cùng cách với sheet Mau_11A: dòng cuối nhận phần làm tròn

# quan hệ với Mẫu 11A của HSDT: trước ký hợp đồng bảng phân bổ là số liệu Mẫu 11A; sau đường cơ sở hợp đồng HSDT đã khóa,
# Mẫu 11A giữ giá dự thầu ban đầu, giá hợp đồng điều chỉnh qua phụ lục và bảng phân bổ là phân bổ giá hợp đồng hiện hành
KY_HD = loi.da_ky_hop_dong()
hm_cr = {}
for g in la:
    for c in [n for n in g.get('nguon', []) if str(n).startswith('CR')]:
        hm_cr.setdefault(c, set()).update(g.get('hang_muc_01a') or [])
if KY_HD:
    doan_11a = (f'Hồ sơ dự thầu đã khóa từ đường cơ sở hợp đồng, Mẫu 11A của hồ sơ dự thầu giữ giá dự thầu ban đầu {loi.tien(gia_ban_dau)} đồng. '
                'Bảng phân bổ ở mục 6 là phân bổ giá hợp đồng hiện hành theo hạng mục Mẫu 01A, dùng để theo dõi khối lượng và thanh toán, '
                'không thay thế Mẫu 11A. Giá hợp đồng chỉ điều chỉnh qua phụ lục hợp đồng.')
    if PHU_LUC:
        doan_11a += ' ' + ' '.join(f'Phần việc theo {x["ma_cr"]} ({x.get("so", "")}) được gán vào hạng mục '
                                   + ', '.join(str(h) for h in sorted(hm_cr.get(x['ma_cr'], set()))) + ' của Mẫu 01A theo khai báo của gói công việc.'
                                   for x in PHU_LUC if hm_cr.get(x['ma_cr']))
else:
    doan_11a = 'Bảng phân bổ ở mục 6 là số liệu điền vào Mẫu 11A của hồ sơ dự thầu, tổng phân bổ bằng giá chào.'
NHAN_GIA = 'Giá hợp đồng sau phụ lục' if PHU_LUC else ('Giá hợp đồng' if KY_HD else 'Giá chào trọn gói')

# dòng tiền theo sprint
dgn_id = {x['id']: x.pop('_dgn') for x in dcn}  # đơn giá ngày chưa làm tròn, khớp công thức Excel
sp = [s['ma'] for s in lich['sprint']]
SAU = 'Sau cut over'


def ky_dong_tien(ky):
    """Cột dòng tiền của một kỳ phan_bo: sprint giữ mã, kỳ sau cut over (KT của giai đoạn Kết thúc) vào cột Sau cut over."""
    return ky if ky in sp else SAU


chi_sp = {}
for r in lich['phan_bo']:
    k = ky_dong_tien(r['sprint']); chi_sp[k] = chi_sp.get(k, 0) + r['md'] * dgn_id[r['thanh_vien']]
quy_sp = {t['id']: loi.quy_thoi_gian(t, lich['sprint']) for t in tvs}
chi_hd_sp = {s: sum(quy_sp[p][s] * dgn_id[p] for p in quy_sp) for s in sp}  # lương trả theo thời gian huy động
tre_tt = int(hsmt['hop_dong'].get('thanh_toan_trong_ngay') or 0)  # thời hạn thanh toán (ngày) theo E-ĐKC 13.1


def sprint_cua(ngay_iso):
    return next((x['ma'] for x in lich['sprint'] if x['den'] >= ngay_iso), SAU)


def ngay_thu_cua(ngay_sk):
    return (dt.date.fromisoformat(ngay_sk) + dt.timedelta(days=tre_tt)).isoformat() if ngay_sk else ''


def phu_luc_tru_khoi_dot_goc(ngay_thu):
    """Phụ lục có giá điều chỉnh không tính vào cơ sở của một đợt thanh toán gốc: phụ lục khai báo đợt thanh toán riêng
    (dot_thanh_toan), hoặc phụ lục có ngày sau ngày thu của đợt (đợt đã thu trước phụ lục không tính lại)."""
    return [x for x in PHU_LUC if x.get('dot_thanh_toan') or (ngay_thu and ngay_thu < x.get('ngay', ''))]


moc_dot = {int(m['dot_thanh_toan']): m for m in lich['moc'] if str(m.get('dot_thanh_toan', '')).strip()}
moc_ma = {m['ma']: m for m in lich['moc']}
TEN_GIA_GOC = 'Giá hợp đồng ban đầu' if PHU_LUC else 'Giá hợp đồng'
tt_rows, chenh_chua_khai = [], {}
for x in hsmt['hop_dong']['thanh_toan']:
    m = moc_dot.get(x['dot'])
    if x['dot'] == 1:
        ngay_sk = lich['ngay_bat_dau']
    elif m:
        ngay_sk = m['ngay']
    else:
        ngay_sk = ''
        if x['dot'] < len(hsmt['hop_dong']['thanh_toan']):
            add('C40', f'Đợt thanh toán {x["dot"]} chưa gắn mốc trong lịch (thêm dot_thanh_toan vào mốc)', 'canh_bao')
    ngay_thu = ngay_thu_cua(ngay_sk)
    tru = phu_luc_tru_khoi_dot_goc(ngay_thu)
    co_so = gia - sum(float(p['gia_dieu_chinh']) for p in tru)
    for p in tru:
        if not p.get('dot_thanh_toan'):  # đợt đã thu trước phụ lục chưa khai báo cách thanh toán phần chênh
            chenh_chua_khai[p['ma_cr']] = chenh_chua_khai.get(p['ma_cr'], 0) + float(x['phan_tram'])
    tt_rows.append({'dot': x['dot'], 'phan_tram': x['phan_tram'], 'su_kien': x['su_kien'], 'ngay_su_kien': ngay_sk,
                    'ngay': ngay_thu, 'sprint': sprint_cua(ngay_thu) if ngay_thu else SAU, 'so_tien': round(co_so * x['phan_tram'] / 100),
                    'co_so': round(co_so), 'co_so_ten': TEN_GIA_GOC if len(tru) == len(PHU_LUC) else 'Giá hợp đồng sau '
                    + ', '.join(p.get('so', '') for p in PHU_LUC if p not in tru), 'ma_cr': '', '_tru': [p['ma_cr'] for p in tru]})
for p in PHU_LUC:
    dots = p.get('dot_thanh_toan') or []
    if dots and abs(sum(float(d.get('phan_tram_gia_dieu_chinh') or 0) for d in dots) - 100) > 1e-6:
        add('C43', f'Đợt thanh toán của phụ lục {p.get("so", "")} ({p["ma_cr"]}) cộng '
                   f'{loi.so(sum(float(d.get("phan_tram_gia_dieu_chinh") or 0) for d in dots), 2)}% giá điều chỉnh, phải bằng 100%')
    for d in dots:
        m = moc_ma.get(d.get('moc', ''))
        if not m:
            add('C40', f'Đợt thanh toán {d.get("dot")} của phụ lục {p.get("so", "")} gắn mốc {d.get("moc")} không có trong lịch', 'canh_bao')
        ngay_sk = m['ngay'] if m else ''
        ngay_thu = ngay_thu_cua(ngay_sk)
        tt_rows.append({'dot': d.get('dot') or p['ma_cr'], 'phan_tram': float(d['phan_tram_gia_dieu_chinh']), 'su_kien': d.get('su_kien', ''),
                        'ngay_su_kien': ngay_sk, 'ngay': ngay_thu, 'sprint': sprint_cua(ngay_thu) if ngay_thu else SAU,
                        'so_tien': round(float(p['gia_dieu_chinh']) * float(d['phan_tram_gia_dieu_chinh']) / 100),
                        'co_so': round(float(p['gia_dieu_chinh'])), 'co_so_ten': f'Giá điều chỉnh {p["ma_cr"]} theo {p.get("so", "")}',
                        'ma_cr': p['ma_cr'], 'moc': d.get('moc', ''), 'nguon': d.get('nguon', ''), '_tru': []})
    if p['ma_cr'] in chenh_chua_khai:
        pt = chenh_chua_khai[p['ma_cr']]
        add('C43', f'Phụ lục {p.get("so", "")} ({p["ma_cr"]}) chưa khai báo đợt thanh toán (phu_luc.dot_thanh_toan), phần giá điều chỉnh tương ứng '
                   f'{loi.so(pt, 0)}% của các đợt đã thu trước phụ lục tạm tính thu sau cut over. Ghi đợt bằng yeu_cau_thay_doi.py phu_luc', 'canh_bao')
        tt_rows.append({'dot': p['ma_cr'], 'phan_tram': pt, 'su_kien': f'Phần giá điều chỉnh theo {p.get("so", "")} tương ứng các đợt đã thu trước phụ lục',
                        'ngay_su_kien': '', 'ngay': '', 'sprint': SAU, 'so_tien': round(float(p['gia_dieu_chinh']) * pt / 100),
                        'co_so': round(float(p['gia_dieu_chinh'])), 'co_so_ten': f'Giá điều chỉnh {p["ma_cr"]} theo {p.get("so", "")}',
                        'ma_cr': p['ma_cr'], '_tru': []})
if abs(sum(r['so_tien'] for r in tt_rows) - gia) > len(tt_rows):
    add('C43', f'Tổng các đợt thanh toán {loi.tien(sum(r["so_tien"] for r in tt_rows))} đồng khác giá hợp đồng {loi.tien(gia)} đồng')
TRU_DOT = [r.pop('_tru') for r in tt_rows]
sp_dt = sp + ([SAU] if any(r['sprint'] == SAU for r in tt_rows) or chi_sp.get(SAU) else [])
dong_tien, lk_c, lk_t = [], 0, 0
for s in sp_dt:
    thu = sum(r['so_tien'] for r in tt_rows if r['sprint'] == s)
    lk_c += chi_hd_sp.get(s, 0) if THEO_HD else chi_sp.get(s, 0); lk_t += thu
    dong_tien.append({'sprint': s, 'chi_luong': round(chi_sp.get(s, 0)), 'chi_luong_huy_dong': round(chi_hd_sp.get(s, 0)), 'thu': thu,
                      'luy_ke_chi': round(lk_c), 'luy_ke_thu': lk_t, 'chenh_lech': round(lk_t - lk_c)})
am = [x['sprint'] for x in dong_tien if x['chenh_lech'] < 0]
if am:
    add('C41', f'Dòng tiền lũy kế âm ở các sprint {", ".join(am)}, cần vốn lưu động', 'canh_bao')

out = {'nhan_cong': nc, 'A': round(A), 'chi_phi_khac': dg.get('chi_phi_khac', []), 'B': round(B), 'D': D, 'cach_tinh_du_phong': cdp,
       'du_phong_hien_co': round(du_phong_co), 'du_phong_ca_khoan_b': du_phong_ca_khoan_b, 'tong_emv_rui_ro': round(emv), 'tong_emv_rui_ro_ban_dau': round(emv_ban_dau),
       'so_rui_ro_da_xay_ra': rr['tong_hop'].get('so_da_xay_ra', 0), 'so_rui_ro_da_dong': rr['tong_hop'].get('so_da_dong', 0),
       'tac_dong_thuc_te_tien': rr['tong_hop'].get('tac_dong_thuc_te_tien', 0), 'tac_dong_thuc_te_md': rr['tong_hop'].get('tac_dong_thuc_te_md', 0),
       'rui_ro_da_xay_ra': [{k: r.get(k) for k in ('ma', 'ten', 'trang_thai', 'ngay_xay_ra', 'ngay_dong', 'tac_dong_thuc_te_md', 'tac_dong_thuc_te_tien', 'emv_tien')}
                            for r in rr['danh_sach'] if r.get('trang_thai') in ('Đã xảy ra', 'Đã đóng') or r.get('ngay_xay_ra')], 'chi_phi_da_biet': round(chi_phi_biet),
       'gia_chao': round(gia), 'gia_truoc_thue': round(truoc_thue), 'thue': round(thue), 'thue_suat': vat, 'can_cu_thue': dg['thue_gtgt']['can_cu'],
       'C': round(C), 'ty_le_C': round(C / truoc_thue, 4), 'gia_goi_thau': round(ggt), 'thap_hon_goi_thau': round(ggt - gia_ban_dau), 'ty_le_thap_hon': round((ggt - gia_ban_dau) / ggt, 4),
       'gia_chao_ban_dau': round(gia_ban_dau), 'gia_dieu_chinh_phu_luc': round(gia_dc),
       'gia_hop_dong_sau_phu_luc': round(gia_hd_sau) if PHU_LUC else None,
       'doan_phu_luc': ([f'Giá chào trên là giá hợp đồng sau {len(PHU_LUC)} phụ lục hợp đồng ('
                         + ', '.join(f'{x.get("so", "")} theo {x["ma_cr"]}' for x in PHU_LUC)
                         + '), gồm giá dự thầu ban đầu {cp#gia_chao_ban_dau|tien} đồng và giá điều chỉnh {cp#gia_dieu_chinh_phu_luc|tien} đồng. '
                         'Giá gói thầu là trần của giá dự thầu ban đầu, không áp cho giá sau phụ lục.'] if PHU_LUC else []),
       'phu_luc_hop_dong': [{k: x.get(k) for k in ('ma_cr', 'so', 'ngay', 'gia_dieu_chinh', 'gia_hop_dong_truoc', 'gia_hop_dong_sau')} for x in PHU_LUC],
       'do_nhay_thue': {'thue_suat': vat2, 'gia_truoc_thue': round(tt2), 'C': round(C2), 'ty_le_C': round(C2 / tt2, 4),
                        'C_sau_luong': round(C2_luong), 'ty_le_C_sau_luong': round(C2_luong / tt2, 4)},
       'C_sau_luong': round(truoc_thue) - tong_nguoi - round(B) - round(D), 'ty_le_C_sau_luong': round(C_luong / truoc_thue, 4),
       'cach_tra_luong': CACH, 'cach_tra_luong_ten': thuc_te.CACH_TRA_LUONG.get(CACH, ''), 'co_so_luong': CO_SO,
       'luong_binh_quan_md': round(luong_md), 'luong_binh_quan_md_huy_dong': round(luong_md_hd),
       'nguong_chiu_dung_md': round(C_luong / luong_md, 1) if luong_md else 0, 'md_hoa_von': round(md_hoa_von, 2),
       'md_huy_dong_hoa_von': round(md_hd_hoa_von, 2), 'ty_le_su_dung': round(wbs['tong_md'] / tong_md_hd, 4) if tong_md_hd else 0,
       'ty_le_su_dung_hoa_von': round(wbs['tong_md'] / md_hd_hoa_von, 4) if md_hd_hoa_von else 0, 'thoi_han_thanh_toan_ngay': tre_tt,
       'tong_md': wbs['tong_md'], 'tong_mm': round(wbs['tong_md'] / MDMM, 2), 'md_moi_mm': MDMM,
       'doi_chieu_nguoi': dcn, 'tong_chi_phi_nguoi': tong_nguoi, 'tong_md_huy_dong': round(sum(x['md_huy_dong'] for x in dcn), 2),
       'tong_md_giao': round(sum(x['md_giao'] for x in dcn), 2), 'tong_md_tinh_luong': round(tong_md_tl, 2), 'tong_md_san_sang': round(tong_md_hd - sum(x['md_giao'] for x in dcn), 2),
       'tong_chi_phi_huy_dong': sum(x['chi_phi_huy_dong'] for x in dcn), 'tong_chi_phi_cong_viec': sum(x['chi_phi_cong_viec'] for x in dcn),
       'ghi_chu_huy_dong': ghi_chu_hd, 'quy_tac_du_phong': quy_tac_dp,
       'chenh_lech_A_nguoi': round(A - tong_nguoi), 'ty_le_chenh_lech_A': round((A - tong_nguoi) / A, 4), 'nang_suat_hoa_von': round(tong_nguoi / A, 4),
       'he_so_bao_hiem': hs, 'ngay_lam_viec_moi_thang': ngay_thang, 'nguon_luong': dg['nguon_luong'], 'nguon_don_gia': dg['nguon_don_gia_tron_goi'],
       'doan_mau_11a': doan_11a, 'nhan_gia_chao': NHAN_GIA, 'theo_wbs': theo_wbs, 'mau_11a': m11, 'thanh_toan': tt_rows, 'dong_tien': dong_tien,
       'kiem_tra': {'phat_hien': F}}
loi_ = [f for f in F if f['muc_do'] == 'loi']
for f in F:
    print(f'[{f["muc_do"]}] {f["ma"]}: {f["mo_ta"]}')
if loi_:
    sys.exit(1)
loi.ghi_json('du_lieu/chi_phi.json', out)

# ---------------- Excel
OUT = loi.duong_dan('ket_qua/06_Cost_estimate/Cost_estimate.xlsx')
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb = Workbook()
wb.properties.creator = wb.properties.lastModifiedBy = loi.du_an().get('nguoi_lap', '')
ts = wb.active; ts.title = 'Tham_so'
ts['A1'] = 'Tham số dự toán'; ts['A1'].font = X.F_TEN
tham = [('md_moi_mm', 'Số MD mỗi người-tháng', MDMM, '0'), ('he_so', 'Hệ số bảo hiểm phần doanh nghiệp', hs, '0.000'),
        ('ngay_thang', 'Ngày làm việc mỗi tháng', ngay_thang, '0'), ('gia_muc_tieu', 'Giá chào mục tiêu, đồng', float(dg.get('gia_chao_muc_tieu') or 0), X.SO),
        ('thue', 'Thuế suất GTGT áp dụng', vat, X.PT), ('thue2', 'Thuế suất GTGT tính độ nhạy', vat2, X.PT),
        ('tre_tt', 'Thời hạn thanh toán theo HSMT (Mục 13.1 E-ĐKCT), ngày', tre_tt, '0'),
        ('ty_le_min', 'Tỷ lệ chi phí chung và lợi nhuận tối thiểu', float(dg.get('ty_le_chi_phi_chung_va_loi_nhuan_toi_thieu', 0)), X.PT),
        ('ty_le_cl', 'Tỷ lệ chi phí chung và lợi nhuận (cách cộng dồn)', float(dg.get('ty_le_chi_phi_chung_va_loi_nhuan', 0)), X.PT),
        ('gia_goi_thau', 'Giá gói thầu theo HSMT, đồng', ggt, X.SO),
        ('gia_dc', 'Giá điều chỉnh theo phụ lục hợp đồng đã cộng vào giá chào mục tiêu, đồng', gia_dc, X.SO),
        *[(f'gia_dc_{x["ma_cr"]}', f'Giá điều chỉnh theo {x.get("so", "")} ({x["ma_cr"]}), đồng', float(x['gia_dieu_chinh']), X.SO) for x in PHU_LUC],
        ('emv', 'EMV còn lại của sổ rủi ro (rủi ro chưa xảy ra, chưa đóng), đồng', emv, X.SO),
        ('dp_ty_le', 'Tỷ lệ dự phòng (cách tỷ lệ)', float(dg.get('du_phong_rui_ro', {}).get('ty_le', 0)), X.PT),
        ('dp_so_tien', 'Số tiền dự phòng rủi ro trong khoản B (để trống là cả khoản B, 0 là không có dự phòng)',
         None if so_tien_dp in (None, '') else float(so_tien_dp), X.SO),
        ('cach_luong', 'Cách trả lương (cong_suc: MD công việc, huy_dong hoặc ket_hop: MD huy động)', CACH, None)]
R = {}
X.tieu_de(ts, 3, ['Tham số', 'Giá trị', 'Nguồn'], [46, 18, 70])
for i, (k, ten, v, f) in enumerate(tham, 4):
    X.o(ts, i, 1, ten); X.o(ts, i, 2, v, fmt=f, nhap=True)
    X.o(ts, i, 3, {'thue': dg['thue_gtgt']['can_cu'], 'he_so': dg.get('nguon_he_so_bao_hiem', ''), 'gia_goi_thau': 'du_lieu/hsmt.json',
                   'gia_dc': 'du_lieu/thay_doi.json phu_luc (' + ', '.join(f'{x["ma_cr"]} {x.get("so", "")}' for x in PHU_LUC) + ')' if PHU_LUC else 'Chưa có phụ lục hợp đồng điều chỉnh giá',
                   **{f'gia_dc_{x["ma_cr"]}': f'du_lieu/thay_doi.json, {x["ma_cr"]} phu_luc' for x in PHU_LUC},
                   'emv': 'du_lieu/rui_ro.json tong_hop.tong_emv_tien_con_lai'}.get(k, 'dau_vao/don_gia.json'), wrap=True)
    R[k] = f'Tham_so!$B${i}'
r0 = len(tham) + 6
X.o(ts, r0 - 1, 1, 'Đơn giá trọn gói theo loại công việc', dam=True)
X.tieu_de(ts, r0, ['Mã loại', 'Tên', 'Đơn giá, đồng/người-tháng'])
for i, (k, v) in enumerate(LCV.items(), r0 + 1):
    X.o(ts, i, 1, k, nhap=True); X.o(ts, i, 2, v['ten'], nhap=True); X.o(ts, i, 3, v['don_gia_tron_goi_thang'], fmt=X.SO, nhap=True)
dg_rng = (f'Tham_so!$A${r0 + 1}:$A${r0 + len(LCV)}', f'Tham_so!$C${r0 + 1}:$C${r0 + len(LCV)}')
r1 = r0 + len(LCV) + 3
X.o(ts, r1 - 1, 1, 'Lương tham chiếu theo vai trò (' + dg['nguon_luong'] + ')', dam=True)
X.tieu_de(ts, r1, ['Mã vai trò', 'Tên', 'Lương, đồng/tháng'])
for i, (k, v) in enumerate(dg['luong_tham_chieu_thang'].items(), r1 + 1):
    X.o(ts, i, 1, k, nhap=True); X.o(ts, i, 2, v['ten'], nhap=True); X.o(ts, i, 3, v['luong'], fmt=X.SO, nhap=True)
lg_rng = (f'Tham_so!$A${r1 + 1}:$A${r1 + len(dg["luong_tham_chieu_thang"])}', f'Tham_so!$C${r1 + 1}:$C${r1 + len(dg["luong_tham_chieu_thang"])}')
ts.freeze_panes = None

wm = wb.create_sheet('WBS_MD')
X.tieu_de(wm, 1, ['Mã gói', 'Tên gói', 'Loại', 'MD', 'Nhóm mức 1', 'Chi phí A của gói, đồng'], [9, 52, 8, 8, 10, 18])
for i, g in enumerate(la, 2):
    X.o(wm, i, 1, g['ma'], nhap=True); X.o(wm, i, 2, g['ten'], nhap=True); X.o(wm, i, 3, g['loai'], nhap=True)
    X.o(wm, i, 4, float(g['md']), fmt=X.SO2, nhap=True); X.o(wm, i, 5, g['ma'].split('.')[0], nhap=True)
    X.o(wm, i, 6, f'=D{i}/{R["md_moi_mm"]}*INDEX({dg_rng[1]},MATCH(C{i},{dg_rng[0]},0))', fmt=X.SO)
nla = len(la) + 1

nh = wb.create_sheet('Nhan_cong')
X.tieu_de(nh, 1, ['Loại công việc', 'Tên', 'MD', 'Người-tháng', 'Đơn giá, đồng/người-tháng', 'Thành tiền, đồng'], [12, 34, 10, 12, 20, 18])
for i, x in enumerate(nc, 2):
    X.o(nh, i, 1, x['loai'], nhap=True); X.o(nh, i, 2, x['ten'], nhap=True)
    X.o(nh, i, 3, f'=SUMIFS(WBS_MD!$D:$D,WBS_MD!$C:$C,A{i})', fmt=X.SO2)
    X.o(nh, i, 4, f'=C{i}/{R["md_moi_mm"]}', fmt=X.SO2)
    X.o(nh, i, 5, f'=INDEX({dg_rng[1]},MATCH(A{i},{dg_rng[0]},0))', fmt=X.SO, lien_ket=True)
    X.o(nh, i, 6, f'=D{i}*E{i}', fmt=X.SO)
nn = len(nc) + 1
X.o(nh, nn + 1, 2, 'A. Tổng nhân công', dam=True)
for c in (3, 4, 6):
    X.o(nh, nn + 1, c, f'=SUM({L(c)}2:{L(c)}{nn})', fmt=X.SO2 if c < 6 else X.SO, dam=True)
A_REF = f'Nhan_cong!$F${nn + 1}'

kh = wb.create_sheet('Chi_phi_khac')
X.tieu_de(kh, 1, ['Mã', 'Khoản mục', 'Số tiền, đồng', 'Căn cứ'], [8, 50, 18, 50])
for i, x in enumerate(dg.get('chi_phi_khac', []), 2):
    X.o(kh, i, 1, x['ma'], nhap=True); X.o(kh, i, 2, x['ten'], nhap=True); X.o(kh, i, 3, float(x['so_tien']), fmt=X.SO, nhap=True)
    X.o(kh, i, 4, x.get('can_cu', ''), nhap=True, wrap=True)
nk = len(dg.get('chi_phi_khac', [])) + 1
X.o(kh, nk + 1, 2, 'B. Tổng chi phí khác', dam=True); X.o(kh, nk + 1, 3, f'=SUM(C2:C{nk})', fmt=X.SO, dam=True)
B_REF = f'Chi_phi_khac!$C${nk + 1}'

th = wb.create_sheet('Tong_hop', 0)
th['A1'] = 'Cost estimate'; th['A1'].font = X.F_TEN
X.o(th, 2, 1, loi.du_an()['ten_du_an'])
D_F = {'emv': f'={R["emv"]}', 'ty_le': f'=({A_REF}+{B_REF})*{R["dp_ty_le"]}'}.get(cdp, '=0')
rows = [('A. Nhân công theo loại công việc', f'={A_REF}', X.SO), ('B. Chi phí khác', f'={B_REF}', X.SO),
        ('D. Dự phòng rủi ro tách riêng', D_F, X.SO), ('Chi phí đã biết (A + B + D)', '=B4+B5+B6', X.SO),
        ('Giá chào trọn gói, gồm thuế', f'={R["gia_muc_tieu"]}' if dg.get('cach_tinh_gia') == 'muc_tieu' else f'=B7*(1+{R["ty_le_cl"]})*(1+{R["thue"]})', X.SO),
        ('Giá trước thuế', f'=B8/(1+{R["thue"]})', X.SO), ('Thuế GTGT', '=B8-B9', X.SO),
        ('C. Chi phí chung và lợi nhuận', '=B9-B7', X.SO), ('Tỷ lệ C trên giá trước thuế', '=IF(B9=0,0,B11/B9)', X.PT),
        ('Giá gói thầu', f'={R["gia_goi_thau"]}', X.SO), ('Giá dự thầu ban đầu thấp hơn giá gói thầu (giá chào trừ điều chỉnh theo phụ lục)', f'=B13-(B8-{R["gia_dc"]})', X.SO),
        ('Tỷ lệ thấp hơn', '=IF(B13=0,0,B14/B13)', X.PT),
        ('Kiểm tra giá dự thầu ban đầu không vượt giá gói thầu', f'=IF(B8-{R["gia_dc"]}<=B13,"Đạt","Vượt giá gói thầu")', None),
        ('Kiểm tra tỷ lệ C tối thiểu', f'=IF(B12>={R["ty_le_min"]},"Đạt","Dưới mức tối thiểu")', None),
        ('Tổng EMV sổ rủi ro', f'={R["emv"]}', X.SO),
        ('Dự phòng hiện có (phần dự phòng trong B, trống thì cả B, hoặc D)',
         f'=IF(ISBLANK({R["dp_so_tien"]}),B5,{R["dp_so_tien"]})' if cdp == 'nam_trong_chi_phi_khac' else '=B6', X.SO),
        ('Kiểm tra dự phòng so với EMV', '=IF(B19>=B18,"Đủ","Thiếu")', None),
        ('Độ nhạy: giá trước thuế nếu áp thuế suất tính độ nhạy', f'=B8/(1+{R["thue2"]})', X.SO), ('Độ nhạy: khoản C tương ứng', '=B21-B7', X.SO),
        ('Chi phí lương kế hoạch theo người (theo cách trả lương)', f'=Doi_chieu_nguoi!$J${len(dcn) + 2}', X.SO),
        ('Khoản còn lại sau chi phí lương kế hoạch (giá trước thuế - lương - B - D)', '=B9-B23-B5-B6', X.SO),
        ('Tỷ lệ khoản còn lại trên giá trước thuế', '=IF(B9=0,0,B24/B9)', X.PT),
        ('Kiểm tra giá trước thuế đủ bù chi phí lương kế hoạch', '=IF(B24>=0,"Đạt","Không đủ")', None),
        ('Kiểm tra tỷ lệ khoản còn lại tối thiểu', f'=IF(B25>={R["ty_le_min"]},"Đạt","Dưới mức tối thiểu")', None),
        ('Độ nhạy: khoản còn lại sau chi phí lương kế hoạch', '=B21-B23-B5-B6', X.SO)]
X.tieu_de(th, 3, ['Khoản mục', 'Giá trị, đồng'], [62, 22])
for i, (k, f, nf) in enumerate(rows, 4):
    X.o(th, i, 1, k, dam=k[:2] in ('A.', 'B.', 'C.', 'D.') or k.startswith('Giá chào trọn')); X.o(th, i, 2, f, fmt=nf, lien_ket=f.startswith('=Tham') or 'Nhan_cong' in f or 'Chi_phi' in f)
th.conditional_formatting.add('B16:B17', FormulaRule(formula=['AND(B16<>"Đạt",B16<>"")'], fill=PatternFill('solid', fgColor='F8CBAD')))
th.conditional_formatting.add('B20', FormulaRule(formula=['B20="Thiếu"'], fill=PatternFill('solid', fgColor='F8CBAD')))
th.conditional_formatting.add('B26:B27', FormulaRule(formula=['AND(B26<>"Đạt",B26<>"")'], fill=PatternFill('solid', fgColor='F8CBAD')))
X.ghi_chu_bang(th, 30, 'Căn cứ thuế: ' + dg['thue_gtgt']['can_cu'])
X.ghi_chu_bang(th, 31, 'Dự phòng hiện có lấy số tiền dự phòng rủi ro khai báo trong khoản B. Để trống thì cả khoản B được coi là dự phòng, 0 là không có dự phòng.')
X.ghi_chu_bang(th, 32, 'Chữ xanh dương: dữ liệu đầu vào. Chữ xanh lá: tham chiếu sang sheet khác. Chữ đen: công thức.')

dn = wb.create_sheet('Doi_chieu_nguoi')
X.tieu_de(dn, 1, ['Mã', 'Họ tên', 'Nhóm', 'Vai trò lương', 'Lương tham chiếu, đồng/tháng', 'Đơn giá ngày gồm bảo hiểm', 'MD huy động', 'MD công việc giao',
                  'MD tính lương', 'Chi phí lương kế hoạch, đồng', 'Lương của ngày huy động, đồng'], [8, 26, 16, 12, 18, 16, 12, 12, 12, 18, 18])
for i, x in enumerate(dcn, 2):
    X.o(dn, i, 1, x['id'], nhap=True); X.o(dn, i, 2, x['ho_ten'], nhap=True); X.o(dn, i, 3, x['loai'], nhap=True); X.o(dn, i, 4, x['vai_tro'], nhap=True)
    X.o(dn, i, 5, f'=INDEX({lg_rng[1]},MATCH(D{i},{lg_rng[0]},0))', fmt=X.SO, lien_ket=True)
    X.o(dn, i, 6, f'=E{i}*{R["he_so"]}/{R["ngay_thang"]}', fmt=X.SO)
    X.o(dn, i, 7, x['md_huy_dong'], fmt=X.SO2, nhap=True); X.o(dn, i, 8, x['md_giao'], fmt=X.SO2, nhap=True)
    X.o(dn, i, 9, f'=IF({R["cach_luong"]}="cong_suc",H{i},G{i})', fmt=X.SO2)
    X.o(dn, i, 10, f'=F{i}*I{i}', fmt=X.SO); X.o(dn, i, 11, f'=F{i}*G{i}', fmt=X.SO)
nd = len(dcn) + 1
X.o(dn, nd + 1, 2, 'Tổng', dam=True)
for c in (7, 8, 9, 10, 11):
    X.o(dn, nd + 1, c, f'=SUM({L(c)}2:{L(c)}{nd})', fmt=X.SO2 if c < 10 else X.SO, dam=True)
X.o(dn, nd + 2, 2, 'Chênh lệch dòng A so với chi phí lương kế hoạch theo người'); X.o(dn, nd + 2, 10, f'={A_REF}-J{nd + 1}', fmt=X.SO)
X.o(dn, nd + 3, 2, 'Năng suất tối thiểu để dòng A đủ chi trả'); X.o(dn, nd + 3, 10, f'=IF({A_REF}=0,0,J{nd + 1}/{A_REF})', fmt=X.PT)
X.o(dn, nd + 4, 2, 'MD tính lương hòa vốn (dòng A chia lương bình quân mỗi MD tính lương)'); X.o(dn, nd + 4, 10, f'=IF(J{nd + 1}=0,0,{A_REF}/(J{nd + 1}/I{nd + 1}))', fmt=X.SO2)
X.o(dn, nd + 5, 2, 'Năng lực: tỷ lệ sử dụng quỹ huy động (MD WBS chia MD huy động)'); X.o(dn, nd + 5, 10, f'=IF(G{nd + 1}=0,0,Nhan_cong!$C${nn + 1}/G{nd + 1})', fmt=X.PT)
X.o(dn, nd + 6, 2, 'Năng lực: MD huy động hòa vốn (dòng A chia lương bình quân mỗi MD huy động)'); X.o(dn, nd + 6, 10, f'=IF(K{nd + 1}=0,0,{A_REF}/(K{nd + 1}/G{nd + 1}))', fmt=X.SO2)
X.o(dn, nd + 7, 2, 'Năng lực: tỷ lệ sử dụng hòa vốn khi trả theo huy động'); X.o(dn, nd + 7, 10, f'=IF(J{nd + 6}=0,0,Nhan_cong!$C${nn + 1}/J{nd + 6})', fmt=X.PT)
X.ghi_chu_bang(dn, nd + 9, 'MD tính lương theo cách trả lương ở sheet Tham_so. Lương của ngày huy động dùng để theo dõi năng lực, chỉ là chi phí khi trả theo huy động.')

g1 = wb.create_sheet('Goi_01A')
X.tieu_de(g1, 1, ['Mã gói', 'Hạng mục Mẫu 01A', 'Tỷ lệ chia', 'Chi phí A phân bổ, đồng'], [10, 14, 10, 20])
rr_ = 2
for i, g in enumerate(la, 2):
    hm = g.get('hang_muc_01a') or []
    for h in hm:
        X.o(g1, rr_, 1, g['ma'], nhap=True); X.o(g1, rr_, 2, h, nhap=True); X.o(g1, rr_, 3, 1 / len(hm), fmt='0.0000', nhap=True)
        X.o(g1, rr_, 4, f'=INDEX(WBS_MD!$F$2:$F${nla},MATCH(A{rr_},WBS_MD!$A$2:$A${nla},0))*C{rr_}', fmt=X.SO)
        rr_ += 1
m1 = wb.create_sheet('Mau_11A')
m1['A1'] = ('Phân bổ giá hợp đồng hiện hành theo hạng mục Mẫu 01A (Mẫu 11A của HSDT đã khóa, giữ giá dự thầu ban đầu)' if KY_HD
            else 'Phân bổ giá chào theo hạng mục Mẫu 01A (căn cứ lập Mẫu 11A)'); m1['A1'].font = X.F_TEN
X.tieu_de(m1, 3, ['STT', 'Danh mục dịch vụ', 'Khối lượng', 'Đơn vị', 'Chi phí A phân bổ, đồng', 'Thành tiền, đồng'], [6, 52, 10, 8, 20, 20])
for i, x in enumerate(m11, 4):
    X.o(m1, i, 1, x['stt'], nhap=True); X.o(m1, i, 2, x['danh_muc'], nhap=True, wrap=True); X.o(m1, i, 3, x['khoi_luong'], nhap=True); X.o(m1, i, 4, x['don_vi'], nhap=True)
    X.o(m1, i, 5, f'=SUMIFS(Goi_01A!$D:$D,Goi_01A!$B:$B,A{i})', fmt=X.SO)
    X.o(m1, i, 6, f'=ROUND(E{i}/{A_REF}*Tong_hop!$B$8,0)' if i < 3 + len(m11) else f'=Tong_hop!$B$8-SUM(F4:F{i - 1})', fmt=X.SO)
n11 = 3 + len(m11)
X.o(m1, n11 + 1, 2, 'Tổng cộng', dam=True)
X.o(m1, n11 + 1, 5, f'=SUM(E4:E{n11})', fmt=X.SO, dam=True); X.o(m1, n11 + 1, 6, f'=SUM(F4:F{n11})', fmt=X.SO, dam=True)
X.ghi_chu_bang(m1, n11 + 3, 'Thành tiền phân bổ theo tỷ trọng chi phí nhân công của hạng mục. Dòng cuối nhận phần làm tròn để tổng bằng đúng giá chào.')

dt_ = wb.create_sheet('Dong_tien')
X.tieu_de(dt_, 1, ['Đợt', 'Tỷ lệ', 'Sự kiện', 'Sprint thu', 'Ngày thu dự kiến', 'Số tiền, đồng', 'Ngày sự kiện', 'Giá làm cơ sở, đồng', 'Cơ sở'],
          [6, 8, 60, 12, 14, 18, 14, 18, 34])
for i, (x, tru) in enumerate(zip(tt_rows, TRU_DOT), 2):
    X.o(dt_, i, 1, x['dot'], nhap=True); X.o(dt_, i, 2, x['phan_tram'] / 100, fmt=X.PT, nhap=True); X.o(dt_, i, 3, x['su_kien'], nhap=True, wrap=True)
    X.o(dt_, i, 4, x['sprint'], nhap=True)
    X.o(dt_, i, 5, f'=IF(G{i}="","",G{i}+{R["tre_tt"]})' if x['ngay_su_kien'] else None, fmt=X.NGAY)
    X.o(dt_, i, 6, f'=B{i}*H{i}', fmt=X.SO)
    X.o(dt_, i, 7, dt.date.fromisoformat(x['ngay_su_kien']) if x['ngay_su_kien'] else None, fmt=X.NGAY, nhap=True)
    # đợt gốc: giá hợp đồng trừ giá điều chỉnh của phụ lục có đợt riêng hoặc ký sau ngày thu; đợt của phụ lục: giá điều chỉnh của phụ lục
    X.o(dt_, i, 8, f'={R["gia_dc_" + x["ma_cr"]]}' if x['ma_cr'] else '=Tong_hop!$B$8' + ''.join(f'-{R["gia_dc_" + m]}' for m in tru), fmt=X.SO)
    X.o(dt_, i, 9, x['co_so_ten'], nhap=True, wrap=True)
r2 = len(tt_rows) + 4
X.tieu_de(dt_, r2, ['Sprint', 'Chi lương theo MD công việc', 'Chi lương theo huy động', 'Thu theo đợt', 'Lũy kế chi (theo cách trả lương)', 'Lũy kế thu', 'Chênh lệch lũy kế'])
X.ghi_chu_bang(dt_, r2 - 1, f'Ngày thu dự kiến = ngày sự kiện + {tre_tt} ngày (thời hạn thanh toán theo HSMT). Lũy kế chi tính theo cách trả lương ở sheet Tham_so. '
               'Đợt gốc tính trên giá hợp đồng ban đầu khi phụ lục có đợt thanh toán riêng, đợt đã thu trước ngày phụ lục không tính lại.')
hd_sheet = wb.create_sheet('Huy_dong_luong')
X.tieu_de(hd_sheet, 1, ['Mã thành viên', 'Sprint', 'MD huy động', 'Đơn giá ngày', 'Chi lương, đồng'], [12, 8, 12, 14, 16])
hd_rows = [(p, s, quy_sp[p][s]) for p in quy_sp for s in sp if quy_sp[p][s]]
for i, (p, s, q) in enumerate(hd_rows, 2):
    X.o(hd_sheet, i, 1, p, nhap=True); X.o(hd_sheet, i, 2, s, nhap=True); X.o(hd_sheet, i, 3, q, fmt=X.SO2, nhap=True)
    X.o(hd_sheet, i, 4, f'=INDEX(Doi_chieu_nguoi!$F$2:$F${nd},MATCH(A{i},Doi_chieu_nguoi!$A$2:$A${nd},0))', fmt=X.SO, lien_ket=True)
    X.o(hd_sheet, i, 5, f'=C{i}*D{i}', fmt=X.SO)
pb_sheet = wb.create_sheet('Phan_bo_luong')
X.tieu_de(pb_sheet, 1, ['Mã thành viên', 'Sprint', 'MD kế hoạch', 'Đơn giá ngày', 'Chi lương, đồng', 'Kỳ dòng tiền'], [12, 8, 12, 14, 16, 14])
for i, x in enumerate(lich['phan_bo'], 2):
    X.o(pb_sheet, i, 1, x['thanh_vien'], nhap=True); X.o(pb_sheet, i, 2, x['sprint'], nhap=True); X.o(pb_sheet, i, 3, x['md'], fmt=X.SO2, nhap=True)
    X.o(pb_sheet, i, 4, f'=INDEX(Doi_chieu_nguoi!$F$2:$F${nd},MATCH(A{i},Doi_chieu_nguoi!$A$2:$A${nd},0))', fmt=X.SO, lien_ket=True)
    X.o(pb_sheet, i, 5, f'=C{i}*D{i}', fmt=X.SO)
    X.o(pb_sheet, i, 6, ky_dong_tien(x['sprint']), nhap=True)
X.ghi_chu_bang(pb_sheet, len(lich['phan_bo']) + 3, 'Kỳ dòng tiền: sprint giữ mã, kỳ KT (giai đoạn Kết thúc sau cut over) tính vào cột Sau cut over của sheet Dong_tien.')
for j, s in enumerate(sp_dt, r2 + 1):
    X.o(dt_, j, 1, s, nhap=True)
    X.o(dt_, j, 2, f'=SUMIFS(Phan_bo_luong!$E:$E,Phan_bo_luong!$F:$F,A{j})', fmt=X.SO)
    X.o(dt_, j, 3, f'=SUMIFS(Huy_dong_luong!$E:$E,Huy_dong_luong!$B:$B,A{j})', fmt=X.SO)
    X.o(dt_, j, 4, f'=SUMIFS($F$2:$F${len(tt_rows) + 1},$D$2:$D${len(tt_rows) + 1},A{j})', fmt=X.SO)
    X.o(dt_, j, 5, f'=IF({R["cach_luong"]}="cong_suc",B{j},C{j})' + (f'+E{j - 1}' if j > r2 + 1 else ''), fmt=X.SO)
    X.o(dt_, j, 6, f'=D{j}' + (f'+F{j - 1}' if j > r2 + 1 else ''), fmt=X.SO)
    X.o(dt_, j, 7, f'=F{j}-E{j}', fmt=X.SO)
dt_.conditional_formatting.add(f'G{r2 + 1}:G{r2 + len(sp_dt)}', FormulaRule(formula=[f'G{r2 + 1}<0'], fill=PatternFill('solid', fgColor='F8CBAD')))
wb.save(OUT)

kq = X.tinh_lai(OUT)
ghi = 'Máy không có LibreOffice, Excel tự tính khi mở.'
if kq and kq.get('status'):
    lc = X.dem_loi_cong_thuc(OUT)
    if lc:
        loi.thoat_loi('Công thức lỗi: ' + ', '.join(lc[:10]))
    v = X.doc_gia_tri(OUT)
    so = [('A', v['Tong_hop']['B4'].value, A), ('B', v['Tong_hop']['B5'].value, B), ('giá chào', v['Tong_hop']['B8'].value, gia),
          ('C', v['Tong_hop']['B11'].value, C), ('lương theo người', v['Doi_chieu_nguoi'][f'J{nd + 1}'].value, tong_nguoi),
          ('dự phòng hiện có', v['Tong_hop']['B19'].value, du_phong_co),
          ('tổng Mẫu 11A', v['Mau_11A'][f'F{n11 + 1}'].value, round(gia))]
    so += [(f'Mẫu 11A dòng {x["stt"]}', v['Mau_11A'][f'F{4 + i}'].value, x['thanh_tien']) for i, x in enumerate(m11)]
    so += [('khoản còn lại sau chi phí lương', v['Tong_hop']['B24'].value, C_luong),
           ('giá dự thầu ban đầu thấp hơn giá gói thầu', v['Tong_hop']['B14'].value, ggt - gia_ban_dau)]
    so += [(f'dòng tiền {x["sprint"]}', v['Dong_tien'][f'G{r2 + 1 + i}'].value, x['chenh_lech']) for i, x in enumerate(dong_tien)]
    for ten, a_, b_ in so:
        if a_ is None or abs(float(a_) - float(b_)) > (0.5 if ten.startswith('Mẫu 11A') else 2):
            loi.thoat_loi(f'Đối chiếu Excel và Python lệch ở {ten}: {a_} khác {b_}')
    ghi = 'Đã tính lại công thức, 0 lỗi, giá trị Excel khớp tính toán Python (A, B, C, giá chào, lương theo người, tổng Mẫu 11A, khoản còn lại sau lương, dòng tiền).'
print(f'A {loi.trieu(A)} triệu, B {loi.trieu(B)} triệu, giá chào {loi.trieu(gia)} triệu, C {loi.trieu(C)} triệu ({C / truoc_thue:.2%}). '
      f'Lương kế hoạch theo người ({thuc_te.CACH_TRA_LUONG[CACH]}) {loi.trieu(tong_nguoi)} triệu, khoản còn lại sau lương {loi.trieu(C_luong)} triệu. '
      f'Quỹ huy động {loi.md(tong_md_hd)} MD, tỷ lệ sử dụng {loi.pt(wbs["tong_md"] / tong_md_hd if tong_md_hd else 0)}. {ghi}')
pbs = loi.hoan_tat_buoc('chi_phi', 'Lập cost estimate và Excel công thức. ' + ghi, bo_qua=('doc_chi_phi',),
                        hash_tep={'ket_qua/06_Cost_estimate/Cost_estimate.xlsx': loi.bam_obj(out)})
for k, v in pbs.items():
    print(f'  {k}: phiên bản {v}')
print('Tiếp theo: python .claude/skills/qlda-loi/scripts/chay.py chuan_bi chi_phi (thuyết minh dự toán)')
