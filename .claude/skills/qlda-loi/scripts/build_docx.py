"""Sinh file Word từ danh sách khối nội dung (dùng chung cho mọi tài liệu của bộ skill).

Dùng: python build_docx.py <noi_dung.json> <file_ra.docx> [--co-dinh mau5a_fixed.json]
noi_dung.json: {"meta": {...}, "blocks": [...]}
Khối hỗ trợ:
  {"t":"center","text":..., "size":13}         dòng căn giữa (bìa, tiêu đề phụ)
  {"t":"h1"|"h2"|"h3","text":...}              tiêu đề các cấp (không in đậm)
  {"t":"p","text":..., "dam":false}            đoạn văn
  {"t":"blank"} | {"t":"pagebreak"}
  {"t":"section","huong":"ngang"|"doc"}        sang trang mới với hướng giấy khác
  {"t":"table","rows":[[...]],"widths":[...],"center_cols":[],"right_cols":[],"dong_dam":[chỉ số dòng]}
  {"t":"fixed_table","key":...,"widths":[...]} | {"t":"fixed_paras","key":...}   phần nguyên văn của mẫu
Quy tắc trình bày: Times New Roman, cỡ mặc định 13, không in đậm (trừ dong_dam, dam), không viết hoa toàn bộ,
gạch dài đổi thành gạch ngang đơn, bảng kẻ lưới đơn, dòng tiêu đề bảng lặp lại khi sang trang.
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import argparse


def lam_sach_thuoc_tinh(path):
    """Bỏ ảnh thu nhỏ và thông tin ứng dụng của mẫu mặc định trong tệp Word."""
    import re as _re, shutil as _sh, zipfile as _zf
    tmp = path + '.tmp'
    with _zf.ZipFile(path) as zi, _zf.ZipFile(tmp, 'w', _zf.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            if it.filename == 'docProps/thumbnail.jpeg':
                continue
            data = zi.read(it.filename)
            if it.filename == '_rels/.rels':
                data = _re.sub(rb'<Relationship [^>]*thumbnail\.jpeg"[^>]*/>', b'', data)
            elif it.filename == 'docProps/app.xml':
                data = _re.sub(rb'<Application>[^<]*</Application>', b'<Application></Application>', data)
                data = _re.sub(rb'<Template>[^<]*</Template>', b'<Template></Template>', data)
            zo.writestr(it, data)
    _sh.move(tmp, path)

import datetime
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import loi  # noqa: E402

ap = argparse.ArgumentParser()
ap.add_argument('noi_dung'); ap.add_argument('ra'); ap.add_argument('--co-dinh')
a = ap.parse_args()
data = json.load(open(a.noi_dung, encoding='utf-8'))
fixed = json.load(open(a.co_dinh, encoding='utf-8')) if a.co_dinh else {}
meta = data.get('meta', {})
FONT = meta.get('font', 'Times New Roman')
SIZE = meta.get('size_pt', 13)
A4 = (11907, 16840)
MARGIN = meta.get('margin_twip', {'top': 1134, 'bottom': 1134, 'left': 1417, 'right': 1134})


def clean(s):
    return re.sub('[–—]', '-', str(s))


def set_font(run, size=SIZE, italic=False, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.all_caps = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(k), FONT)


def fmt(par, align=None, before=0, after=0, line=None, keep_next=False):
    pf = par.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True


doc = Document()
for sname, sz in (('Normal', SIZE), ('Heading 1', SIZE + 2), ('Heading 2', SIZE), ('Heading 3', SIZE)):
    st = doc.styles[sname]
    st.font.name = FONT; st.font.size = Pt(sz); st.font.bold = False; st.font.italic = False
    st.font.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(k), FONT)
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]


def setup_section(sec, huong):
    w, h = A4
    if huong == 'ngang':
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Twips(h), Twips(w)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Twips(w), Twips(h)
    sec.top_margin, sec.bottom_margin = Twips(MARGIN['top']), Twips(MARGIN['bottom'])
    sec.left_margin, sec.right_margin = Twips(MARGIN['left']), Twips(MARGIN['right'])


def text_width(huong):
    w, h = A4
    return (h if huong == 'ngang' else w) - MARGIN['left'] - MARGIN['right']


sec = doc.sections[0]
cur_huong = meta.get('huong', 'doc')
setup_section(sec, cur_huong)
hp = sec.header.paragraphs[0]; fmt(hp, WD_ALIGN_PARAGRAPH.RIGHT)
set_font(hp.add_run(clean(meta.get('header', ''))), 10, italic=True)
fp = sec.footer.paragraphs[0]; fmt(fp, WD_ALIGN_PARAGRAPH.CENTER)
r = fp.add_run(); set_font(r, 11)
for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if tag:
        el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag)
    else:
        el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
    r._element.append(el)


def para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=SIZE, before=0, after=5, line=1.25, style=None, bold=False, keep_next=False):
    pr = doc.add_paragraph(style=style)
    fmt(pr, align, before, after, line, keep_next)
    set_font(pr.add_run(clean(text)), size, bold=bold)
    return pr


def cell_margins(tbl):
    tblPr = tbl._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 40), ('left', 70), ('bottom', 40), ('right', 70)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)


def table(rows, widths, center_cols=(), right_cols=(), dong_dam=(), size=None):
    size = size or meta.get('size_bang_pt', SIZE)
    tw = text_width(cur_huong)
    total = sum(widths)
    if total > tw:  # co lại cho vừa trang
        widths = [int(w * tw / total) for w in widths]
    ncol = len(widths)
    tbl = doc.add_table(rows=0, cols=ncol)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    cell_margins(tbl)
    for ri, row in enumerate(rows):
        r_ = tbl.add_row()
        if ri == 0:
            repeat_header(r_)
        cells = r_.cells
        for i in range(ncol):
            content = row[i] if i < len(row) else ''
            items = content if isinstance(content, list) else [content]
            c = cells[i]; c.width = Twips(widths[i])
            first = True
            for t in items:
                pr = c.paragraphs[0] if first else c.add_paragraph()
                first = False
                al = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else (WD_ALIGN_PARAGRAPH.RIGHT if i in right_cols else WD_ALIGN_PARAGRAPH.LEFT)
                fmt(pr, al, 0, 2)
                set_font(pr.add_run(clean('' if t is None else t)), size, bold=ri in dong_dam)
    for i, w in enumerate(widths):
        tbl.columns[i].width = Twips(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


PB = meta.get('phien_ban', ''); NG = meta.get('ngay', '')
_txt = json.dumps(data['blocks'], ensure_ascii=False).replace('{{PHIEN_BAN}}', PB).replace('{{NGAY}}', NG)
data['blocks'] = json.loads(_txt)
for b in data['blocks']:
    t = b['t']
    if t == 'center':
        para(b['text'], WD_ALIGN_PARAGRAPH.CENTER, b.get('size', SIZE), 0, 6, None)
    elif t == 'h1':
        para(b['text'], WD_ALIGN_PARAGRAPH.LEFT, SIZE + 2, 12, 6, None, 'Heading 1', keep_next=True)
    elif t == 'h2':
        para(b['text'], WD_ALIGN_PARAGRAPH.LEFT, SIZE, 10, 4, None, 'Heading 2', keep_next=True)
    elif t == 'h3':
        para(b['text'], WD_ALIGN_PARAGRAPH.LEFT, SIZE, 8, 4, None, 'Heading 3', keep_next=True)
    elif t == 'p':
        pr_ = para(b['text'], bold=b.get('dam', False), align=WD_ALIGN_PARAGRAPH.LEFT if b.get('thut') is not None else WD_ALIGN_PARAGRAPH.JUSTIFY,
                   after=2 if b.get('thut') is not None else 5)
        if b.get('thut'):
            pr_.paragraph_format.left_indent = Twips(360 * int(b['thut']))
    elif t == 'blank':
        doc.add_paragraph()
    elif t == 'pagebreak':
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    elif t == 'muc_luc':
        para('Mục lục', WD_ALIGN_PARAGRAPH.CENTER, SIZE + 2, 0, 8)
        # trường TOC, nội dung tạm là danh sách tiêu đề, Word cập nhật số trang khi mở tệp
        for i, txt in enumerate(b.get('muc', []) or ['']):
            pr = doc.add_paragraph(); fmt(pr, WD_ALIGN_PARAGRAPH.LEFT, 0, 2, 1.15, False)
            if i == 0:
                r0 = pr.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r0._r.append(f1)
                r1 = pr.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'TOC \\o "1-2" \\h \\z \\u'; r1._r.append(it)
                r2 = pr.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate'); r2._r.append(f2)
            set_font(pr.add_run(clean(txt)), SIZE)
            if i == len(b.get('muc', []) or ['']) - 1:
                r3 = pr.add_run(); f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end'); r3._r.append(f3)
        st = doc.settings.element
        if st.find(qn('w:updateFields')) is None:
            u = OxmlElement('w:updateFields'); u.set(qn('w:val'), 'true'); st.append(u)
    elif t == 'section':
        cur_huong = b.get('huong', 'doc')
        ns = doc.add_section(WD_SECTION.NEW_PAGE)
        setup_section(ns, cur_huong)
    elif t == 'table':
        table(b['rows'], b['widths'], b.get('center_cols', []), b.get('right_cols', []), b.get('dong_dam', []), b.get('size'))
    elif t == 'fixed_table':
        table([[l, r_] for l, r_ in fixed[b['key']]], b['widths'])
    elif t == 'fixed_paras':
        for txt in fixed[b['key']]:
            para(txt)
    else:
        raise ValueError('Khối không hợp lệ: ' + t)

try:
    tac_gia = meta.get('tac_gia') or loi.du_an().get('nguoi_lap', '')
except Exception:
    tac_gia = meta.get('tac_gia', '')
cp = doc.core_properties
cp.author = tac_gia; cp.last_modified_by = tac_gia
cp.title = meta.get('ten', '')
cp.created = cp.modified = datetime.datetime.now()
cp.revision = 1
cp.comments = ''; cp.subject = ''; cp.keywords = ''; cp.category = ''
os.makedirs(os.path.dirname(os.path.abspath(a.ra)), exist_ok=True)
doc.save(a.ra)
lam_sach_thuoc_tinh(a.ra)
print('Đã tạo', loi.tuong_doi(a.ra))
