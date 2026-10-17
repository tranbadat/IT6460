"""Tiện ích Excel dùng chung: kiểu chữ, ô nhập liệu, công thức, độ rộng cột, tính lại công thức."""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import os, shutil, subprocess, sys, json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = 'Times New Roman'
F_THUONG = Font(name=FONT, size=11)
F_NHAP = Font(name=FONT, size=11, color='0000FF')       # giá trị nhập (đầu vào)
F_LIEN_KET = Font(name=FONT, size=11, color='008000')   # tham chiếu sang sheet khác
F_TIEU_DE = Font(name=FONT, size=11, bold=True)
F_TEN = Font(name=FONT, size=14, bold=True)
NEN_TIEU_DE = PatternFill('solid', fgColor='D9E1F2')
NEN_CAN_DIEN = PatternFill('solid', fgColor='FFFF00')
NEN_NHOM = PatternFill('solid', fgColor='F2F2F2')
VIEN = Border(*(Side(style='thin', color='808080'),) * 4)
SO = '#,##0;(#,##0);-'
SO2 = '#,##0.00;(#,##0.00);-'
PT = '0.0%;(0.0%);-'
NGAY = 'dd/mm/yyyy'


def tieu_de(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F_TIEU_DE; c.fill = NEN_TIEU_DE; c.border = VIEN
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def o(ws, row, col, value, kieu=None, fmt=None, nhap=False, lien_ket=False, dam=False, wrap=False, nen=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = F_NHAP if nhap else (F_LIEN_KET if lien_ket else (F_TIEU_DE if dam else F_THUONG))
    c.border = VIEN
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical='top')
    if nen:
        c.fill = nen
    return c


def ghi_chu_bang(ws, row, text, col=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name=FONT, size=10, italic=True)
    return c


def _tinh_lai_goc(path, timeout=60):
    """Tính lại công thức bằng LibreOffice nếu có. Trả về dict kết quả hoặc None nếu máy không có LibreOffice."""
    helper = '/mnt/skills/public/xlsx/scripts/recalc.py'
    if os.path.exists(helper):
        r = subprocess.run([sys.executable, helper, path, str(timeout)], capture_output=True, text=True, encoding='utf-8', errors='replace')
        try:
            return json.loads(r.stdout)
        except Exception:
            return {'error': r.stdout + r.stderr}
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return None
    outdir = os.path.dirname(os.path.abspath(path))
    tmp = os.path.join(outdir, '_tinh_lai')
    os.makedirs(tmp, exist_ok=True)
    subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, path], capture_output=True, timeout=timeout)
    out = os.path.join(tmp, os.path.basename(path))
    if os.path.exists(out):
        shutil.move(out, path)
    shutil.rmtree(tmp, ignore_errors=True)
    return {'status': 'success'}


def doc_gia_tri(path):
    """Đọc giá trị đã tính (sau tinh_lai)."""
    from openpyxl import load_workbook
    return load_workbook(path, data_only=True)


def dem_loi_cong_thuc(path):
    wb = doc_gia_tri(path)
    loi = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('#') and c.value.rstrip('!?0/').upper() in ('#REF', '#NAME', '#VALUE', '#DIV', '#N/A', '#NUM', '#NULL'):
                    loi.append(f'{ws.title}!{c.coordinate}={c.value}')
    return loi


def lam_sach_thuoc_tinh(path):
    """Bỏ tên ứng dụng tính lại (LibreOffice) trong thuộc tính tệp Excel."""
    import re as _re, zipfile as _zf
    tmp = path + '.tmp'
    with _zf.ZipFile(path) as zi, _zf.ZipFile(tmp, 'w', _zf.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == 'docProps/app.xml':
                data = _re.sub(rb'<Application>[^<]*</Application>', b'<Application></Application>', data)
                data = _re.sub(rb'<AppVersion>[^<]*</AppVersion>', b'', data)
            zo.writestr(it, data)
    shutil.move(tmp, path)


def tinh_lai(path, timeout=60):
    """Tính lại công thức bằng LibreOffice nếu có, sau đó làm sạch thuộc tính tệp. Trả về dict kết quả hoặc None nếu máy không có LibreOffice."""
    kq = _tinh_lai_goc(path, timeout)
    try:
        lam_sach_thuoc_tinh(path)
    except Exception:
        pass
    return kq
