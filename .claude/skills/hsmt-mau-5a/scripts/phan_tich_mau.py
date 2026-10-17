"""Bước 1. Phân tích Mẫu số 5A (TT 79/2025/TT-BTC).

Dùng: python3 phan_tich_mau.py <Mau_5A.docx> <thu_muc_lam_viec>
Ra:
- mau5a_fixed.json: Chương I, Chương VI, Mục 1 Chương III nguyên văn (chép vào HSMT, không sửa).
- mau5a_cau_truc.json: các mục phải điền kèm hướng dẫn của mẫu: E-BDL, E-ĐKCT (Phần 3A), Bảng số 01,
  Bảng số 02, danh mục biểu mẫu Chương IV, cột Mẫu 01A, hướng dẫn Chương III và Chương V, danh mục Phần 4.
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import json, os, re, sys
import docx

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
d = docx.Document(src)
N = lambda s: re.sub(r'\s+', ' ', s).strip()


def cell_pars(c):
    return [t for t in (N(p.text) for p in c.paragraphs) if t]


def first_table(pred, start=0):
    for i, tb in enumerate(d.tables):
        if i < start:
            continue
        try:
            if pred(tb):
                return i
        except Exception:
            pass
    raise SystemExit('Không tìm thấy bảng cần thiết trong mẫu, kiểm tra lại file Mẫu 5A')


def dump2(ti):
    rows, seen = [], set()
    for r in d.tables[ti].rows:
        left = cell_pars(r.cells[0])
        right = cell_pars(r.cells[-1]) if len(r.cells) > 1 else []
        k = ('|'.join(left), '|'.join(right))
        if k in seen:
            continue
        seen.add(k)
        rows.append([left, right])
    return rows


def uniq_rows(ti):
    rows = []
    for r in d.tables[ti].rows:
        vals, last = [], None
        for c in r.cells:
            if c._tc is last:
                continue
            last = c._tc
            vals.append(N(c.text))
        if not rows or vals != rows[-1]:
            rows.append(vals)
    return rows


c1 = first_table(lambda t: t.rows[0].cells[0].text.strip().startswith('1. Phạm vi gói thầu'))
c6 = first_table(lambda t: t.rows[0].cells[0].text.strip().startswith('1. Định nghĩa'))
bdl = first_table(lambda t: t.rows[0].cells[0].text.strip().startswith('E-CDNT'))
dkct = first_table(lambda t: t.rows[0].cells[0].text.strip().startswith('E-ĐKC'))  # bảng đầu tiên là Phần 3A
b01 = first_table(lambda t: 'Bảng số 01' in t.rows[0].cells[0].text)
b02 = first_table(lambda t: N(t.rows[0].cells[1].text) == 'Vị trí công việc')
forms = first_table(lambda t: N(t.rows[0].cells[1].text) == 'Biểu mẫu')
m01a = first_table(lambda t: N(t.rows[0].cells[1].text) == 'Danh mục dịch vụ', start=forms)

paras = [N(p.text) for p in d.paragraphs]


def idx(prefix, start=0):
    return next(i for i, t in enumerate(paras) if i >= start and t.startswith(prefix))


i1 = idx('Mục 1. Đánh giá tính hợp lệ')
i2 = idx('Mục 2. Tiêu chuẩn đánh giá về năng lực')
i4 = idx('Chương IV. BIỂU MẪU MỜI THẦU', i2)
i5 = idx('Chương V. YÊU CẦU VỀ KỸ THUẬT', i4)
i6 = idx('Phần 3', i5)
ip4 = max(i for i, t in enumerate(paras) if t.startswith('Phần 4. BIỂU MẪU HỢP ĐỒNG'))

phan4 = []
for j in range(ip4, len(paras)):
    m = re.match(r'^Mẫu số (1[3-7][AB]?)$', paras[j])
    if m:
        tieu_de = next((t for t in paras[j + 1:j + 6] if t and not t.startswith('(')), '')
        phan4.append({'ma': 'Mẫu số ' + m.group(1), 'tieu_de': tieu_de})

fixed = {'c1': dump2(c1), 'c6': dump2(c6), 'muc1': [t for t in paras[i1 + 1:i2] if t]}
cau_truc = {
    'nguon': os.path.basename(src),
    'bia': [N(r.cells[0].text) for r in d.tables[0].rows],
    'e_bdl': [{'ma': N(r.cells[0].text), 'noi_dung_mau': N(r.cells[-1].text)} for r in d.tables[bdl].rows],
    'e_dkct': [{'ma': N(r.cells[0].text), 'noi_dung_mau': N(r.cells[-1].text)} for r in d.tables[dkct].rows],
    'bang01': uniq_rows(b01),
    'bang02': uniq_rows(b02),
    'chuong4_bieu_mau': uniq_rows(forms),
    'mau01a_cot': uniq_rows(m01a)[0],
    'huong_dan_chuong3': [t for t in paras[i2:i4] if t],
    'huong_dan_chuong5': [t for t in paras[i5 + 1:i6] if t],
    'phan4': phan4,
}
json.dump(fixed, open(os.path.join(out, 'mau5a_fixed.json'), 'w', encoding='utf-8'), ensure_ascii=False)
json.dump(cau_truc, open(os.path.join(out, 'mau5a_cau_truc.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f"Chương I {len(fixed['c1'])} dòng, Chương VI {len(fixed['c6'])} dòng, Mục 1 Chương III {len(fixed['muc1'])} đoạn")
print(f"E-BDL {len(cau_truc['e_bdl'])} mục, E-ĐKCT {len(cau_truc['e_dkct'])} mục, Bảng 01 {len(cau_truc['bang01'])} dòng, "
      f"biểu mẫu Chương IV {len(cau_truc['chuong4_bieu_mau'])} dòng, Phần 4 {len(phan4)} mẫu")
