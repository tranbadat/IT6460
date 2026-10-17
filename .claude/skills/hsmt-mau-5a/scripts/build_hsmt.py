"""Sinh file Word E-HSMT từ nội dung, dùng được trên Claude và ChatGPT (Code Interpreter).

Dùng: python3 build_hsmt.py <hsmt_noi_dung.json> <mau5a_fixed.json> <file_ra.docx> [--vong N]
- hsmt_noi_dung.json: nội dung HSMT sau khi ghép khung và bản soạn (dien_khung.py).
- mau5a_fixed.json: Chương I, Chương VI, Mục 1 Chương III trích nguyên văn từ Mẫu 5A (phan_tich_mau.py).
- --vong N: vòng kiểm tra tạo ra tệp, ghi vào thuộc tính version. Thuộc tính identifier ghi md5 của hsmt_noi_dung.json
  để bước trích dữ liệu đối chiếu tệp Word giao đi với nội dung được trích.
Chỉ cần thư viện python-docx.
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import hashlib, json, os, re, sys


def lam_sach_thuoc_tinh(path):
    """Bỏ ảnh thu nhỏ và thông tin ứng dụng của mẫu mặc định trong tệp Word."""
    import re as _re, shutil as _sh, zipfile as _zf
    tmp = path + '.tmp'
    with _zf.ZipFile(path) as zi, _zf.ZipFile(tmp, 'w', _zf.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            if it.filename == 'docProps/thumbnail.jpeg':
                continue
            data = zi.read(it.filename)
            if it.filename == '_rels/.rels':
                data = _re.sub(rb'<Relationship [^>]*thumbnail\.jpeg"[^>]*/>', b'', data)
            elif it.filename == 'docProps/app.xml':
                data = _re.sub(rb'<Application>[^<]*</Application>', b'<Application></Application>', data)
                data = _re.sub(rb'<Template>[^<]*</Template>', b'<Template></Template>', data)
            zo.writestr(it, data)
    _sh.move(tmp, path)

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, RGBColor

data_path, fixed_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
data = json.load(open(data_path, encoding='utf-8'))
fixed = json.load(open(fixed_path, encoding='utf-8'))
meta = data['meta']
FONT, SIZE = meta['font'], meta['size_pt']

def clean(s):
    return re.sub('[\u2013\u2014]', '-', str(s))

def set_font(run, size=SIZE, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.italic = italic
    run.font.all_caps = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), FONT)

def fmt(par, align=None, before=0, after=0, line=None):
    pf = par.paragraph_format
    if align is not None: pf.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line: pf.line_spacing = line

doc = Document()
# kiểu mặc định và kiểu tiêu đề: Times New Roman, không đậm, màu đen
for sname, sz in (('Normal', SIZE), ('Heading 1', SIZE + 2), ('Heading 2', SIZE)):
    st = doc.styles[sname]
    st.font.name = FONT; st.font.size = Pt(sz); st.font.bold = False; st.font.italic = False
    st.font.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]

sec = doc.sections[0]
sec.page_width, sec.page_height = Twips(11907), Twips(16840)
m = meta['margin_twip']
sec.top_margin, sec.bottom_margin = Twips(m['top']), Twips(m['bottom'])
sec.left_margin, sec.right_margin = Twips(m['left']), Twips(m['right'])

def truong(run, tag=None, lenh=None):
    """Thêm ký hiệu trường (begin, separate, end) hoặc lệnh trường vào run."""
    if tag:
        el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag)
    else:
        el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = lenh
    run._element.append(el)


hp = sec.header.paragraphs[0]; fmt(hp, WD_ALIGN_PARAGRAPH.RIGHT)
set_font(hp.add_run(clean(meta['header'])), 10, italic=True)
fp = sec.footer.paragraphs[0]; fmt(fp, WD_ALIGN_PARAGRAPH.CENTER)
r = fp.add_run(); set_font(r, 11)
truong(r, 'begin'); truong(r, lenh='PAGE'); truong(r, 'end')

def para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=SIZE, before=0, after=5, line=1.25, style=None):
    pr = doc.add_paragraph(style=style)
    fmt(pr, align, before, after, line)
    set_font(pr.add_run(clean(text)), size)
    return pr

def cell_margins(tbl):
    tblPr = tbl._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 50), ('left', 80), ('bottom', 50), ('right', 80)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)

def table(rows, widths, center_cols=()):
    ncol = len(widths)
    tbl = doc.add_table(rows=0, cols=ncol)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    cell_margins(tbl)
    for row in rows:
        cells = tbl.add_row().cells
        for i, content in enumerate(row):
            items = content if isinstance(content, list) else [content]
            c = cells[i]; c.width = Twips(widths[i])
            first = True
            for t in items:
                pr = c.paragraphs[0] if first else c.add_paragraph()
                first = False
                fmt(pr, WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT, 0, 3)
                set_font(pr.add_run(clean(t)))
    for i, w in enumerate(widths):
        tbl.columns[i].width = Twips(w)
    if len(rows) > 1 and str(rows[0][0]) in ('STT', 'TT', 'Nội dung'):  # dòng tiêu đề lặp lại khi bảng sang trang
        tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return tbl

def muc_luc(dong):
    """Trường mục lục (TOC) theo Heading 1, 2, có số trang. Word cập nhật khi mở tệp. Danh sách dòng là nội dung tạm."""
    for i, txt in enumerate(dong):
        pr = doc.add_paragraph(); fmt(pr, WD_ALIGN_PARAGRAPH.LEFT, 0, 3)
        if i == 0:
            r0 = pr.add_run(); truong(r0, 'begin'); truong(r0, lenh=' TOC \\o "1-2" \\h \\z \\u '); truong(r0, 'separate')
        set_font(pr.add_run(clean(txt)))
        if i == len(dong) - 1:
            truong(pr.add_run(), 'end')
    st = doc.settings.element
    if st.find(qn('w:updateFields')) is not None:
        return
    uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true')
    # đúng thứ tự phần tử của w:settings: updateFields đứng trước hdrShapeDefaults, footnotePr, endnotePr, compat
    sau = next((c for c in st if c.tag in {qn('w:' + n) for n in ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids')}), None)
    sau.addprevious(uf) if sau is not None else st.append(uf)


for b in data['blocks']:
    t = b['t']
    if t == 'toc':
        muc_luc(b['muc'])
    elif t == 'center':
        para(b['text'], WD_ALIGN_PARAGRAPH.CENTER, b.get('size', SIZE), 0, 6, None)
    elif t == 'h1':
        para(b['text'], WD_ALIGN_PARAGRAPH.CENTER, SIZE + 2, 12, 8, None, 'Heading 1')
    elif t == 'h2':
        para(b['text'], WD_ALIGN_PARAGRAPH.LEFT, SIZE, 10, 6, None, 'Heading 2')
    elif t == 'p':
        para(b['text'])
    elif t == 'blank':
        doc.add_paragraph()
    elif t == 'pagebreak':
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    elif t == 'table':
        table(b['rows'], b['widths'], b.get('center_cols', []))
    elif t == 'fixed_table':
        table([[l, r] for l, r in fixed[b['key']]], b['widths'])
    elif t == 'fixed_paras':
        for txt in fixed[b['key']]:
            para(txt)
    else:
        raise ValueError('Khối không hợp lệ: ' + t)

try:
    _tg = json.load(open(os.path.join(os.getcwd(), 'dau_vao', 'du_an.json'), encoding='utf-8')).get('nguoi_lap', '')
except Exception:
    _tg = meta.get('tac_gia', '')
doc.core_properties.author = _tg
doc.core_properties.last_modified_by = _tg
doc.core_properties.title = meta['ten']
doc.core_properties.comments = ''
doc.core_properties.version = sys.argv[sys.argv.index('--vong') + 1] if '--vong' in sys.argv else ''
doc.core_properties.identifier = hashlib.md5(open(data_path, 'rb').read()).hexdigest()
doc.save(out_path)
lam_sach_thuoc_tinh(out_path)
print('Đã tạo', out_path)
