"""Kiểm tra độc lập file HSMT (vai Kiểm tra, lớp A).

Dùng: python3 verify_hsmt.py <HSMT.docx> --lam-viec <thu_muc_lam_viec> [--mau Mau_5A.docx] [--json bao_cao.json]
Thư mục làm việc cần có: mau5a_cau_truc.json, mau5a_fixed.json, rfp_yeu_cau.json, gia_tri_tinh.json, tham_so_chu_dau_tu.json.
Chỉ đọc file Word đầu ra và các nguồn chuẩn. Không đọc khung, bản soạn hay nội dung trung gian của bên tạo.
Đọc thêm sổ phát hiện so_phat_hien.json (nếu có, do bên kiểm tra ghi qua chay_quy_trinh.py ghi_nhan) để mang
các phát hiện còn mở của vòng trước sang vòng này.
Mọi giá trị kỳ vọng lấy từ nguồn chuẩn, không ghi cứng trong script.
Kết luận (quy_uoc.py): dat; chua_dat khi còn lỗi; cho_nguoi_dung khi còn tham số chờ Chủ đầu tư quyết định (TS01)
hoặc phát hiện mở của vòng trước chưa giải trình. ket_luan_lop_a chỉ xét lỗi mới của lớp A.
Mã thoát: 0 đạt, 1 còn lỗi, 2 chưa hoàn tất, chờ người dùng.
"""
import os as _os, sys as _sys  # noqa: E401
_os.environ.setdefault('PYTHONUTF8', '1')  # tiến trình con đọc, ghi UTF-8 (Windows)
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass
import argparse, json, os, re, sys, zipfile
import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from quy_uoc import (BEN, DAU_CHO, DOAN_2_2_C, DOAN_2_2_NHAN_SU, HA_TANG, KIEM_THU_VAN_HANH, LIEN_HE_CDT, THAY_DOI_HE_THONG, doc_so, hang_muc_theo_dot, ket_luan,  # noqa: E402
                     ket_luan_lop_a, lay, mang_sang, md5_tep, tham_so_thieu, vong_cua_word)

ap = argparse.ArgumentParser()
ap.add_argument('hsmt'); ap.add_argument('--lam-viec', required=True); ap.add_argument('--mau'); ap.add_argument('--json')
a = ap.parse_args()
J = lambda n: json.load(open(os.path.join(a.lam_viec, n), encoding='utf-8'))
ct, fixed, rfp, gv, ts = J('mau5a_cau_truc.json'), J('mau5a_fixed.json'), J('rfp_yeu_cau.json'), J('gia_tri_tinh.json'), J('tham_so_chu_dau_tu.json')

x = zipfile.ZipFile(a.hsmt).read('word/document.xml').decode('utf-8')
t = ' '.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', x))
d = docx.Document(a.hsmt)
N = lambda s: re.sub(r'\s+', ' ', s).strip()
norm = lambda s: re.sub('[\u2013\u2014]', '-', re.sub(r'\s+', '', s))
findings = []

# Phần do bên lập soạn (bỏ Chương I, Chương VI, Mục 1 Chương III là nguyên văn mẫu)
FIXED_FIRST = ('1. Phạm vi gói thầu', '1. Định nghĩa')
own, in_muc1 = [], False
for p in d.paragraphs:
    s0 = p.text.strip()
    if s0.startswith('Mục 1. Đánh giá tính hợp lệ'):
        in_muc1 = True; own.append(s0); continue
    if s0.startswith('Mục 2. Tiêu chuẩn đánh giá về năng lực'):
        in_muc1 = False
    if not in_muc1:
        own.append(s0)
tables = []
for tb in d.tables:
    rows = [[N(c.text) for c in r.cells] for r in tb.rows]
    tables.append(rows)
    if not tb.rows[0].cells[0].text.strip().startswith(FIXED_FIRST):
        own += [c for r in rows for c in r]
t_own = ' '.join(own)
t_own_low = t_own.lower()


def add(ma, nhom, mo_ta, vi_tri='', cach_sua='can_nguoi_dung', de_xuat='', muc_do='loi'):
    findings.append(dict(ma=ma, muc_do=muc_do, nhom=nhom, vi_tri=vi_tri, mo_ta=mo_ta, cach_sua=cach_sua, de_xuat=de_xuat))


def ctx(m, w=50, src=None):
    src = t if src is None else src
    return src[max(0, m.start() - w): m.end() + w]


def find_table(pred):
    for rows in tables:
        try:
            if pred(rows): return rows
        except Exception: pass
    return None


# 1. Pháp lý
for m in re.finditer(r'24/2024', t):
    add('PL01', 'phap_ly', 'Viện dẫn Nghị định 24/2024 đã hết hiệu lực', ctx(m), 'tu_dong', 'Thay bằng Nghị định số 214/2025/NĐ-CP')
if 'đặt cọc' in t_own: add('PL02', 'phap_ly', 'Đấu thầu qua mạng không dùng đặt cọc')
if re.search(r'\d+ bản (in|giấy)|bản in', t_own): add('PL03', 'phap_ly', 'Mẫu 5A chỉ cho nộp một bộ E-HSDT qua Hệ thống')
for key in ('luat', 'nghi_dinh', 'thong_tu'):
    so = re.search(r'số (\S+)', ts['can_cu_phap_ly'][key]).group(1)
    if so not in t: add('PL04', 'phap_ly', 'Thiếu căn cứ ' + so)

# 2. Chỗ giữ chỗ và dấu vết của mẫu chưa điền
for pat, mo in ((r'<<SOẠN', 'Còn chỗ giữ chỗ chưa soạn'), (r'_{3,}', 'Còn ký hiệu ___ của mẫu'), (r'\[(ghi|nêu|Hệ thống)', 'Còn hướng dẫn trong ngoặc vuông của mẫu'), (r'\w\(\d\)', 'Còn ký hiệu chú thích (n) của mẫu')):
    for m in re.finditer(pat, t_own):
        add('CT00', 'cau_truc', mo, ctx(m, 40, t_own))

# 3. Cấu trúc theo Mẫu 5A
for key in ['Chương I. Chỉ dẫn nhà thầu', 'Chương II. Bảng dữ liệu đấu thầu', 'Chương III. Tiêu chuẩn đánh giá E-HSDT',
            'Chương IV. Biểu mẫu mời thầu và dự thầu', 'Chương V. Yêu cầu về kỹ thuật', 'Chương VI. Điều kiện chung của hợp đồng',
            'Chương VII. Điều kiện cụ thể của hợp đồng', 'Phần 4. Biểu mẫu hợp đồng', 'Bảng số 01', 'Bảng số 02', 'Mẫu số 01A',
            'Mục 3. Tiêu chuẩn đánh giá về kỹ thuật', 'Mục 4. Tiêu chuẩn đánh giá về tài chính', 'Mục 5. Phương án kỹ thuật thay thế', 'Mục 6. Trường hợp gói thầu chia',
            'Bảng tiêu chí chi tiết', '3.5. Phân định trách nhiệm và giả định']:
    if key not in t: add('CT01', 'cau_truc', 'Thiếu ' + key)
firsts = {norm(r[0]) for rows in tables for r in rows if r}
for it in ct['e_bdl'] + ct['e_dkct']:
    if norm(it['ma']) not in firsts: add('CT02', 'cau_truc', 'Thiếu mục so với Mẫu 5A: ' + it['ma'])
for r in ct['chuong4_bieu_mau']:
    if r and r[0].isdigit():
        ma = re.match(r'(Mẫu số \w+)\.', r[1]).group(1)
        if ma + '.' not in t: add('CT03', 'cau_truc', 'Danh mục Chương IV thiếu ' + ma)
for f in ct['phan4']:
    if f['ma'] + '.' not in t: add('CT03', 'cau_truc', 'Phần 4 thiếu ' + f['ma'])
# Bảng 01: các cột cố định của mẫu (nhà thầu độc lập, liên danh, tài liệu) phải giữ nguyên
b01_mau = [r for r in ct['bang01'] if r and r[0] in ('1', '2', '3', '4')]
b01 = find_table(lambda rows: rows[0][0] == 'TT' and 'Liên danh' in ' '.join(rows[0]))
if not b01:
    add('CT04', 'cau_truc', 'Không tìm thấy Bảng số 01')
else:
    for rm in b01_mau:
        rr = next((r for r in b01 if r[0] == rm[0]), None)
        if not rr: add('CT04', 'cau_truc', f'Bảng số 01 thiếu dòng {rm[0]}'); continue
        if rr[3:7] != rm[3:7]:
            add('CT05', 'cau_truc', f'Bảng số 01 dòng {rm[0]} lệch cột cố định của mẫu: {rr[3:7]} thay vì {rm[3:7]}', rm[1])

# 4. Nguyên văn
have = norm(t)
for name, key in (('Chương I', 'c1'), ('Chương VI', 'c6')):
    miss = [s for l, r in fixed[key] for s in l + r if norm(s) not in have]
    if miss: add('NV01', 'nguyen_van', f'{name} lệch nguyên văn mẫu ở {len(miss)} đoạn', miss[0][:120], 'tu_dong', 'Chạy lại build_hsmt.py với mau5a_fixed.json')
miss = [s for s in fixed['muc1'] if norm(s) not in have]
if miss: add('NV01', 'nguyen_van', f'Mục 1 Chương III lệch nguyên văn mẫu ở {len(miss)} đoạn', miss[0][:120], 'tu_dong')

# 5. Số liệu theo giá trị đã tính và tham số
L, T, NL, TH, HD, KT = gv['lich'], gv['tien'], gv['nang_luc'], gv['thoi_han'], gv['hop_dong'], gv['ky_thuat']
for k in ('phat_hanh', 'dong_thau', 'han_nha_thau_de_nghi_lam_ro', 'han_chu_dau_tu_tra_loi', 'hop_dong_hieu_luc_du_kien', 'cut_over'):
    if L[k] not in t_own: add('SL01', 'so_lieu', f'Thiếu mốc {k} = {L[k]}')
if L.get('hoi_nghi_tien_dau_thau') and L['hoi_nghi_tien_dau_thau']['ngay'] not in t_own: add('SL01', 'so_lieu', 'Thiếu ngày hội nghị tiền đấu thầu')
for k, v in T.items():
    if v + ' đồng' not in t_own: add('SL02', 'so_lieu', f'Không tìm thấy {k} = {v} đồng')
for k in ('moc_lich_su_hop_dong', 'moc_hop_dong_tuong_tu'):
    if 'Từ ngày ' + NL[k] not in t_own: add('SL03', 'so_lieu', f'Mốc {k} phải là {NL[k]}')
for m in re.finditer(r'Từ ngày (\d{2} tháng \d{2} năm \d{4}) đến thời điểm đóng thầu', t_own):
    if m.group(1) not in (NL['moc_lich_su_hop_dong'], NL['moc_hop_dong_tuong_tu']):
        add('SL03', 'so_lieu', 'Mốc năng lực không khớp giá trị đã tính', m.group(0))
if f"≥ {TH['hieu_luc_hsdt_ngay']} ngày" not in t_own: add('SL04', 'so_lieu', 'Hiệu lực E-HSDT không khớp')
if f"≥ {TH['hieu_luc_bao_dam_du_thau_ngay']} ngày" not in t_own: add('SL04', 'so_lieu', 'Hiệu lực bảo đảm dự thầu không khớp')
if TH['hieu_luc_bao_dam_thuc_hien_den'] not in t_own: add('SL04', 'so_lieu', 'Hiệu lực bảo đảm thực hiện hợp đồng không khớp')
if f"Giá trị bảo đảm thực hiện hợp đồng: {HD['bao_dam_thuc_hien_phan_tram']}%" not in t_own: add('SL05', 'so_lieu', 'Bảo đảm thực hiện hợp đồng không khớp tham số')
if f"{HD['nha_thau_phu_toi_da_phan_tram']}% giá dự thầu" not in t_own: add('SL05', 'so_lieu', 'Tỷ lệ nhà thầu phụ không khớp tham số')
if HD['tam_ung_phan_tram'] and f"Tạm ứng: {HD['tam_ung_phan_tram']}%" not in t_own: add('SL06', 'so_lieu', 'Mục 12.1 không khớp tỷ lệ tạm ứng')
m = re.search(r'E-ĐKC 13\.1(.*?)E-ĐKC 13\.2', t_own)
if m:
    pct = [int(v) for v in re.findall(r'Đợt \d+: (\d+)%', m.group(1))]
    ky_vong = [x['phan_tram'] for x in HD['thanh_toan']]
    if pct != ky_vong: add('SL06', 'so_lieu', f'Các đợt thanh toán {pct} khác tham số {ky_vong}')
else:
    add('SL06', 'so_lieu', 'Không tìm thấy Mục 13.1 E-ĐKCT')


def kiem_hang_muc_dot(van_13_1):
    # Mỗi hạng mục Mẫu số 01A gắn đúng một đợt thanh toán, ghi trong Mục 13.1 E-ĐKCT dạng (hạng mục 1, 2 Mẫu số 01A)
    m01 = find_table(lambda rows: rows[0][0] == 'STT' and re.sub(r'\(\d\)$', '', rows[0][1]).strip() == 'Danh mục dịch vụ')
    stt = {int(r[0]) for r in (m01 or [])[1:] if r[0].isdigit()}
    gan = [s for v in hang_muc_theo_dot(van_13_1).values() for s in v]
    de_xuat = 'Soạn lại DKC_13_1: mỗi đợt gắn sản phẩm ghi (hạng mục a, b Mẫu số 01A), mỗi hạng mục thuộc đúng một đợt'
    if not gan:
        add('SL11', 'so_lieu', 'Mục 13.1 E-ĐKCT chưa nêu hạng mục Mẫu số 01A gắn từng đợt thanh toán', van_13_1[:90], 'can_nguoi_dung', de_xuat); return
    for mo_ta, ds in (('dẫn chiếu hạng mục không có trong Mẫu số 01A', set(gan) - stt), ('gắn hạng mục vào nhiều đợt', {s for s in gan if gan.count(s) > 1}),
                      ('chưa gắn đợt thanh toán cho hạng mục Mẫu số 01A', stt - set(gan))):
        if ds: add('SL11', 'so_lieu', f'Mục 13.1 E-ĐKCT {mo_ta}: {sorted(ds)}', van_13_1[:90], 'can_nguoi_dung', de_xuat)


if m:
    kiem_hang_muc_dot(m.group(1))
b02 = find_table(lambda rows: rows[0][:2] == ['STT', 'Vị trí công việc'])
if b02:
    have02 = [(r[1], r[2]) for r in b02[1:]]
    want02 = [(x['vi_tri'], str(x['so_luong'])) for x in ts['nhan_su_chu_chot']]
    if have02 != want02: add('SL07', 'so_lieu', f'Bảng số 02 {have02} khác tham số {want02}')
bkt = find_table(lambda rows: rows[0][:2] == ['STT', 'Tiêu chí tổng quát'])
if bkt:
    rows = [r for r in bkt[1:] if r[0]]
    if [int(r[2]) for r in rows] != KT['trong_so'] or [int(r[3]) for r in rows] != KT['diem_toi_thieu']:
        add('SL08', 'so_lieu', 'Điểm tiêu chí kỹ thuật không khớp giá trị đã tính')
    if any(re.search(r'giá dự thầu|chi phí ước tính', r[1], re.I) for r in rows):
        add('SL09', 'so_lieu', 'Tiêu chí kỹ thuật không được chấm giá')

# 5b. Tiêu chí kỹ thuật chi tiết (Mục 3, 3.1 Chương III Mẫu 5A), tiêu chí 5 không chấm điều kiện của nhà thầu
KHONG_TINH_DIEM = r'điều kiện khác|đề nghị (của nhà thầu )?đối với Chủ đầu tư'
TAI_CHINH = r'chi phí ước tính|giá dự thầu|Mẫu số 11A'
MUC_CHAM = re.compile(r'^(.+?) \((\d+(?:,\d+)?) điểm\): \S')
so_dk = lambda s: float(s.replace(',', '.'))
o_cell = lambda c: [N(x.text) for x in c.paragraphs if x.text.strip()]


def doan_muc(bat_dau, ket_thuc):
    """Đoạn văn và bảng giữa hai tiêu đề Heading 2, theo thứ tự trong thân văn bản."""
    doan, bang, trong = [], [], False
    for el in d.element.body.iterchildren():
        if el.tag == qn('w:p'):
            par = Paragraph(el, d)
            s = par.text.strip()
            if par.style is not None and par.style.name == 'Heading 2':
                if s.startswith(bat_dau):
                    trong = True; continue
                if trong and s.startswith(ket_thuc):
                    break
            if trong and s:
                doan.append(s)
        elif el.tag == qn('w:tbl') and trong:
            bang.append(Table(el, d))
    return doan, bang


def loi_dong_con(cells):
    try:
        diem = so_dk(' '.join(cells[2]))
    except ValueError:
        return ['điểm tối đa không phải là số']
    out = [] if ' '.join(cells[3]) == 'Không' or ' '.join(cells[3]).startswith('Có') else ['thiếu đánh dấu tiêu chí cơ bản (Có hoặc Không)']
    muc = [MUC_CHAM.match(x) for x in cells[4]]
    if len(muc) < 3 or not all(muc):
        return out + ['cần ít nhất 3 mức chấm dạng: mức (n điểm): mô tả']
    ds = [so_dk(m.group(2)) for m in muc]
    if ds[0] != diem or ds[-1] != 0 or any(x <= y for x, y in zip(ds, ds[1:])):
        out.append(f'điểm các mức {ds} phải giảm dần từ điểm tối đa {diem} đến 0')
    return out


def kiem_bang_chi_tiet(tb):
    dong = [[o_cell(c) for c in r.cells] for r in tb.rows[1:]]
    for j, w in enumerate(KT['trong_so'], 1):
        con = [r for r in dong if re.fullmatch(rf'{j}\.\d+', ' '.join(r[0]))]
        if not 3 <= len(con) <= 6:
            add('KT02', 'ky_thuat', f'Tiêu chí tổng quát {j} có {len(con)} tiêu chí chi tiết, Mục 3 yêu cầu từ 3 đến 6', f'Tiêu chí {j}', 'can_nguoi_dung', 'Soạn lại TIEU_CHI_CHI_TIET')
        for r in con:
            for e in loi_dong_con(r):
                add('KT04', 'ky_thuat', f"Tiêu chí chi tiết {' '.join(r[0])}: {e}", ' '.join(r[1])[:90], 'can_nguoi_dung', 'Soạn lại TIEU_CHI_CHI_TIET')
        try:
            tong = sum(so_dk(' '.join(r[2])) for r in con)
        except ValueError:
            continue
        if con and abs(tong - w) > 1e-9:
            add('KT03', 'ky_thuat', f'Tiêu chí tổng quát {j}: tổng điểm tiêu chí chi tiết {tong:g} khác điểm tối đa {w}', f'Tiêu chí {j}', 'can_nguoi_dung', 'Soạn lại TIEU_CHI_CHI_TIET')
    # chỉ dòng tiêu chí chi tiết j.i được tính là có chấm điểm (dòng tiêu chí tổng quát chỉ là tên nhóm)
    return ' '.join(x for r in dong if re.fullmatch(r'\d+\.\d+', ' '.join(r[0])) for c in r for x in c)


def kiem_do_phu(bang_txt, doan_txt):
    for yc in rfp['yeu_cau']:
        if yc['nhom'] not in ('de_xuat', 'chinh_sach') or re.search(TAI_CHINH, ' '.join(yc['tu_khoa_kiem_tra']), re.I):
            continue
        if re.search(KHONG_TINH_DIEM, ' '.join(yc['tu_khoa_kiem_tra']), re.I):
            thieu = [k for k in yc['tu_khoa_kiem_tra'] if k.lower() not in doan_txt.lower()]
            if thieu: add('KT06', 'ky_thuat', f"Mục 3 chưa nêu nội dung {yc['ma']} không tính điểm: thiếu {thieu}", yc['nguyen_van'][:90])
            continue
        thieu = [k for k in yc['tu_khoa_kiem_tra'] if k.lower() not in bang_txt.lower()]
        if thieu:
            add('KT05', 'ky_thuat', f"Yêu cầu {yc['ma']} chưa được tiêu chí chi tiết nào chấm: thiếu {thieu}", yc['nguyen_van'][:90], 'can_nguoi_dung', 'Bổ sung tiêu chí chi tiết phủ yêu cầu này (TIEU_CHI_CHI_TIET)')


def kiem_khong_cham_dieu_kien(ten_bang, txt):
    ms = list(re.finditer(KHONG_TINH_DIEM, txt, re.I))
    if ms:
        add('KT06', 'ky_thuat', f'Bảng tiêu chí {ten_bang} chấm điều kiện, đề nghị của nhà thầu đối với Chủ đầu tư ({len(ms)} chỗ), Mục 24, 25 E-CDNT coi là đặt điều kiện',
            ctx(ms[0], 60, txt), 'can_nguoi_dung', 'Bỏ khỏi tiêu chí (TIEU_CHI_5, TIEU_CHI_CHI_TIET), chỉ chấm đề xuất có lợi và thông tin theo Điều 19, 20 Nghị định số 214/2025/NĐ-CP')


def kiem_muc_3():
    doan, bang = doan_muc('Mục 3. Tiêu chuẩn đánh giá về kỹ thuật', 'Mục 4. Tiêu chuẩn đánh giá về tài chính')
    kiem_khong_cham_dieu_kien('tổng quát', ' '.join(r[1] for r in (bkt or [])[1:]))
    if not re.search(r'Mục 24 và Mục 25 E-CDNT', ' '.join(doan)):
        add('KT06', 'ky_thuat', 'Mục 3 thiếu câu: điều kiện, đề nghị của nhà thầu xét theo Mục 24, 25 E-CDNT, không tính điểm', 'Mục 3 Chương III', 'can_nguoi_dung', 'Lập lại khung (lap_khung.py)')
    tb = next((x for x in bang if [N(c.text) for c in x.rows[0].cells[:2]] == ['STT', 'Tiêu chí chi tiết']), None)
    if tb is None:
        add('KT01', 'ky_thuat', 'Mục 3 Chương III thiếu bảng tiêu chí chi tiết có điểm và mức chấm (Mục 3, 3.1 Mẫu 5A)', 'Mục 3 Chương III', 'can_nguoi_dung', 'Lập lại khung và soạn TIEU_CHI_CHI_TIET')
        return
    bang_txt = kiem_bang_chi_tiet(tb)
    if re.search(r'giá dự thầu|chi phí ước tính', bang_txt, re.I): add('SL09', 'so_lieu', 'Tiêu chí kỹ thuật chi tiết không được chấm giá')
    kiem_khong_cham_dieu_kien('chi tiết', bang_txt)
    kiem_do_phu(bang_txt, ' '.join(doan))


kiem_muc_3()

# 5c. Tham số chờ Chủ đầu tư quyết định, câu mẫu E-ĐKC 15, phân định trách nhiệm Chương V mục 3.5
thieu_ts = tham_so_thieu(ts)
for k, vt in thieu_ts:
    add('TS01', 'cho_nguoi_dung', f'Tham số {k} chưa có, chờ Chủ đầu tư quyết định', vt, 'can_nguoi_dung',
        'Bổ sung vào tham_so_chu_dau_tu.json (references/tham_so.md), chạy lại từ bước chuẩn bị', 'canh_bao')
dau = list(re.finditer(r'\[' + DAU_CHO + r'[^\]]*\]', t))
if dau and not thieu_ts:
    add('TS02', 'tham_so', 'Còn dấu chờ Chủ đầu tư quyết định trong khi tham số đã đủ', ctx(dau[0]), 'can_nguoi_dung', 'Chạy lại chay_quy_trinh.py chuan_bi rồi hoan_thien')
dkc15 = next((' '.join(r[1:]) for rows in tables for r in rows if r and norm(r[0]) == norm('E-ĐKC 15')), None)
for cau in ('Trường hợp không kịp thời khắc phục sai sót theo thông báo của Chủ đầu tư', 'Bồi thường thiệt hại trên cơ sở'):
    if dkc15 is not None and cau not in dkc15 and not (cau.startswith('Bồi') and 'Bồi thường thiệt hại: Không áp dụng' in dkc15):
        add('NV02', 'nguyen_van', f'E-ĐKC 15 thiếu câu mẫu: {cau} ...', dkc15[:90], 'can_nguoi_dung', 'Giữ câu mẫu, tham số còn trống thì ghi dấu chờ Chủ đầu tư quyết định (lap_khung.py)')
GPD = gv.get('phan_dinh') or {}
for ten, han in [('hạn cung cấp thông tin hệ thống hiện có', GPD.get('han_cung_cap_thong_tin_he_thong'))] + [(f'thời hạn {k}', v) for k, v in (GPD.get('thoi_han_ha_tang') or {}).items()]:
    if han and han not in t_own: add('SL10', 'so_lieu', f'Chương V mục 3.5 thiếu {ten} {han}')
bpd = find_table(lambda rows: rows[0] == ['STT', 'Nội dung', 'Phân định'])
dong_pd = lambda dau: next((r for r in bpd or [] if r[1].startswith(dau)), None)  # noqa: E731
ky_vong_pd = [(ten, BEN.get(lay(ts, 'phan_dinh.ha_tang_ben_cung_cap.' + k))) for k, ten in HA_TANG]
ky_vong_pd += [('Thay đổi phía hệ thống hiện có', THAY_DOI_HE_THONG.get(lay(ts, 'phan_dinh.thay_doi_he_thong_hien_co'))),
               ('Môi trường kiểm thử vận hành', KIEM_THU_VAN_HANH.get(lay(ts, 'phan_dinh.kiem_thu_van_hanh_tren')))]
for ten, gt_pd in ky_vong_pd:
    r = dong_pd(ten)
    if gt_pd and (r is None or not r[2].startswith(gt_pd)):
        add('SL10', 'so_lieu', f'Chương V mục 3.5: {ten} không khớp tham số ({gt_pd})')
# Thông tin Chủ đầu tư tại E-CDNT 5.1 (c), 7.5, 38.1, E-ĐKC 1.3, 4.1 (Mẫu 5A yêu cầu ghi đủ trước khi phát hành)
for k, ten in LIEN_HE_CDT:
    v = str(ts['chu_dau_tu'].get(k) or '').strip()
    if v and v not in t_own: add('SL12', 'so_lieu', f'Thiếu {ten} của Chủ đầu tư theo tham số: {v}', 'E-CDNT 5.1 (c), 7.5, 38.1, E-ĐKC 1.3, 4.1')
dkc13 = next((' '.join(r[1:]) for rows in tables for r in rows if r and norm(r[0]) == norm('E-ĐKC 1.3')), '')
if re.search(r'ghi khi hoàn thiện hợp đồng', dkc13):
    add('SL12', 'so_lieu', 'E-ĐKC 1.3 hoãn thông tin Chủ đầu tư đến khi hoàn thiện hợp đồng, Mẫu 5A yêu cầu ghi đủ trước khi phát hành', dkc13[:90], 'can_nguoi_dung', 'Bổ sung tham số chu_dau_tu, lập lại khung')
# Đoạn quy phạm Mục 2.2 Chương III chép nguyên văn từ mẫu
for dau in DOAN_2_2_NHAN_SU + (DOAN_2_2_C,):
    mau22 = next((x for x in ct['huong_dan_chuong3'] if x.startswith(dau)), None)
    if mau22 and norm(mau22) not in have:
        add('NV03', 'nguyen_van', f'Mục 2.2 Chương III thiếu đoạn quy phạm của Mẫu 5A: {dau} ...', mau22[:90], 'tu_dong', 'Lập lại khung (lap_khung.py chép nguyên văn)')
    elif mau22 and have.count(norm(mau22)) > 1:
        add('NV03', 'nguyen_van', f'Mục 2.2 Chương III lặp đoạn quy phạm của Mẫu 5A: {dau} ...', mau22[:90], 'can_nguoi_dung',
            'Soạn lại NHAN_SU_LY_DO chỉ gồm lý do cần nhân sự chủ chốt, script đã chép các đoạn quy phạm')
# Bảng số 01 dòng 3 (thuế GTGT, ghi chú (6)), dòng 4 (cộng gộp hạng mục, ghi chú (11)) theo tham số
so_vn = lambda v: str(v).replace('.', ',')  # noqa: E731
CS = ts['chinh_sach']
d3 = next((r[2] for r in (b01 or []) if r[0] == '3'), '')
d4 = next((r[2] for r in (b01 or []) if r[0] == '4'), '')
THUE = NL.get('thue_gtgt') or {}
if THUE.get('loai') == 'khong_chiu_thue' and THUE.get('can_cu') and THUE['can_cu'][1:] not in d3:
    add('SL13', 'so_lieu', 'Bảng số 01 dòng 3 thiếu căn cứ không chịu thuế GTGT theo tham số', d3[-90:], 'can_nguoi_dung', 'Lập lại khung')
if THUE.get('loai') == 'chiu_thue' and f"thuế suất {so_vn(THUE['thue_suat_phan_tram'])}%" not in d3:
    add('SL13', 'so_lieu', 'Bảng số 01 dòng 3 thiếu thuế suất GTGT dùng để trừ khi tính doanh thu', d3[-90:], 'can_nguoi_dung', 'Lập lại khung')
if 'theo hướng dẫn của Mẫu số 5A' in d4:
    add('SL13', 'so_lieu', 'Bảng số 01 dòng 4 chỉ viện dẫn chung hướng dẫn của Mẫu số 5A, chưa nêu cụ thể', d4[-90:], 'can_nguoi_dung', 'Lập lại khung')
if CS.get('hop_dong_tuong_tu_cong_gop') is False and 'trong các hợp đồng khác nhau' in d4:
    add('SL13', 'so_lieu', 'Bảng số 01 dòng 4 cho cộng gộp hạng mục trong khi tham số hop_dong_tuong_tu_cong_gop = false', d4[-90:], 'can_nguoi_dung', 'Lập lại khung')
for hm_ in (NL.get('hang_muc_tuong_tu') or []) if CS.get('hop_dong_tuong_tu_cong_gop') else []:
    if f"{hm_['ten']} {hm_['gia_tri']} đồng" not in d4:
        add('SL13', 'so_lieu', f"Bảng số 01 dòng 4 thiếu giá trị hạng mục tương tự {hm_['ten']} {hm_['gia_tri']} đồng (ghi chú (11))", d4[-90:], 'can_nguoi_dung', 'Lập lại khung')
# E-ĐKC 15 phạt không huy động nhân sự chủ chốt, E-ĐKC 18 (e) lùi mốc khi Chủ đầu tư cung cấp chậm
ns_ = CS.get('phat_khong_huy_dong_nhan_su_phan_tram_ngay')
if ns_ == 0 and dkc15 is not None and 'nhân sự chủ chốt không được huy động' in dkc15:
    add('SL14', 'so_lieu', 'E-ĐKC 15 còn câu phạt không huy động nhân sự chủ chốt trong khi tham số bằng 0 (không áp dụng)', dkc15[-90:], 'can_nguoi_dung', 'Lập lại khung')
if ns_ and dkc15 is not None and f'{so_vn(ns_)}% giá hợp đồng cho mỗi ngày không huy động' not in dkc15:
    add('SL14', 'so_lieu', f'E-ĐKC 15 thiếu mức phạt {so_vn(ns_)}% mỗi ngày không huy động nhân sự chủ chốt', dkc15[-90:], 'can_nguoi_dung', 'Lập lại khung')
dkc18e = next((' '.join(r[1:]) for rows in tables for r in rows if r and norm(r[0]) == norm('E-ĐKC 18 (e)')), '')
gh_ = CS.get('gia_han_khi_chu_dau_tu_cung_cap_cham')
if gh_ is True and ('lùi tương ứng số ngày chậm' not in dkc18e or 'Mục 18 (e) E-ĐKCT' not in t_own):
    add('SL14', 'so_lieu', 'E-ĐKC 18 (e) và câu dẫn mục 3.5 Chương V chưa thống nhất việc lùi mốc khi Chủ đầu tư cung cấp chậm', dkc18e[:90], 'can_nguoi_dung', 'Lập lại khung')
if gh_ is False and 'không có' not in dkc18e:
    add('SL14', 'so_lieu', 'E-ĐKC 18 (e) ghi trường hợp lùi mốc trong khi tham số gia_han_khi_chu_dau_tu_cung_cap_cham = false', dkc18e[:90], 'can_nguoi_dung', 'Lập lại khung')
# Vòng của tệp Word (build_hsmt.py ghi khi hoan_thien) phải là vòng đang kiểm tra
VONG_WORD, MD5_NOI_DUNG = vong_cua_word(a.hsmt)
VONG_HIEN = doc_so(a.lam_viec)['vong'] + 1
if not VONG_WORD:
    add('CT07', 'cau_truc', 'Tệp Word không ghi vòng kiểm tra (tạo ngoài chay_quy_trinh.py hoan_thien)', os.path.basename(a.hsmt), 'can_nguoi_dung', 'Chạy chay_quy_trinh.py hoan_thien rồi kiểm tra tệp mới', 'canh_bao')
elif VONG_WORD != str(VONG_HIEN):
    add('CT07', 'cau_truc', f'Tệp Word thuộc vòng {VONG_WORD}, vòng đang kiểm tra là {VONG_HIEN}', os.path.basename(a.hsmt), 'can_nguoi_dung', 'Kiểm tra tệp HSMT_ket_qua.docx hoặc HSMT_du_thao.docx do hoan_thien của vòng này tạo')

# 6. Yêu cầu RFP
for yc in rfp['yeu_cau']:
    thieu = [k for k in yc['tu_khoa_kiem_tra'] if k.lower() not in t_own_low]
    if thieu:
        add('RF01', 'rfp', f"Yêu cầu {yc['ma']} chưa thể hiện: thiếu {thieu}", yc['nguyen_van'][:100], 'can_nguoi_dung', 'Soạn bổ sung tại ' + yc['vi_tri_hsmt'])
# Sản phẩm bàn giao của RFP được kiểm tra qua từ khóa của các yêu cầu nhóm san_pham.

# 7. Định dạng
if re.search(r'<w:b/>|<w:b w:val="(1|true)"/>', x): add('DD01', 'dinh_dang', 'Có chữ in đậm', cach_sua='tu_dong', de_xuat='Chạy lại build_hsmt.py')
if re.search(r'<w:caps/>|<w:caps w:val="(1|true)"/>', x): add('DD02', 'dinh_dang', 'Có chữ viết hoa toàn bộ', cach_sua='tu_dong')
for m in re.finditer('[\u2013\u2014]', t_own): add('DD03', 'dinh_dang', 'Có gạch dài', ctx(m, 30, t_own), 'tu_dong', 'Đổi thành gạch ngang đơn')
fonts = set(re.findall(r'w:ascii="([^"]+)"', x)) - {'Times New Roman'}
if fonts: add('DD04', 'dinh_dang', 'Có font khác Times New Roman: ' + ', '.join(fonts), cach_sua='tu_dong')
for m in re.finditer(r'[“”"]', t_own): add('DD05', 'dinh_dang', 'Có ngoặc kép, kiểm tra có dùng để nhấn mạnh không', ctx(m, 30, t_own), muc_do='canh_bao')

# 8. Văn phong
STYLE = {r'(?<!chủ )\b[Cc]hốt\b': None, r'\bxong\b': None, r'chạy được': 'hoạt động được', r'dùng được': 'có thể sử dụng',
         r'lỗi chặn': 'lỗi nghiêm trọng làm gián đoạn vận hành', r'đóng lỗi': 'khắc phục dứt điểm lỗi',
         r'go-live': 'đưa hệ thống vào vận hành', r'\brồi\b': None, r'dưới đây là': None, r'[→✓✔•]': None}
for pat, rep in STYLE.items():
    for m in re.finditer(pat, t_own):
        add('VP01', 'van_phong', 'Từ hoặc ký hiệu không phù hợp văn phong: ' + m.group(0), ctx(m, 40, t_own),
            'tu_dong' if rep else 'can_nguoi_dung', rep or 'Chọn cách diễn đạt trang trọng theo ngữ cảnh')
for s in own:
    if s.count(';') > 3:
        add('VP02', 'van_phong', f'Câu có {s.count(";")} dấu chấm phẩy', s[:100], 'can_nguoi_dung', 'Tách câu hoặc dùng từ nối', 'canh_bao')

# 9. Mang phát hiện còn mở của vòng trước sang (sổ phát hiện), kết luận
so = doc_so(a.lam_viec)
findings = mang_sang(findings, so)
loi = [f for f in findings if f['muc_do'] == 'loi' and not f.get('giai_trinh')]
kl = ket_luan(findings)
cho_nd = [f for f in findings if f['ma'] in ('TS01', 'PH01') or f.get('mang_sang')]
report = {'file': os.path.basename(a.hsmt), 'md5': md5_tep(a.hsmt), 'vong_word': VONG_WORD, 'md5_noi_dung': MD5_NOI_DUNG, 'vong': so['vong'] + 1, 'ket_luan': kl, 'ket_luan_lop_a': ket_luan_lop_a(findings),
          'so_loi': len(loi), 'so_canh_bao': len(findings) - len(loi), 'so_cho_nguoi_dung': len(cho_nd), 'phat_hien': findings}
if a.json:
    json.dump(report, open(a.json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
for f in findings:
    gt = ' | đã giải trình' if f.get('giai_trinh') else ''
    print(f"[{f['muc_do']}] {f['ma']} {f['nhom']}: {f['mo_ta']} | {f['vi_tri'][:90]} | {f['cach_sua']}{gt}")
ghi_chu = {'dat': '', 'chua_dat': '', 'cho_nguoi_dung': f', chưa hoàn tất, chờ người dùng: {len(cho_nd)} mục (TS01 tham số Chủ đầu tư, PH01 phát hiện vòng trước)'}[kl]
print(f"Kết luận: {kl} ({report['so_loi']} lỗi, {report['so_canh_bao']} cảnh báo{ghi_chu})")
sys.exit({'dat': 0, 'chua_dat': 1, 'cho_nguoi_dung': 2}[kl])
