"""Cập nhật Slide_bao_cao_ket_qua.pptx theo bộ tài liệu 657 MD (du_lieu/ của dự án IT6460).

Giữ nguyên thiết kế: chỉ thay chữ trong các khung có sẵn, thêm hoặc bớt hàng bảng bằng cách nhân bản hàng mẫu,
thay dữ liệu biểu đồ, co giãn thanh MD theo tỷ lệ và viết lại ghi chú người trình bày.
"""
import copy
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)
S = list(p.slides)


def shp(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def _set_para(par, text):
    runs = par.runs
    if not runs:
        r = par.add_run()
        r.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_text(tf, lines):
    """Thay nội dung khung chữ theo danh sách dòng, giữ định dạng của đoạn tương ứng (đoạn thêm chép đoạn cuối)."""
    if isinstance(lines, str):
        lines = [lines]
    paras = list(tf.paragraphs)
    # giữ các đoạn có chữ làm mẫu, bỏ đoạn rỗng ở cuối nếu có
    mau = [q for q in paras if q.text.strip()] or paras
    for q in paras:
        if q not in mau:
            q._p.getparent().remove(q._p)
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):
        moi = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(moi)
        paras = list(tf.paragraphs)
    for q in paras[len(lines):]:
        q._p.getparent().remove(q._p)
    for q, t in zip(tf.paragraphs, lines):
        _set_para(q, t)


def t(slide, sid, lines):
    set_text(shp(slide, sid).text_frame, lines)


def table_rows(slide, sid, rows, mau_than=1, giu_cuoi=True):
    """Ghi bảng: hàng 0 là tiêu đề giữ nguyên, rows là các hàng nội dung (gồm hàng tổng nếu giu_cuoi).
    Hàng thêm nhân bản từ hàng mẫu mau_than, hàng tổng nhân bản từ hàng cuối cũ."""
    tbl = shp(slide, sid).table
    trs = list(tbl._tbl.tr_lst)
    tieu_de, than, cuoi = trs[0], trs[1:-1] if giu_cuoi else trs[1:], trs[-1] if giu_cuoi else None
    so_than = len(rows) - (1 if giu_cuoi else 0)
    mau = than[mau_than - 1]
    for tr in than:
        tbl._tbl.remove(tr)
    if cuoi is not None:
        tbl._tbl.remove(cuoi)
    for i in range(so_than):
        tbl._tbl.append(copy.deepcopy(than[i] if i < len(than) else mau))
    if cuoi is not None:
        tbl._tbl.append(cuoi)
    for r, vals in zip(list(tbl.rows)[1:], rows):
        for c, v in zip(r.cells, vals):
            set_text(c.text_frame, v)


def header(slide, sid, vals):
    tbl = shp(slide, sid).table
    for c, v in zip(tbl.rows[0].cells, vals):
        set_text(c.text_frame, v)


def bars(slide, ids, vals, legend_ids, nhan):
    """Co giãn các đoạn thanh MD liền nhau theo tỷ lệ, ghi số lên đoạn (kể cả đoạn cuối) và chú thích."""
    segs = [shp(slide, i) for i in ids]
    x0 = segs[0].left
    tong_w = segs[-1].left + segs[-1].width - x0
    tong = float(sum(vals))
    x = x0
    for k, (s, v) in enumerate(zip(segs, vals)):
        w = round(tong_w * v / tong) if k < len(segs) - 1 else x0 + tong_w - x
        s.left, s.width = x, w
        x += w
        if not s.text_frame.text.strip():
            # đoạn cuối không có chữ: chép đoạn văn của đoạn trước để cùng kiểu chữ
            src = segs[k - 1].text_frame
            dst = s.text_frame
            for q in list(dst.paragraphs):
                q._p.getparent().remove(q._p)
            dst._txBody.append(copy.deepcopy(src.paragraphs[0]._p))
        set_text(s.text_frame, str(v))
    for i, (lid, n) in enumerate(zip(legend_ids, nhan)):
        t(slide, lid, n)


def notes(slide, thoi_luong, noi_dung, nguon):
    tf = slide.notes_slide.notes_text_frame
    set_text(tf, [f'Thời lượng: {thoi_luong} giây', noi_dung, f'Nguồn đối chiếu: {nguon}'])


# ---------------- Slide 1: bìa
s = S[0]
t(s, 30, ['Hồ sơ mời thầu theo Mẫu 5A, Thông tư số 79/2025/TT-BTC',
          'WBS, kế hoạch theo sprint và phương án nhân sự',
          'Ước lượng nỗ lực, dự toán, giá chào, kế hoạch kiểm soát và chuyển giao'])
notes(s, 25, 'Nhóm xin báo cáo kết quả bài tập ABC Wear. Nội dung gồm hai phần: hồ sơ mời thầu ở vai trò bên mời thầu '
      'và hồ sơ dự thầu cùng bộ kế hoạch dự án ở vai trò nhà thầu PM_26I_02: WBS, tiến độ, nhân sự, ước lượng, chi phí, '
      'rủi ro, thay đổi và chuyển giao. Mọi số liệu trình bày lấy từ bộ tài liệu đã qua kiểm tra lớp A và lớp B.',
      '[S1], [S2], [S5].')

# ---------------- Slide 2: phạm vi
s = S[1]
t(s, 39, 'ABC Wear lựa chọn nhà thầu xây dựng website bán hàng trực tuyến, kết nối với Main System. '
         'Nhóm đóng vai trò bên mời thầu và nhà thầu PM_26I_02.')
t(s, 44, 'Thành viên')
t(s, 48, 'Sản phẩm')
t(s, 52, 'Đặt hàng')
t(s, 56, 'Tích hợp')
t(s, 57, 'Ngoài phạm vi: quản trị nội bộ, thanh toán trực tuyến qua cổng thanh toán.')
table_rows(s, 58, [
    ['Hồ sơ mời thầu, dự thầu', 'E-HSMT theo Mẫu 5A và HSDT đáp ứng 5 tiêu chí kỹ thuật'],
    ['Phạm vi và WBS', '8 nhóm công việc, 144 gói, 657 MD, có sản phẩm đầu ra'],
    ['Tiến độ và nhân sự', '7 sprint, 12 nhân sự, giao việc theo gói và ma trận RACI'],
    ['Ước lượng và chi phí', 'Ước lượng theo danh mục chức năng, đơn giá theo loại công việc, dự toán, đối chiếu lương và giá chào'],
    ['Kiểm soát dự án', 'Master plan, kế hoạch rủi ro, thay đổi, cấu hình và quản lý phiên bản'],
    ['Chuyển giao', 'Kiểm thử chấp nhận, kế hoạch cut over và điều kiện bàn giao'],
], giu_cuoi=False)
notes(s, 45, 'Bài tập yêu cầu lập phương án triển khai website bán hàng cho ABC Wear. Ở vai trò bên mời thầu, nhóm lập '
      'hồ sơ mời thầu theo Mẫu 5A. Ở vai trò nhà thầu PM_26I_02, nhóm phân rã ba phân hệ chức năng là tìm kiếm sản phẩm, '
      'thành viên, đặt hàng cùng cơ sở dữ liệu bán hàng và tích hợp Main System thành tám nhóm WBS, sau đó lập tiến độ, '
      'phân công, dự toán, hồ sơ dự thầu và các kế hoạch kiểm soát. Quản trị nội bộ và thanh toán trực tuyến qua cổng '
      'thanh toán không thuộc phạm vi.', '[S1], [S2], [S3] mục giả định ngoài phạm vi, [S5].')

# ---------------- Slide 3: HSMT
s = S[2]
t(s, 88, ['Dịch vụ phi tư vấn, đấu thầu rộng rãi trong nước qua mạng',
          'Một giai đoạn một túi hồ sơ, hợp đồng trọn gói',
          'Giá gói thầu 2.000 triệu đồng',
          'Bảo đảm dự thầu 30 triệu đồng (1,5%), cam kết trong đơn dự thầu'])
t(s, 91, ['Năng lực, kinh nghiệm: doanh thu bình quân từ 2 tỷ đồng, hợp đồng tương tự từ 600 triệu đồng',
          'Nhân sự chủ chốt: 01 quản trị dự án, 03 kỹ sư phát triển, 01 kỹ sư phân tích nghiệp vụ hoặc kiểm thử',
          'Kỹ thuật: 5 tiêu chí, tối thiểu 70 trên 100 điểm',
          'Tài chính: phương pháp giá thấp nhất'])
t(s, 94, ['Tạm ứng 15%, thanh toán theo mốc 25%, 30%, 25%, 5%',
          'Bảo đảm thực hiện hợp đồng 3%',
          'Phạt chậm tiến độ 0,1% mỗi ngày, tối đa 8%',
          'Bảo hành 12 tháng, bảo lãnh bảo hành 5%'])
t(s, 95, 'Đủ 23 mục E-BDL và 25 mục E-ĐKCT theo Mẫu 5A. Điều chỉnh so với RFP có căn cứ: nộp qua Hệ thống thay cho '
         '10 bản in, hội nghị tiền đấu thầu thay buổi trình bày, hạn làm rõ ngày 03/3.')
notes(s, 55, 'Ở vai trò bên mời thầu, nhóm lập E-HSMT theo Mẫu 5A cho gói thầu dịch vụ phi tư vấn, đấu thầu rộng rãi '
      'trong nước qua mạng, một giai đoạn một túi hồ sơ, hợp đồng trọn gói, giá gói thầu 2 tỷ đồng. Bảo đảm dự thầu '
      '30 triệu đồng, thấp hơn 50 triệu đồng nên nhà thầu cam kết trong đơn dự thầu. Lịch đấu thầu bảo đảm hợp đồng có '
      'hiệu lực ngày 18/3 và giữ mốc cut over 01/7/2027. Đánh giá theo ba bước: năng lực và kinh nghiệm đạt hoặc không '
      'đạt, kỹ thuật năm tiêu chí tối thiểu 70 điểm, tài chính theo giá thấp nhất. Hợp đồng tạm ứng 15%, bốn đợt thanh '
      'toán gắn mốc xác nhận đặc tả, chức năng chính và tích hợp, đưa vào vận hành và nộp bảo lãnh bảo hành. Hồ sơ đủ '
      '23 mục E-BDL, 25 mục E-ĐKCT và ghi sáu điểm điều chỉnh so với RFP có căn cứ pháp lý.',
      '[S1] E-TBMT, Chương II, III, VII và bảng điều chỉnh so với RFP.')

# ---------------- Slide 4: WBS tổng thể
s = S[3]
cd = CategoryChartData()
cd.categories = ['1. Quản lý dự án', '2. Phân tích và thiết kế chung', '3. Tìm kiếm và hiển thị sản phẩm',
                 '4. Phân hệ thành viên', '5. Phân hệ đặt hàng', '6. CSDL bán hàng và tích hợp',
                 '7. Kiểm thử hệ thống và vận hành', '8. Triển khai và chuyển giao']
cd.add_series('Ba phân hệ chức năng, CSDL và tích hợp', (0, 0, 95, 91, 129, 93, 0, 0))
cd.add_series('Quản lý, thiết kế chung, kiểm thử, chuyển giao', (60, 66, 0, 0, 0, 0, 73, 50))
shp(s, 104).chart.replace_data(cd)
t(s, 106, 'Tổng nỗ lực của 8 nhóm, 144 gói công việc')
t(s, 107, '408')
t(s, 108, 'MD tìm kiếm, thành viên, đặt hàng, tích hợp')
t(s, 109, '249')
t(s, 111, ['Mỗi gói có mã WBS, MD, loại công việc và sản phẩm đầu ra.',
           'Tổng các gói con khớp với nhóm cha theo quy tắc 100%.',
           'Tổng từng nhóm khớp ước lượng theo danh mục chức năng, lệch tối đa 2 MD.'])
notes(s, 50, 'WBS gồm tám nhóm, 144 gói công việc, tổng 657 MD. Ba phân hệ chức năng cùng cơ sở dữ liệu bán hàng và '
      'tích hợp chiếm 408 MD, quản lý dự án, thiết kế chung, kiểm thử hệ thống và chuyển giao chiếm 249 MD. Tổng từng '
      'nhóm bám ước lượng theo danh mục 106 chức năng và 25 đầu việc phi chức năng, nhóm 6 thấp hơn 2 MD và nhóm 7 cao '
      'hơn 2 MD, đều trong dung sai 10%. Mỗi gói lá có sản phẩm đầu ra và truy vết về yêu cầu RFP.',
      '[S3] WBS tổng hợp mục căn cứ ước lượng và bảng đối chiếu.')

# ---------------- Slide 5 đến 7: ba phân hệ
def phan_he(s, tieu_de, tong, nhom, goi, tong_33, pha, ghi_chu, ids_bar, ids_leg, id_tbl, id_md, id_nhom, id_note):
    t(s, s.shapes.title.shape_id, tieu_de)
    rows = [[m, ten, str(md)] for m, ten, md in goi] + [['', f'Tổng gói {nhom}.3 - Lập trình và kiểm thử đơn vị', str(tong_33)]]
    table_rows(s, id_tbl, rows)
    t(s, id_md, f'{tong} MD')
    t(s, id_nhom, f'Tổng WBS {nhom} - ' + {3: 'Tìm kiếm và hiển thị sản phẩm', 4: 'Phân hệ thành viên', 5: 'Phân hệ đặt hàng'}[nhom])
    ten_pha = ['Phân tích yêu cầu', 'Thiết kế', 'Lập trình và kiểm thử đơn vị', 'Kiểm thử phân hệ', 'Nghiệm thu nội bộ, hỗ trợ UAT']
    bars(s, ids_bar, pha, ids_leg, [f'{n}: {v} MD' for n, v in zip(ten_pha, pha)])
    t(s, id_note, ghi_chu)


s = S[4]
phan_he(s, 'Phân hệ tìm kiếm và hiển thị sản phẩm', 95, 3, [
    ('3.3.1', 'Dịch vụ chỉ mục tìm kiếm và cập nhật chỉ mục từ Selling DB', 5),
    ('3.3.2', 'Dịch vụ tìm kiếm theo từ khóa và gợi ý tìm kiếm', 5),
    ('3.3.3', 'Dịch vụ lọc và sắp xếp kết quả tìm kiếm', 5),
    ('3.3.4', 'Dịch vụ danh mục, chi tiết sản phẩm, biến thể, tồn kho và giá khuyến mại hiển thị', 6),
    ('3.3.5', 'Giao diện trang chủ và điều hướng danh mục theo Menu Layer', 4),
    ('3.3.6', 'Giao diện tìm kiếm, kết quả và bộ lọc', 5),
    ('3.3.7', 'Giao diện chi tiết sản phẩm', 5),
    ('3.3.8', 'Bộ nhớ đệm và tối ưu thời gian phản hồi trang danh mục, chi tiết sản phẩm', 4),
], 39, [8, 9, 39, 30, 9], 'Chỉ mục tìm kiếm và bộ nhớ đệm phục vụ chỉ tiêu phản hồi không quá 2 giây (KT02).',
    [124, 125, 126, 127, 128], [130, 132, 134, 136, 138], 121, 122, 123, 139)
notes(s, 50, 'Phân hệ tìm kiếm và hiển thị sản phẩm gồm chỉ mục tìm kiếm, tìm kiếm và gợi ý, lọc và sắp xếp, danh mục '
      'và chi tiết sản phẩm, cùng ba nhóm màn hình. Lập trình và kiểm thử đơn vị là 39 MD trong tổng 95 MD. Dữ liệu '
      'sản phẩm, giá, tồn kho lấy từ Main System qua Selling DB, phân hệ chỉ hiển thị. Bộ nhớ đệm và tối ưu thời gian '
      'phản hồi nhằm đạt chỉ tiêu KT02.', '[S3] WBS nhóm 3, mã 3.1 đến 3.5.')

s = S[5]
phan_he(s, 'Phân hệ thành viên', 91, 4, [
    ('4.3.1', 'Dịch vụ đăng ký thành viên và xác thực địa chỉ thư điện tử', 5),
    ('4.3.2', 'Dịch vụ đăng nhập, phiên làm việc và giới hạn đăng nhập sai', 6),
    ('4.3.3', 'Dịch vụ khôi phục và đổi mật khẩu', 4),
    ('4.3.4', 'Dịch vụ thông tin thành viên và sổ địa chỉ giao hàng', 5),
    ('4.3.5', 'Dịch vụ đồng ý điều khoản và bảo vệ dữ liệu cá nhân của thành viên', 4),
    ('4.3.6', 'Giao diện đăng ký, đăng nhập và khôi phục mật khẩu', 6),
    ('4.3.7', 'Giao diện trang thành viên và sổ địa chỉ', 5),
], 35, [8, 8, 35, 32, 8], ['Bảo mật tài khoản kiểm thử theo OWASP Top 10 (4.4.5).',
                           'Có đồng ý điều khoản và bảo vệ dữ liệu cá nhân.'],
    [151, 152, 153, 154, 155], [157, 159, 161, 163, 165], 148, 149, 150, 166)
nt = shp(s, 166)
nt.left, nt.top, nt.width, nt.height = 8001000, 4937760, 3657600, 822960
notes(s, 45, 'Phân hệ thành viên gồm đăng ký và xác thực thư điện tử, đăng nhập và phiên làm việc có giới hạn đăng '
      'nhập sai, khôi phục mật khẩu, thông tin thành viên, sổ địa chỉ và đồng ý điều khoản bảo vệ dữ liệu cá nhân. '
      'Lập trình và kiểm thử đơn vị là 35 MD trong tổng 91 MD. Kiểm thử phân hệ 32 MD gồm kiểm thử bảo mật tài khoản '
      'theo OWASP Top 10.', '[S3] WBS nhóm 4, mã 4.1 đến 4.5.')

s = S[6]
phan_he(s, 'Phân hệ đặt hàng', 129, 5, [
    ('5.3.1', 'Dịch vụ giỏ hàng', 5),
    ('5.3.2', 'Dịch vụ giữ tồn kho tạm thời khi đặt hàng và giải phóng tồn', 5),
    ('5.3.3', 'Dịch vụ tính phí vận chuyển', 5),
    ('5.3.4', 'Dịch vụ khuyến mại và mã giảm giá', 5),
    ('5.3.5', 'Dịch vụ đặt hàng, xác nhận đơn và thư điện tử xác nhận', 6),
    ('5.3.6', 'Dịch vụ lịch sử, trạng thái đơn hàng và hủy đơn', 5),
    ('5.3.7', 'Dịch vụ yêu cầu đổi trả và hoàn tiền', 6),
    ('5.3.8', 'Giao diện giỏ hàng', 5),
    ('5.3.9', 'Giao diện đặt hàng gồm địa chỉ, phí vận chuyển, mã giảm giá và xác nhận', 6),
    ('5.3.10', 'Giao diện theo dõi đơn hàng, hủy đơn và đổi trả', 5),
], 53, [11, 12, 53, 42, 11], 'Biểu phí vận chuyển, mã giảm giá, giữ tồn, hủy và đổi trả theo quy tắc ABC Wear xác nhận.',
    [178, 179, 180, 181, 182], [184, 186, 188, 190, 192], 175, 176, 177, 193)
notes(s, 55, 'Phân hệ đặt hàng có nỗ lực lớn nhất, 129 MD, vì gồm bảy chức năng: giỏ hàng, giữ tồn, phí vận chuyển, '
      'khuyến mại và mã giảm giá, đặt hàng và xác nhận, theo dõi và hủy đơn, đổi trả và hoàn tiền. Lập trình và kiểm '
      'thử đơn vị 53 MD, kiểm thử phân hệ 42 MD gồm luồng đặt hàng đầu cuối với dữ liệu biên và đồng thời. Các quy tắc '
      'nghiệp vụ do ABC Wear xác nhận tại gói 2.1.5, thanh toán trực tuyến không thuộc phạm vi.', '[S3] WBS nhóm 5, mã 5.1 đến 5.5.')

# ---------------- Slide 8: CSDL và tích hợp
s = S[7]
t(s, 201, 'Cơ sở dữ liệu bán hàng và tích hợp')
header(s, 202, ['Mã', 'Nội dung công việc', 'MD'])
table_rows(s, 202, [
    ['6.1', 'Selling DB, khung ứng dụng và dịch vụ thư điện tử', '22'],
    ['6.2.1, 6.2.2, 6.2.13', 'Đặc tả, thiết kế và xác nhận giao diện tích hợp', '10'],
    ['6.2.3, 6.2.4', 'Đồng bộ sản phẩm, giá và tồn kho từ Main System', '10'],
    ['6.2.5, 6.2.6', 'Chuyển đơn hàng, nhận trạng thái xử lý và vận chuyển', '10'],
    ['6.2.7', 'Giám sát, cảnh báo lỗi đồng bộ và gửi lại', '4'],
    ['6.2.8-6.2.12', 'Giả lập Main System, kiểm thử tích hợp và sửa lỗi', '24'],
    ['6.3', 'Chuyển đổi dữ liệu và nội dung sản phẩm', '13'],
    ['', 'Tổng WBS 6', '93'],
])
cols = shp(s, 202).table.columns
cols[0].width, cols[1].width = 1554480, 5394950 + 822950 - 1554480
tb = shp(s, 205)
tb.top, tb.height = 4572000, 1234440
t(s, 203, '93 MD')
t(s, 204, 'Tổng WBS 6 - CSDL bán hàng và tích hợp')
t(s, 205, ['Năm luồng tích hợp với Main System: sản phẩm và giá, tồn kho, chuyển đơn hàng, trạng thái vận chuyển, giám sát và gửi lại.',
           'Trạng thái vận chuyển nhận qua Main System. Thư điện tử gửi qua máy chủ thư của ABC Wear.'])
t(s, 211, 'Selling DB')
t(s, 212, 'Chuyển đổi dữ liệu')
t(s, 213, 'Máy chủ thư')
notes(s, 50, 'Nhóm 6 gồm Selling DB và nền tảng dùng chung 22 MD, tích hợp Main System 58 MD và chuyển đổi dữ liệu 13 MD, '
      'tổng 93 MD. Tích hợp có năm luồng: đồng bộ sản phẩm và giá, tồn kho, chuyển đơn hàng, nhận trạng thái vận chuyển, '
      'giám sát và gửi lại. Dịch vụ giả lập Main System cho phép phát triển trước khi có môi trường thử ngày 20/5. Đặc '
      'tả giao diện phụ thuộc tài liệu ABC Wear cung cấp trước 14/4, đây là rủi ro hạng A RR02.',
      '[S3] WBS nhóm 6, [S7] RR02, cam kết CK1, CK2 trong [S4].')

# ---------------- Slide 9: sprint
s = S[8]
t(s, 223, '18/3-01/7/2027, 73 ngày làm việc. S3 gồm 13 ngày, các sprint khác 10 ngày. Dự trữ đến cut over 4 ngày.')
table_rows(s, 224, [
    ['S1', '18/3-31/3', '52,5', 'Biên bản khởi động, báo cáo khảo sát, khung ứng dụng triển khai tự động'],
    ['S2', '01/4-14/4', '87', 'Cấu trúc Selling DB, sơ đồ menu, mẫu giao diện ưu tiên'],
    ['S3', '15/4-06/5', '145', 'Tìm kiếm, đăng ký, đăng nhập và giỏ hàng trên môi trường phát triển'],
    ['S4', '07/5-20/5', '112', 'Đăng ký, tìm kiếm, đặt hàng có phí vận chuyển, mã giảm giá'],
    ['S5', '21/5-03/6', '103,5', 'Đặt hàng đầu cuối, chuyển đơn sang Main System thử'],
    ['S6', '04/6-17/6', '79', 'Bản phát hành ứng viên cho UAT, báo cáo kiểm thử hệ thống'],
    ['S7', '18/6-01/7', '78', 'Hệ thống trên môi trường vận hành, biên bản UAT và bàn giao'],
    ['Tổng', '7 sprint', '657', ''],
])
t(s, 225, 'Đường găng 16 gói, dài 52 ngày. Sử dụng 82,5% quỹ huy động 796 MD.')
t(s, 230, 'Trình đặc tả, giao diện ưu tiên (M2, hạn 10/5)')
t(s, 232, '16/6')
t(s, 233, 'Chức năng chính và tích hợp đạt (M3, M4)')
t(s, 236, 'Hoàn thành UAT và kiểm thử vận hành (M5)')
t(s, 239, 'Cut over, nghiệm thu, bàn giao (M6, M7)')
notes(s, 65, 'Kế hoạch gồm bảy sprint, mỗi sprint có sản phẩm trình diễn. S1, S2 khảo sát, đặc tả và nền tảng, S3 dài 13 '
      'ngày để trình đặc tả và giao diện ưu tiên ngày 06/5, hai bên xác nhận trước 10/5 (mốc M2, đợt thanh toán 2). '
      'S4, S5 phát triển và tích hợp, ABC Wear kiểm tra chức năng chính ngày 16/6 (mốc M3, đợt 3). S7 dành cho kiểm '
      'thử chấp nhận 18/6-29/6 và cut over 01/7. Đường găng 16 gói, dự trữ 4 ngày, nhóm chấp nhận phần EMV thời gian '
      'vượt dự trữ và báo ABC Wear khi dự trữ dưới 2 ngày.', '[S4] bảng sprint, mốc và đường găng, [S7] mục dự trữ.')

# ---------------- Slide 10: nhóm chính
s = S[9]
table_rows(s, 248, [
    ['Trần Bá Đạt', 'PM kiêm kiến trúc sư', 'Quản lý dự án (WBS 1), đặc tả và nghiệm thu', '63,33'],
    ['Phạm Tuấn Anh', 'Backend', 'Phân hệ thành viên, đặt hàng (WBS 4, 5)', '59'],
    ['Vũ Đức Minh', 'Frontend kiêm UX/UI', 'Phân hệ tìm kiếm, thiết kế giao diện (WBS 3, 2.3)', '50'],
    ['Ninh Đức Toàn', 'Backend kiêm DevOps', 'CSDL, tích hợp, môi trường, cut over (WBS 6, 8)', '59'],
    ['Nguyễn Thị Phương Liên', 'BA kiêm kiểm thử', 'Đặc tả, kiểm thử hệ thống, UAT (WBS 2, 7)', '68'],
    ['Tổng nhóm chính', '5 người', '', '299,33'],
])
t(s, 249, ['Năm thành viên giữ 5 vị trí nhân sự chủ chốt theo Bảng số 02 của HSMT.',
           'Mỗi hạng mục WBS mức 1, mức 2 có đúng một người chịu trách nhiệm (RACI 38 dòng).',
           'MD giao theo người và sprint, không ô nào vượt quỹ.'])
t(s, 252, 'MD quỹ của 5 thành viên')
t(s, 254, '299,33')
notes(s, 55, 'Nhóm chính gồm năm thành viên, giữ đủ năm vị trí nhân sự chủ chốt theo Bảng số 02: một quản trị dự án, ba '
      'kỹ sư phát triển và một kỹ sư phân tích nghiệp vụ kiêm kiểm thử. Ma trận RACI 38 dòng xác định người chịu trách '
      'nhiệm từng hạng mục. Quỹ của nhóm chính là 365 MD, MD giao 299,33 MD, thấp hơn nhiều so với 657 MD cần thực '
      'hiện, nên phải huy động thêm bảy vị trí.', '[S9] mục RACI và tải theo sprint, [S2] Mẫu 06A.')

# ---------------- Slide 11: vị trí bổ sung
s = S[10]
table_rows(s, 267, [
    ['K', 'Kiến trúc sư kiêm BA, rà soát bảo mật', 'S1-S7', '55,67'],
    ['D1', 'Frontend kiêm backend, sửa lỗi, UAT', 'S2-S7', '51'],
    ['D2', 'Backend, API và kiểm thử hiệu năng', 'S2-S7', '53'],
    ['D3', 'Backend kiêm DevOps, giám sát', 'S2-S7', '56'],
    ['Q1', 'Kiểm thử kỹ thuật, tự động, dữ liệu giả lập', 'S2-S7', '45'],
    ['Q2', 'Kiểm thử chức năng, hiệu năng, an toàn', 'S3-S7', '48'],
    ['D4', 'Backend kiêm DevOps, chuyển đổi, cut over', 'S3-S7', '49'],
    ['Tổng', '7 vị trí bổ sung', '', '357,67'],
])
cd = CategoryChartData()
cd.categories = ['S1', 'S2', 'S3', 'S4', 'S5', 'S6', 'S7']
cd.add_series('Nhóm chính', (5, 5, 5, 5, 5, 5, 5))
cd.add_series('Vị trí bổ sung', (1, 5, 7, 7, 7, 7, 7))
shp(s, 269).chart.replace_data(cd)
tb = shp(s, 268)
for r in list(shp(s, 267).table.rows):
    r.height = 338328
tb.top, tb.height = 5257800, 1005840
t(s, 268, ['Quỹ huy động 796 MD, khối lượng giao 657 MD (82,5%).',
           'Ngày sẵn sàng 139 MD không tính lương (trả theo công sức).',
           'MD giao mỗi người không vượt quỹ của sprint.'])
t(s, 270, 'Quy mô: 6 người ở S1, 10 người ở S2, 12 người ở S3-S7.')
notes(s, 65, 'Bảy vị trí bổ sung từ nhân sự của nhà thầu, tổng 357,67 MD. K tham gia từ S1, D1, D2, D3, Q1 từ S2, Q2 và '
      'D4 từ S3. D4 được bổ sung để lịch 657 MD kịp cut over, đồng thời là nguồn lực điều chuyển khi cần đẩy nhanh. '
      'Tổng quỹ 796 MD so với 657 MD giao, phần 139 MD là ngày sẵn sàng, không tính lương vì nhóm trả lương theo công '
      'sức. Rủi ro vị trí bổ sung không có mặt đúng sprint là RR04 hạng A, biện pháp là ký hợp đồng chậm nhất 10 ngày '
      'làm việc trước sprint huy động.', '[S4] quỹ và tải theo sprint, [S6] đối chiếu lương, [S7] RR04.')

# ---------------- Slide 12: ước lượng
s = S[11]
t(s, 280, 'Sau hiệu chỉnh 657 MD (550,3 và 106,7)')
table_rows(s, 282, [
    ['PM', '49,00', '55', '134,75', 'Đạt'],
    ['BA', '90,00', '43', '193,50', 'K, Liên, Đạt'],
    ['Tester/QA', '160,00', '40', '320,00', 'Q2, Q1, Liên, Toàn, Minh'],
    ['Dev', '318,00', '42', '667,80', 'Tuấn Anh, D1-D4, Minh, Toàn, K, Đạt'],
    ['DevOps', '40,00', '52', '104,00', 'D3, D4, Toàn'],
    ['Tổng', '657,00', '', '1.420,05', '12 người'],
])
c3 = shp(s, 283).table.columns
c3[0].width, c3[2].width, c3[3].width = 2651760, 4663450 + 3200400 + 1737350 - 2651760 - 1005840, 1005840
header(s, 283, ['Nhóm công việc', 'Nhu cầu, MD', 'Người thực hiện, MD giao', 'Số người'])
table_rows(s, 283, [
    ['Quản lý dự án', '49,00', 'Đạt 49', '1'],
    ['Phân tích và kiểm thử (BA + QA)', '250,00', 'Liên 68, Q2 48, Q1 45, K 40,67, Toàn 24, Minh 14, Đạt 10,33', '7'],
    ['Phát triển (Dev + DevOps)', '358,00', 'Tuấn Anh 59, D3 56, D2 53, D1 51, D4 49, Minh 36, Toàn 35, K 15, Đạt 4', '9'],
], giu_cuoi=False)
t(s, 284, 'Dòng dự toán là loại công việc, không phải chức danh: một người có thể làm nhiều loại công việc. '
          'Chi phí mỗi dòng bằng MD chia 20, nhân đơn giá MM.')
notes(s, 65, 'Ước lượng theo danh mục 106 chức năng và 25 đầu việc phi chức năng: 573,6 MD trước hiệu chỉnh, trừ tái sử '
      'dụng còn 550,3 MD, cộng 106,7 MD phi chức năng, tổng 657 MD, khoảng sai số từ -15% đến +30%. MD được chia theo '
      'năm loại công việc để áp đơn giá trọn gói theo người-tháng, ví dụ Dev 318 MD chia 20 nhân 42 triệu đồng được '
      '667,8 triệu đồng. Bảng dưới cho thấy người thực hiện từng nhóm công việc theo giao việc, cộng đúng nhu cầu.',
      '[S3] mục căn cứ ước lượng, [S6] bảng nhân công theo loại công việc, [S9].')

# ---------------- Slide 13: chi phí
s = S[12]
table_rows(s, 293, [
    ['A - Nhân công theo loại công việc', '1.420,05'],
    ['B - Môi trường, công cụ, bảo hành 12 tháng và dự phòng', '150,00'],
    ['Chi phí đã biết (A + B)', '1.570,05'],
    ['C - Chi phí chung và lợi nhuận', '179,95'],
    ['Giá chào trọn gói (A + B + C)', '1.750,00'],
])
t(s, 294, 'Đối chiếu dòng A với chi phí lương theo công sức')
header(s, 295, ['Nhóm nhân sự', 'MD tính lương', 'Triệu đồng'])
table_rows(s, 295, [
    ['5 thành viên nhóm', '299,33', '561,69'],
    ['7 vị trí bổ sung', '357,67', '624,15'],
    ['Tổng', '657', '1.185,84'],
])
cd = CategoryChartData()
cd.categories = ['Lương và bảo hiểm theo công sức: 1.185,84', 'A - Nhân công: 1.420,05', 'A + B - Chi phí đã biết: 1.570,05',
                 'Giá chào: 1.750,00', 'Giá gói thầu: 2.000,00']
cd.add_series('Triệu đồng', (1185.84, 1420.05, 1570.05, 1750.0, 2000.0))
shp(s, 297).chart.replace_data(cd)
t(s, 296, 'Lương tham chiếu: trung vị ITviec 2025-2026 theo vị trí, hệ số bảo hiểm 1,235, 22 ngày mỗi tháng. '
          'Không phải lương thực tế của từng người.')
t(s, 298, ['Dòng A cao hơn chi phí lương 234,21 triệu đồng (16,5%), đủ chi trả khi năng suất từ 83,5%.',
           'Dự phòng 50 triệu đồng thấp hơn EMV 147 triệu đồng, phần vượt bù từ khoản còn lại sau lương.',
           'Giá chào thấp hơn giá gói thầu 250 triệu đồng (12,5%).',
           'Không chịu thuế GTGT, khoản C là 179,95 triệu đồng (10,28%).'])
notes(s, 75, 'Dòng A là 1.420,05 triệu đồng theo đơn giá trọn gói từng loại công việc, cộng 150 triệu đồng khoản B gồm '
      'môi trường, công cụ, bảo hành và 50 triệu đồng dự phòng, chi phí đã biết 1.570,05 triệu đồng. Để kiểm tra dòng A, '
      'nhóm tính lương theo công sức của từng người với lương tham chiếu và bảo hiểm, được 1.185,84 triệu đồng cho 657 MD. '
      'Dòng A cao hơn 234,21 triệu đồng, đủ chi trả khi năng suất từ 83,5%. Dự phòng thấp hơn EMV rủi ro 147 triệu đồng, '
      'nhóm quyết định giữ 50 triệu đồng và bù phần vượt từ khoản còn lại sau lương 414,16 triệu đồng. Dòng tiền lũy kế '
      'âm từ S1 đến S6, thấp nhất khoảng 464,7 triệu đồng cuối S4, nhóm chuẩn bị vốn lưu động 500 triệu đồng. Giá chào '
      '1.750 triệu đồng thấp hơn giá gói thầu 12,5%, dịch vụ phần mềm không chịu thuế GTGT.',
      '[S6] Cost_estimate.xlsx và thuyết minh dự toán mục 3 đến 8, khoản 21 Điều 5 Luật Thuế GTGT số 48/2024/QH15.')

# ---------------- Slide 14: rủi ro và truyền thông
s = S[13]
t(s, 307, '21 rủi ro (19 đe dọa, 2 cơ hội), 5 rủi ro hạng A. EMV 147 triệu đồng (110-183), dự phòng 50 triệu đồng.')
header(s, 308, ['Rủi ro hạng A tiêu biểu', 'Biện pháp phòng ngừa', 'Chủ trì'])
table_rows(s, 308, [
    ['RR02 Tài liệu giao diện, môi trường thử Main System đến muộn', 'Đề nghị bàn giao tại họp khởi động, dịch vụ giả lập, thống nhất đặc tả bằng biên bản', 'Toàn'],
    ['RR04 Vị trí bổ sung không có mặt đúng sprint huy động', 'Ký hợp đồng chậm nhất 10 ngày làm việc trước sprint, danh sách người dự phòng', 'Đạt'],
    ['RR10 Lỗi UAT vượt khả năng khắc phục trong 18/6-29/6', 'Khắc phục lỗi ngay từng đợt kiểm thử, thống nhất mức nghiêm trọng từ đầu UAT', 'Liên'],
    ['RR13 Lương vị trí bổ sung cao hơn lương tham chiếu', 'Thống nhất lương khi ký hợp đồng, so với lương tham chiếu trước khi ký', 'Đạt'],
], giu_cuoi=False)
t(s, 312, 'Báo cáo cuối sprint')
t(s, 313, 'Gửi đầu mối ABC Wear ngày cuối sprint: SPI, CPI, mốc, rủi ro, thay đổi và vấn đề cần quyết định.')
t(s, 318, 'Họp tiến độ hằng tuần với ABC Wear. Mỗi sprint có lập kế hoạch, trình diễn và họp cải tiến.')
t(s, 319, 'Rà soát rủi ro đường găng hằng tuần, báo ABC Wear khi dự trữ còn dưới 2 ngày.')
notes(s, 60, 'Sổ rủi ro có 21 rủi ro, trong đó 2 cơ hội, xếp theo ma trận 5x5 quy đổi ba hạng. Năm rủi ro hạng A, bảng '
      'nêu bốn rủi ro tiêu biểu với biện pháp lấy nguyên từ sổ rủi ro. Tổng EMV 147 triệu đồng, khoảng 110 đến 183 '
      'triệu đồng. EMV thời gian trên đường găng 5,4 ngày, lớn hơn dự trữ 4 ngày, nhóm chấp nhận và theo dõi hằng tuần. '
      'Báo cáo tiến độ gửi ngày cuối sprint, họp tiến độ hằng tuần với ABC Wear.',
      '[S7] sổ rủi ro, ma trận và mục dự phòng, [S5] mục truyền thông và báo cáo.')

# ---------------- Slide 15: chuyển giao
s = S[14]
t(s, 330, 'UAT')
t(s, 331, '18/6-29/6/2027')
t(s, 332, 'Hai luồng kiểm thử chấp nhận cùng ABC Wear, khắc phục lỗi nghiêm trọng và cao.')
t(s, 336, 'Đến 30/6/2027')
t(s, 337, 'Diễn tập triển khai, khôi phục, danh sách kiểm tra cut over được xác nhận.')
t(s, 342, 'Chuyển đổi dữ liệu chính thức ngày 30/6, kiểm tra chức năng chính sau triển khai.')
t(s, 346, '01/7/2027')
t(s, 347, 'Mã nguồn, tài liệu kỹ thuật, hướng dẫn vận hành, biên bản bàn giao (M7).')
t(s, 352, 'Bảo hành, hỗ trợ tăng cường 02/7-15/7, bảo lãnh bảo hành 5%.')
t(s, 353, 'Ba đường cơ sở, kiểm soát bằng SPI, CPI cuối mỗi sprint')
t(s, 356, 'Phạm vi: 144 gói WBS, đặc tả tại M2')
t(s, 359, 'Tiến độ: cut over 01/7/2027, dự trữ 4 ngày')
t(s, 361, '1.420,05 triệu')
t(s, 362, 'Chi phí: BAC bằng khoản A')
notes(s, 45, 'Điều kiện hoàn thành xác định theo tiêu chí. Kiểm thử chấp nhận và kiểm thử vận hành 18/6-29/6 phải đạt, '
      'lỗi nghiêm trọng và cao được khắc phục. Trước cut over đã diễn tập chuyển đổi dự phòng, diễn tập triển khai và '
      'có danh sách kiểm tra được ABC Wear xác nhận. Ngày 30/6 chuyển đổi dữ liệu chính thức, ngày 01/7 đưa hệ thống '
      'vào vận hành và bàn giao. Ba đường cơ sở là phạm vi 657 MD, tiến độ đến 01/7 và chi phí BAC 1.420,05 triệu đồng '
      'bằng khoản A, giá hợp đồng không dùng làm đường cơ sở chi phí.', '[S4] gói 7.3, 8.2, 8.3, [S5] mục đường cơ sở, [S2] mục bảo hành.')

# ---------------- Slide 16: chất lượng và thay đổi
s = S[15]
t(s, 373, 'Áp dụng cho mọi đề nghị thay đổi, chấm điểm tác động từ 1 đến 10.')
t(s, 377, 'Đánh giá WBS, MD, chi phí, mốc: 1 ngày làm việc (nhỏ), 3 ngày làm việc (trung bình, lớn).')
t(s, 381, 'Nhỏ: quản trị dự án. Trung bình: thêm đầu mối ABC Wear. Lớn: ban kiểm soát, phụ lục hợp đồng.')
t(s, 385, 'Cập nhật tài liệu theo sơ đồ phụ thuộc, tăng phiên bản, đóng CR khi ABC Wear kiểm tra đạt.')
t(s, 390, 'Sản phẩm đúng đặc tả, mã nguồn đạt quy chuẩn, kiểm thử đạt, trình diễn cuối sprint được ABC Wear xác nhận.')
t(s, 395, 'Năm chỉ tiêu KT01-KT05: sẵn sàng 99,5%, phản hồi không quá 2 giây, 500 người dùng đồng thời, OWASP Top 10, sao lưu.')
t(s, 400, 'Đường cơ sở tại M2, nhánh và nhãn phiên bản, kiểm tra cấu hình trước M3 và cut over, sổ phiên bản tự động.')
notes(s, 55, 'Kế hoạch chất lượng dùng năm chỉ tiêu kỹ thuật định lượng KT01 đến KT05 và tiêu chí hoàn thành cho gói, '
      'sprint và mốc. Mọi đề nghị thay đổi được chấm điểm tác động và xếp ba mức. Mức nhỏ là thay đổi nội bộ của nhà '
      'thầu, mức trung bình cần đầu mối ABC Wear đồng phê duyệt, mức lớn chạm cut over, mốc thanh toán, đặc tả đã xác '
      'nhận, phạm vi hoặc giá thì ban kiểm soát thay đổi xem xét và người đại diện hai bên ký phụ lục. Quản lý cấu hình '
      'có đường cơ sở tại M2, sổ phiên bản đăng ký tự động theo băm nội dung.',
      '[S8] bảng phân loại, thẩm quyền và quản lý cấu hình, [S5] mục chất lượng, [S10].')

# ---------------- Slide 17: tổng kết
s = S[16]
t(s, 411, 'Mẫu 5A, giá gói thầu 2 tỷ đồng, bảng dữ liệu 23 mục, điều kiện cụ thể 25 mục')
t(s, 414, '8 nhóm công việc, 144 gói, tổng 657 MD')
t(s, 417, '7 sprint, 5 người chính và 7 vị trí bổ sung, quỹ 796 MD')
t(s, 420, 'Ước lượng theo chức năng, BAC 1,42 tỷ đồng, giá chào 1,75 tỷ đồng')
t(s, 423, 'Rủi ro, thay đổi, cấu hình, phiên bản, báo cáo và điều kiện bàn giao')
t(s, 424, 'Các tài liệu dùng chung một nguồn dữ liệu, liên kết theo mã WBS và đều đạt kiểm tra lớp A, lớp B.')
tf = s.notes_slide.notes_text_frame
set_text(tf, ['Thời lượng: 30 giây',
              'Kết quả gồm hồ sơ mời thầu theo Mẫu 5A, hồ sơ dự thầu và bộ tài liệu kế hoạch liên kết từ WBS đến sprint, '
              'nhân sự, dự toán và kiểm soát dự án. Khi triển khai thực tế, các giả định về nhân sự bổ sung, tài liệu và '
              'môi trường Main System, quy tắc nghiệp vụ của ABC Wear cần được xác nhận lại. Nhóm xin kết thúc phần báo '
              'cáo và sẵn sàng trả lời câu hỏi.',
              'Nguồn đối chiếu: tổng hợp các đầu mục trên.',
              'Danh mục nguồn (thư mục ket_qua):',
              '[S1] 01_HSMT, Hồ sơ mời thầu',
              '[S2] 02_HSDT, Hồ sơ dự thầu',
              '[S3] 03_WBS, WBS tổng hợp và 8 tài liệu nhóm',
              '[S4] 04_Master_schedule, Master schedule',
              '[S5] 05_Master_plan, Kế hoạch tổng thể',
              '[S6] 06_Cost_estimate, Cost estimate và thuyết minh dự toán',
              '[S7] 07_Risk_management_plan, Kế hoạch quản lý rủi ro',
              '[S8] 08_Change_management_plan, Kế hoạch quản lý thay đổi',
              '[S9] 10_Responsibilities, Danh sách phân chia công việc',
              '[S10] 15_Quan_ly_phien_ban, Tài liệu quản lý phiên bản'])

# ---------------- Sắp lại: slide "6. Chất lượng và thay đổi" đứng trước "7. Rủi ro"
lst = p.slides._sldIdLst
ids = list(lst)
cl = ids[15]
lst.remove(cl)
lst.insert(13, cl)

p.save(OUT)
print('Đã ghi', OUT)
